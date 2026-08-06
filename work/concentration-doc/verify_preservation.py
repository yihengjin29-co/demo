from __future__ import annotations

import re
from pathlib import Path

from docx import Document


root = Path(__file__).parent
source = Document(root / "source.docx")
modified = Document(root / "金控并表系统功能需求说明书_风险并表板块_补充集中度.docx")


def normalize_paragraph(text: str):
    if re.match(r"^\s*\d+(?:\.\d+)*\s", text) and "\t" in text:
        return re.sub(r"\t\d+\s*$", "", text)
    return text


source_paragraphs = [normalize_paragraph(paragraph.text) for paragraph in source.paragraphs]
modified_paragraphs = [normalize_paragraph(paragraph.text) for paragraph in modified.paragraphs]
cursor = 0
missing_paragraphs = []
for text in source_paragraphs:
    try:
        cursor = modified_paragraphs.index(text, cursor) + 1
    except ValueError:
        missing_paragraphs.append(text)

source_tables = [[cell.text for row in table.rows for cell in row.cells] for table in source.tables]
modified_tables = [[cell.text for row in table.rows for cell in row.cells] for table in modified.tables]
cursor = 0
missing_tables = []
for table_text in source_tables:
    try:
        cursor = modified_tables.index(table_text, cursor) + 1
    except ValueError:
        missing_tables.append(table_text[:10])

print(f"source_paragraphs={len(source.paragraphs)} modified_paragraphs={len(modified.paragraphs)} missing={len(missing_paragraphs)}")
print(f"source_tables={len(source.tables)} modified_tables={len(modified.tables)} missing={len(missing_tables)}")
print(f"source_images={len(source.inline_shapes)} modified_images={len(modified.inline_shapes)}")
if missing_paragraphs:
    print("missing_paragraph_samples=", missing_paragraphs[:5])
if missing_tables:
    print("missing_table_samples=", missing_tables[:3])
