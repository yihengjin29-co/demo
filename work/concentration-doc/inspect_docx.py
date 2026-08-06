from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


DOCX = Path(__file__).with_name("source.docx")


def iter_blocks(document):
    body = document.element.body
    paragraph_by_element = {paragraph._p: paragraph for paragraph in document.paragraphs}
    table_by_element = {table._tbl: table for table in document.tables}
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "P", paragraph_by_element[child]
        elif child.tag == qn("w:tbl"):
            yield "T", table_by_element[child]


document = Document(DOCX)
print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} inline_shapes={len(document.inline_shapes)} sections={len(document.sections)}")
print("\n=== BODY OUTLINE ===")
for index, (kind, block) in enumerate(iter_blocks(document)):
    if kind == "P":
        paragraph: Paragraph = block
        text = paragraph.text.strip().replace("\n", " | ")
        has_drawing = bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))
        if text or has_drawing:
            print(f"{index:04d} P [{paragraph.style.name}] drawing={has_drawing} :: {text[:240]}")
    else:
        table: Table = block
        rows = []
        for row in table.rows[:3]:
            rows.append(" || ".join(cell.text.strip().replace("\n", " / ")[:120] for cell in row.cells))
        print(f"{index:04d} T rows={len(table.rows)} cols={len(table.columns)} :: {' >>> '.join(rows)}")

print("\n=== KEYWORD MATCHES ===")
keywords = ("集中度", "风险监测", "图片", "截图", "界面样式", "权限清单", "字段规则")
for paragraph_index, paragraph in enumerate(document.paragraphs):
    if any(keyword in paragraph.text for keyword in keywords):
        print(f"P{paragraph_index:04d} [{paragraph.style.name}] {paragraph.text.strip()}")
for table_index, table in enumerate(document.tables):
    text = " | ".join(cell.text for row in table.rows for cell in row.cells)
    if any(keyword in text for keyword in keywords):
        print(f"T{table_index:03d} rows={len(table.rows)} cols={len(table.columns)} {text[:500]}")

with ZipFile(DOCX) as archive:
    media = [name for name in archive.namelist() if name.startswith("word/media/")]
    print(f"\nmedia_count={len(media)}")
    for name in media:
        info = archive.getinfo(name)
        print(f"{name} {info.file_size}")
