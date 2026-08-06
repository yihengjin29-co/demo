from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


SOURCE = Path(__file__).with_name("source.docx")
OUTPUT = Path(__file__).with_name("金控并表系统功能需求说明书_风险并表板块_补充集中度.docx")


def new_paragraph_like(reference: Paragraph, text: str = ""):
    element = OxmlElement("w:p")
    if reference._p.pPr is not None:
        element.append(deepcopy(reference._p.pPr))
    paragraph = Paragraph(element, reference._parent)
    if text:
        paragraph.add_run(text)
    return element


def new_screenshot_placeholder(reference: Paragraph):
    element = new_paragraph_like(reference)
    p_pr = element.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:line"), "3968")
    spacing.set(qn("w:lineRule"), "exact")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "120")
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "D9E2F0")
        borders.append(border)
    p_pr.append(borders)
    return element


def replace_cell_text(cell, text: str):
    text_nodes = cell._tc.xpath(".//w:t")
    if text_nodes:
        text_nodes[0].text = text
        if text.startswith(" ") or text.endswith(" "):
            text_nodes[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for node in text_nodes[1:]:
            node.text = ""
    else:
        cell.paragraphs[0].add_run(text)


def cloned_table(document, template: Table, row_specs, row_templates):
    table_element = deepcopy(template._tbl)
    for row in list(table_element.tr_lst):
        table_element.remove(row)
    table = Table(table_element, document._body)
    for kind, values in row_specs:
        table_element.append(deepcopy(row_templates[kind]._tr))
        row = table.rows[-1]
        if kind in {"title", "section"}:
            replace_cell_text(row.cells[0], values[0])
        else:
            for cell, value in zip(row.cells, values):
                replace_cell_text(cell, value)
    return table_element


document = Document(SOURCE)

paragraphs = document.paragraphs
find_paragraph = lambda value: next(paragraph for paragraph in paragraphs if paragraph.text.strip() == value)
heading3_ref = find_paragraph("二级菜单—预警提示与处置")
heading4_ref = find_paragraph("预警提示与处置")
interface_ref = find_paragraph("界面样式：")
rules_label_ref = find_paragraph("界面规则")
permission_label_ref = find_paragraph("权限清单")
rule_ref = next(paragraph for paragraph in paragraphs if paragraph.text.startswith("查询区–页面上方为预警事项查询栏"))
screenshot_ref = next(paragraph for paragraph in paragraphs if paragraph._p.xpath(".//w:drawing | .//w:pict"))
blank_ref = next(paragraph for paragraph in paragraphs if not paragraph.text and not paragraph._p.xpath(".//w:drawing | .//w:pict"))
insert_before = find_paragraph("一级菜单——指标管理")._p

permission_template = document.tables[45]
field_template = document.tables[46]
operation_template = document.tables[47]
permission_rows = {"header": permission_template.rows[0], "data": permission_template.rows[1]}
field_rows = {"title": field_template.rows[0], "header": field_template.rows[1], "section": field_template.rows[2], "data": field_template.rows[3]}
operation_rows = {"title": operation_template.rows[0], "header": operation_template.rows[1], "data": operation_template.rows[2]}


def permissions(*rows):
    return cloned_table(
        document,
        permission_template,
        [("header", ["权限事项", "集团", "金控公司", "各金融机构"])] + [("data", list(row)) for row in rows],
        permission_rows,
    )


def fields(rows):
    specs = [
        ("title", ["一、操作界面字段规则及要求"]),
        ("header", ["字段名称", "字段类型", "操作方式", "字段值", "备注说明"]),
    ]
    for row in rows:
        specs.append(("section" if len(row) == 1 else "data", list(row)))
    return cloned_table(document, field_template, specs, field_rows)


def operations(rows):
    return cloned_table(
        document,
        operation_template,
        [("title", ["二、操作功能说明"]), ("header", ["操作功能", "说明"])] + [("data", list(row)) for row in rows],
        operation_rows,
    )


def add_screen(title, rules, permission_data, field_data, operation_data, page_break=False):
    heading = new_paragraph_like(heading4_ref, title)
    if page_break:
        heading_p_pr = heading.get_or_add_pPr()
        if heading_p_pr.find(qn("w:pageBreakBefore")) is None:
            heading_p_pr.append(OxmlElement("w:pageBreakBefore"))
    blocks = [
        heading,
        new_paragraph_like(interface_ref, "界面样式："),
        new_screenshot_placeholder(screenshot_ref),
        new_paragraph_like(rules_label_ref, "界面规则"),
        *[new_paragraph_like(rule_ref, text) for text in rules],
        new_paragraph_like(permission_label_ref, "权限清单"),
        permissions(*permission_data),
        new_paragraph_like(blank_ref),
        fields(field_data),
        new_paragraph_like(blank_ref),
        operations(operation_data),
        new_paragraph_like(blank_ref),
    ]
    for block in blocks:
        insert_before.addprevious(block)


insert_before.addprevious(new_paragraph_like(heading3_ref, "二级菜单—集中度风险监测"))

add_screen(
    "集中度风险监测 – 监测总览",
    [
        "查询区–页面上方为集中度监测查询栏——支持按照数据期次和集中度类型筛选；查询结果仅影响当前已配置展示的集中度卡片。",
        "配置区–页面右上角设置“集中度类型设置”入口——可从单一客户集中度、单一集团客户集中度、行业集中度和区域集中度中选择展示类型并调整顺序；默认展示单一集团客户、行业和区域集中度。",
        "卡片区–每个集中度类型以独立卡片展示监测对象数量及红、黄、绿灯数量，灯号使用彩色圆点和文字表达；点击“查看详情”进入对应集中度汇总页面。",
        "数据规则–展示配置按当前角色保存；至少保留一种集中度类型。集团和金控公司查看并表范围数据，各金融机构仅查看本机构权限范围内数据。",
    ],
    [
        ("查看集中度监测总览", "√", "√", "√（本机构）"),
        ("维护展示类型配置", "√", "√", "√（个人配置）"),
        ("查看集中度汇总", "√", "√", "√（本机构）"),
        ("导出集中度数据", "√", "√", "√（本机构）"),
    ],
    [
        ("查询区",),
        ("数据期次", "时间", "下拉选择", "页面可选项", "选择需要查看的集中度数据期次。"),
        ("集中度类型", "文本", "下拉选择", "页面可选项", "从已配置展示的集中度类型中筛选。"),
        ("配置区",),
        ("集中度类型设置", "文本", "按钮操作", "页面操作入口", "点击后打开展示集中度类型配置窗口。"),
        ("展示类型", "文本", "复选框", "单一客户、单一集团客户、行业、区域集中度", "至少选择一项；默认勾选单一集团客户、行业和区域集中度。"),
        ("展示顺序", "数字", "按钮操作", "上移、下移", "调整已选择集中度卡片的展示顺序。"),
        ("卡片区",),
        ("集中度类型", "文本", "自动反显", "页面记录值", "展示当前卡片对应的集中度类型名称。"),
        ("监测对象数量", "数字", "系统计算", "非负整数", "统计当前期次、当前类型下的集中对象数量。"),
        ("红灯数量", "数字", "系统计算", "非负整数", "以红色圆点和黑色文字展示红灯对象数量。"),
        ("黄灯数量", "数字", "系统计算", "非负整数", "以黄色圆点和黑色文字展示黄灯对象数量。"),
        ("绿灯数量", "数字", "系统计算", "非负整数", "以绿色圆点和黑色文字展示绿灯对象数量。"),
        ("当前期次", "时间", "自动反显", "页面记录值", "展示卡片数据所属指标期次。"),
        ("查看详情", "文本", "按钮操作", "页面操作入口", "进入当前集中度类型的汇总页面。"),
    ],
    [
        ("查询", "通过点击“查询”，系统按照数据期次和集中度类型刷新当前监测卡片。"),
        ("重置", "通过点击“重置”，系统恢复默认数据期次和全部已配置展示类型。"),
        ("集中度类型设置", "通过点击“集中度类型设置”，系统打开配置窗口，支持选择展示类型并调整顺序。"),
        ("保存配置", "通过点击“保存配置”，系统保存当前角色的展示类型及顺序，并刷新总览卡片。"),
        ("查看详情", "通过点击“查看详情”，系统进入对应集中度汇总页面。"),
    ],
)

add_screen(
    "集中度风险监测 – 集中度汇总",
    [
        "查询区–页面上方根据集中度类型动态展示集中对象名称、指标期次和集中度区间；不设置所属机构和亮灯情况查询条件。",
        "列表区–页面主体展示集中对象名称、集中度、投融资或业务余额、当期亮灯和指标期次，默认按照集中度从高到低排序。",
        "口径规则–单一客户按客户汇总，单一集团客户按集团客户汇总，行业和区域分别按行业、区域汇总；各集中对象应包含多个往来机构和业务类型。",
        "下钻规则–点击“查看详情”进入对应集中对象分析页面；集团和金控公司查看并表口径，各金融机构仅查看本机构权限范围内记录。",
    ],
    [
        ("查询集中度汇总", "√", "√", "√（本机构）"),
        ("查看集中度分析", "√", "√", "√（本机构）"),
        ("导出汇总数据", "√", "√", "√（本机构）"),
    ],
    [
        ("查询区",),
        ("集中对象名称", "文本", "自主输入", "100个字符以内", "根据类型动态显示客户、集团客户、行业或区域名称，并支持模糊查询。"),
        ("指标期次", "时间", "下拉选择", "页面可选项", "选择集中度指标所属期次。"),
        ("集中度区间", "数字", "区间输入", "最低值至最高值（%）", "按照集中度百分比区间筛选记录。"),
        ("列表区",),
        ("序号", "数字", "系统计算", "正整数", "系统根据当前列表展示顺序自动计算。"),
        ("集中对象名称", "文本", "自动反显", "页面记录值", "根据类型展示客户、集团客户、行业或区域名称。"),
        ("集中度", "数字", "系统计算", "百分比", "展示集中对象业务余额与计量分母计算形成的集中度。"),
        ("投融资或业务余额", "数字", "系统计算", "金额（亿元）", "单一客户和集团客户展示投融资余额；行业和区域展示业务余额。"),
        ("当期亮灯", "文本", "自动反显", "红灯、黄灯、绿灯", "根据生效阈值展示当前期次灯号。"),
        ("指标期次", "时间", "自动反显", "页面记录值", "展示当前集中度记录的数据期次。"),
        ("查看详情", "文本", "链接操作", "页面操作入口", "点击后进入当前集中对象分析页面。"),
        ("分页区",),
        ("分页", "数字", "链接操作", "总数、每页条数、页码", "支持每页条数、上一页、下一页及页面跳转。"),
    ],
    [
        ("查询", "通过点击“查询”，系统按照集中对象名称、指标期次和集中度区间筛选汇总记录。"),
        ("重置", "通过点击“重置”，系统清空查询条件并恢复默认期次。"),
        ("导出", "通过点击“导出”，系统导出当前权限范围和查询条件下的集中度汇总数据。"),
        ("查看详情", "通过点击“查看详情”，系统进入集中对象分析页面。"),
        ("返回总览", "通过点击“返回总览”，系统返回集中度风险监测总览页面。"),
        ("分页", "通过点击分页控件，系统切换当前列表记录。"),
    ],
    page_break=True,
)

add_screen(
    "集中度风险监测 – 集中度分析",
    [
        "基本信息区–页面上方展示集中对象名称、当前集中度、业务余额、计量分母、当前亮灯、往来机构、业务类型和指标期次。",
        "趋势区–页面中部展示集中度与业务余额历史趋势，并同步标识黄灯、红灯阈值。",
        "构成分析区–支持“按往来机构”和“按业务类型”两个页签；按机构至少展示国际AMC、浦发银行、国泰海通证券等三家往来机构，按业务类型至少展示业务类型A至D。",
        "一致性规则–基本信息中的往来机构数量应与按机构构成项一致；机构构成、业务类型构成的余额合计均应等于集中对象业务余额。",
        "下钻规则–点击任一机构或业务类型的“查看业务明细”，进入相应维度的底层业务明细页面；可查看与当前集中度记录关联的预警事项。",
    ],
    [
        ("查看集中度基本信息", "√", "√", "√（本机构）"),
        ("查看趋势及构成分析", "√", "√", "√（本机构）"),
        ("查看业务明细", "√", "√", "√（本机构）"),
        ("查看关联预警", "√", "√", "√（本机构）"),
    ],
    [
        ("集中度基本信息区",),
        ("集中对象名称", "文本", "自动反显", "页面记录值", "展示当前客户、集团客户、行业或区域名称。"),
        ("当前集中度", "数字", "系统计算", "百分比", "展示当前期次集中度计算结果。"),
        ("业务余额", "数字", "系统计算", "金额（亿元）", "展示当前集中对象全部底层业务余额合计。"),
        ("计量分母", "文本/数字", "自动反显", "分母名称及金额", "展示集中度计算采用的计量分母及其数值。"),
        ("当前亮灯", "文本", "自动反显", "红灯、黄灯、绿灯", "展示当前集中度按照阈值判断的灯号。"),
        ("往来机构", "文本", "自动反显", "机构名称集合", "展示当前集中对象涉及的全部往来机构，不得仅展示单一机构。"),
        ("业务类型", "文本", "自动反显", "业务类型集合", "展示当前集中对象涉及的业务类型。"),
        ("指标期次", "时间", "自动反显", "页面记录值", "展示当前记录的数据期次。"),
        ("趋势区",),
        ("历史期次", "时间", "自动反显", "页面记录值", "展示历史趋势横轴期次。"),
        ("集中度趋势", "数字", "图形展示", "百分比", "展示各历史期次集中度变化。"),
        ("业务余额趋势", "数字", "图形展示", "金额（亿元）", "展示各历史期次业务余额变化。"),
        ("预警阈值", "数字", "图形展示", "黄灯、红灯阈值", "在趋势图中标识当前生效预警阈值。"),
        ("构成分析区",),
        ("构成维度", "文本", "页签切换", "按往来机构、按业务类型", "切换机构构成和业务类型构成。"),
        ("构成项", "文本", "自动反显", "机构名称或业务类型", "展示当前维度下的组成项。"),
        ("业务余额", "数字", "系统计算", "金额（亿元）", "汇总当前组成项对应的底层业务余额。"),
        ("占比", "数字", "系统计算", "百分比", "当前组成项业务余额除以集中对象业务余额。"),
        ("操作", "文本", "链接操作", "查看业务明细", "进入当前组成项对应的业务明细页面。"),
    ],
    [
        ("切换构成维度", "通过点击“按往来机构”或“按业务类型”，系统切换饼图和构成明细。"),
        ("查看业务明细", "通过点击构成项后的“查看业务明细”，系统携带当前机构或业务类型条件进入底层业务明细。"),
        ("查看全部业务明细", "通过点击“查看全部业务明细”，系统进入当前集中对象全部底层业务记录。"),
        ("查看关联预警", "通过点击“查看关联预警”，系统进入当前集中度记录关联的预警处置页面。"),
        ("返回汇总", "通过点击“返回汇总”，系统返回当前集中度类型的汇总页面。"),
    ],
    page_break=True,
)

add_screen(
    "集中度风险监测 – 业务明细",
    [
        "查询区–页面上方展示并锁定当前集中对象，并支持按照所属机构、集团客户、业务类型、行业、区域、业务发生日期和数据期次筛选。",
        "列表区–页面主体展示构成分析下钻形成的底层业务记录；业务余额为单笔业务余额，不得以机构或业务类型汇总金额替代。",
        "数据规则–每个集中对象至少涉及三家往来机构，每家机构至少形成五条业务明细；明细余额汇总应与对应机构构成余额一致。",
        "范围规则–集团和金控公司可查看并表范围业务明细，各金融机构仅查看本机构权限范围内明细；导出结果沿用当前查询条件和数据权限。",
    ],
    [
        ("查询业务明细", "√", "√", "√（本机构）"),
        ("查看业务明细", "√", "√", "√（本机构）"),
        ("导出业务明细", "√", "√", "√（本机构）"),
    ],
    [
        ("查询区",),
        ("集中对象名称", "文本", "自动反显", "页面记录值", "展示并锁定当前下钻的集中对象名称。"),
        ("所属机构", "文本", "下拉选择", "页面可选项", "筛选往来机构；由机构维度下钻时自动带入并锁定。"),
        ("集团客户名称", "文本", "下拉选择", "页面可选项", "筛选集团客户；集团客户集中度下钻时自动带入并锁定。"),
        ("业务类型", "文本", "下拉选择", "页面可选项", "筛选业务类型；由业务类型维度下钻时自动带入并锁定。"),
        ("行业", "文本", "下拉选择", "页面可选项", "筛选行业；行业集中度下钻时自动带入并锁定。"),
        ("区域", "文本", "下拉选择", "页面可选项", "筛选区域；区域集中度下钻时自动带入并锁定。"),
        ("业务发生日期", "时间", "日期区间选择", "起始日期至结束日期", "按照业务发生日期范围筛选明细。"),
        ("数据期次", "时间", "下拉选择", "页面可选项", "选择业务明细所属数据期次。"),
        ("列表区",),
        ("序号", "数字", "系统计算", "正整数", "系统根据当前列表展示顺序自动计算。"),
        ("业务发生日期", "时间", "自动反显", "YYYY-MM-DD", "展示单笔业务发生日期。"),
        ("集中对象名称", "文本", "自动反显", "页面记录值", "展示当前业务归属的客户、集团客户、行业或区域。"),
        ("所属机构", "文本", "自动反显", "页面记录值", "展示单笔业务所属往来机构。"),
        ("业务类型", "文本", "自动反显", "页面记录值", "展示单笔业务类型。"),
        ("单笔业务余额", "数字", "自动反显", "金额（亿元）", "展示底层单笔业务金额，不展示汇总金额。"),
        ("行业", "文本", "自动反显", "页面记录值", "展示单笔业务所属行业。"),
        ("区域", "文本", "自动反显", "页面记录值", "展示单笔业务所属区域。"),
        ("数据期次", "时间", "自动反显", "页面记录值", "展示单笔业务所属数据期次。"),
        ("数据来源", "文本", "自动反显", "页面记录值", "展示单笔业务数据来源。"),
        ("分页区",),
        ("分页", "数字", "链接操作", "总数、每页条数、页码", "支持每页条数、上一页、下一页及页面跳转。"),
    ],
    [
        ("查询", "通过点击“查询”，系统按照当前集中对象及已选择的明细条件筛选底层业务记录。"),
        ("重置", "通过点击“重置”，系统恢复下钻时带入的固定条件和默认数据期次。"),
        ("导出", "通过点击“导出”，系统导出当前权限范围和查询条件下的业务明细。"),
        ("返回分析", "通过点击“返回分析”，系统返回当前集中对象的集中度分析页面。"),
        ("分页", "通过点击分页控件，系统切换当前业务明细记录。"),
    ],
    page_break=True,
)

settings = document.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

document.save(OUTPUT)
print(OUTPUT)
