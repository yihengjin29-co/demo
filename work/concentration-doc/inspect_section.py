from __future__ import annotations

from pathlib import Path
import sys

from docx import Document
from docx.oxml.ns import qn


document = Document(Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("source.docx"))
print("=== INLINE SHAPES ===")
for index, shape in enumerate(document.inline_shapes):
    print(f"{index:02d} width_cm={shape.width / 360000:.2f} height_cm={shape.height / 360000:.2f}")
paragraph_by_element = {paragraph._p: paragraph for paragraph in document.paragraphs}
table_by_element = {table._tbl: table for table in document.tables}
blocks = []
for child in document.element.body.iterchildren():
    if child.tag == qn("w:p"):
        blocks.append(("P", paragraph_by_element[child]))
    elif child.tag == qn("w:tbl"):
        blocks.append(("T", table_by_element[child]))

print("=== HEADINGS ===")
for index, (kind, block) in enumerate(blocks):
    if kind == "P" and block.style.name.startswith("Heading"):
        print(f"{index:04d} [{block.style.name}] {block.text}")

start = next(index for index, (kind, block) in enumerate(blocks) if kind == "P" and block.text.strip() == "二级菜单—预警提示与处置")
end = next(index for index, (kind, block) in enumerate(blocks) if index > start and kind == "P" and block.text.strip() == "一级菜单——指标管理")
print(f"\n=== TARGET SECTION {start}..{end} ===")
for index in range(start, end + 1):
    kind, block = blocks[index]
    if kind == "P":
        text = block.text.strip().replace("\n", " | ")
        drawing = bool(block._p.xpath(".//w:drawing | .//w:pict"))
        print(f"{index:04d} P [{block.style.name}] drawing={drawing} :: {text}")
    else:
        print(f"{index:04d} T rows={len(block.rows)} cols={len(block.columns)}")
        for row in block.rows:
            print("    " + " || ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
