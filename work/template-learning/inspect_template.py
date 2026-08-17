import json
import sys
from pathlib import Path

from docx import Document

src = Path(r"C:\Users\Yiheng.jin\Desktop\SHGJ\业务功能说明书\并表系统需求说明书V1.2.docx")
out = Path(r"C:\Users\Yiheng.jin\demo\work\template-learning")
out.mkdir(parents=True, exist_ok=True)

doc = Document(src)
outline = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    style = p.style.name if p.style else ""
    if style.startswith("Heading") or style.lower().startswith("toc") or any(
        k in text for k in ("场景描述", "业务流程", "功能需求", "功能说明", "传导流程")
    ):
        outline.append({"i": i, "style": style, "text": text[:500]})

tables = []
for i, table in enumerate(doc.tables):
    rows = len(table.rows)
    cols = len(table.columns)
    header = [c.text.strip()[:100] for c in table.rows[0].cells] if rows else []
    sample = []
    for row in table.rows[:4]:
        sample.append([c.text.strip()[:300] for c in row.cells])
    tables.append({"i": i, "rows": rows, "cols": cols, "header": header, "sample": sample})

report = {
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "sections": len(doc.sections),
    "outline": outline,
    "table_summary": tables,
}
(out / "template_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({k: report[k] for k in ("paragraphs", "tables", "inline_shapes", "sections")}, ensure_ascii=False))
print("OUTLINE_COUNT", len(outline))
print("TABLE_HEADERS")
for item in tables:
    print(item["i"], item["rows"], item["cols"], item["header"])
if len(sys.argv) > 1 and sys.argv[1] == "full":
    print("OUTLINE")
    for item in outline:
        print(f"{item['i']}\t{item['style']}\t{item['text']}")
    print("TABLE_SAMPLES")
    for item in tables:
        print(json.dumps(item, ensure_ascii=False))
