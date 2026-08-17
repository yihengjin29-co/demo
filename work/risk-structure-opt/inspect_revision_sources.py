import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Yiheng.jin\demo")
DOCS = {
    "original": ROOT / "work" / "risk-structure-opt" / "source.docx",
    "optimized": ROOT / "deliverables" / "金控并表管理系统说明书_风险并表功能优化版.docx",
}
XLSX = Path(r"C:\Users\Yiheng.jin\Desktop\SHGJ\附件2：上海国际集团并表管理系统建设_系统功能清单_CL0728.xlsx")


def paragraph_outline(path: Path):
    doc = Document(path)
    rows = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        style = p.style.name if p.style else ""
        if style.startswith("Heading") or "说明" in text or "流程" in text or i >= 65:
            rows.append({"i": i, "style": style, "text": text[:500]})
    return {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "outline": rows}


doc_report = {name: paragraph_outline(path) for name, path in DOCS.items()}
(ROOT / "work" / "risk-structure-opt" / "doc_outline.json").write_text(
    json.dumps(doc_report, ensure_ascii=False, indent=2), encoding="utf-8"
)

wb = load_workbook(XLSX, read_only=False, data_only=False)
summary = []
for ws in wb.worksheets:
    nonempty = 0
    samples = []
    for row in ws.iter_rows():
        vals = [cell.value for cell in row]
        if any(v not in (None, "") for v in vals):
            nonempty += 1
            if len(samples) < 12:
                samples.append({"row": row[0].row, "values": vals[:20]})
    summary.append(
        {
            "title": ws.title,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "nonempty_rows": nonempty,
            "merged": [str(r) for r in list(ws.merged_cells.ranges)[:20]],
            "samples": samples,
        }
    )
(ROOT / "work" / "risk-structure-opt" / "xlsx_summary.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
)

print(json.dumps({"docs": {k: {"paragraphs": v["paragraphs"], "tables": v["tables"]} for k, v in doc_report.items()}, "sheets": summary}, ensure_ascii=False, indent=2, default=str))
