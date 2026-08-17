from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


SOURCE = Path(__file__).with_name("source.docx")
OUTPUT = Path(__file__).with_name("optimized.docx")


SECTIONS = [
    {
        "title": "（一）门户与工作台",
        "overview": (
            "围绕集团、国资公司及各金融机构不同管理层级和业务职责，系统拟设置并表管理驾驶舱、工作台等功能界面，"
            "支持相关职能部门统一查看核心风险信息、接收待办任务和消息提醒，并按照角色权限进入相关业务事项。"
        ),
        "rows": [
            (
                "并表管理驾驶舱",
                "支持相关职能部门查看纳入机构、风险指标、预警事项、重大风险事件、专项风险事项及历史指标数据等核心信息，形成集团风险并表运行态势概览。",
            ),
            (
                "工作台",
                "支持集团、国资公司及各金融机构按照角色查看待办任务、消息提醒、业务进度及常用业务入口，开展日常业务协同。",
            ),
        ],
    },
    {
        "title": "（二）并表管理核心功能",
        "overview": (
            "围绕集团并表风险偏好、金融机构风险及资本限额和指标运行监测要求，系统拟设置风险偏好及目标、指标新增与维护、"
            "指标版本管理、最新期指标状态和集中度风险监测等功能界面，支持相关职能部门开展方案管理、指标维护、版本追溯、"
            "状态监测及集中度风险识别等业务活动。"
        ),
        "rows": [
            (
                "风险偏好及目标",
                "支持相关职能部门查阅和维护集团并表风险偏好方案及目标，管理适用机构、监测频率和阈值要求，并开展方案审阅、发布、重检及动态调整等业务活动。",
            ),
            (
                "指标新增与维护",
                "支持相关职能部门开展风险及资本指标新增、定义维护、适用机构配置、监测频率维护和阈值设置等业务活动。",
            ),
            (
                "指标版本管理",
                "支持相关职能部门查询指标历史版本、生效期间和调整记录，开展指标版本发布、生效及追溯管理。",
            ),
            (
                "最新期指标状态",
                "支持相关职能部门查看最新期风险及资本指标值、监测状态、亮灯情况和机构数据，掌握指标运行及限额执行情况。",
            ),
            (
                "集中度风险监测",
                "支持相关职能部门开展客户、集团客户、行业和区域等多维度集中度指标监测、趋势分析、汇总查询及业务明细追溯等业务活动。",
            ),
        ],
    },
    {
        "title": "（三）业务规则管理",
        "overview": (
            "针对集团风险并表业务规则配置及运行管理需求，系统拟设置风险预警规则管理功能界面，支持相关职能部门开展预警规则维护、"
            "规则配置、版本管理、审核发布及运行状态管理等业务活动。"
        ),
        "rows": [
            (
                "风险预警规则管理",
                "支持相关职能部门开展风险预警规则查询、维护、配置、审核、发布和历史版本管理，并按照生效规则开展指标预警判断。",
            ),
        ],
    },
    {
        "title": "（四）任务与业务流程管理",
        "overview": (
            "面向集团并表管理过程中的任务协同及流程管理需求，系统通过工作台及相关业务界面承载任务下发、进度跟踪、催办提醒、"
            "流程审批和运行监控等业务活动。当前DEMO未设置独立的任务与业务流程管理页面，本章节不将具体流程节点拆分为独立功能界面。"
        ),
        "rows": [
            (
                "暂无独立功能界面",
                "当前相关能力由工作台及风险偏好、预警处置、重大风险事件、定期报告等业务页面协同承载，待形成独立页面后再补充页面级说明。",
            ),
        ],
    },
    {
        "title": "（五）预警与整改闭环",
        "overview": (
            "基于集团风险监测预警及整改闭环管理要求，系统拟设置预警提示与处置、新建专项风险提示、专项风险提示与管理、"
            "专项风险查看与反馈等功能界面，支持相关职能部门开展预警生成、事项下发、原因反馈、处置方案管理、整改跟踪及解除评价等业务活动。"
        ),
        "rows": [
            (
                "预警提示与处置",
                "支持相关职能部门查询风险预警事项，开展提示函下发、原因反馈、处置方案提交、执行进展跟踪、预警解除和事项关闭等业务活动。",
            ),
            (
                "新建专项风险提示",
                "支持相关职能部门根据日常管理、外部信息或风险线索形成专项风险提示，明确涉及机构、风险事项和反馈要求并发起后续协同。",
            ),
            (
                "专项风险提示与管理",
                "支持相关职能部门查询和管理专项风险提示及相关任务，跟踪机构反馈、处置进度和事项状态，并开展催办、评价及关闭管理。",
            ),
            (
                "专项风险查看与反馈",
                "支持各金融机构查看本机构专项风险提示，填写风险原因、影响情况和应对措施，提交反馈材料并持续更新处置进展。",
            ),
        ],
    },
    {
        "title": "（六）重大风险事件管理",
        "overview": (
            "围绕集团重大风险事件识别、报送和处置跟踪要求，系统拟设置重大风险事件定义管理、重大风险事件列表等功能界面，"
            "支持相关职能部门开展事件标准维护、事件首报、核实、处置跟踪、续报、终报及归档等业务活动。"
        ),
        "rows": [
            (
                "重大风险事件定义管理",
                "支持相关职能部门查询和维护重大风险事件定义、识别标准、适用范围及生效要求，为重大风险事件识别和报送提供统一依据。",
            ),
            (
                "重大风险事件列表",
                "支持相关职能部门查询重大风险事件台账，开展事件首报、核实审阅、处置方案管理、进展跟踪、续报、终报和归档管理。",
            ),
        ],
    },
    {
        "title": "（七）报表与分析中心",
        "overview": (
            "根据集团并表管理报表分析及数据应用要求，系统拟设置历史指标数据查询、定期风险报告管理等功能界面，支持相关职能部门开展指标历史查询、"
            "趋势分析、报告编制、提交审阅、发布归档和数据追溯等业务活动。"
        ),
        "rows": [
            (
                "历史指标数据查询",
                "支持相关职能部门按照指标、机构和期次查询历史指标数据，查看指标变化趋势、亮灯记录和相关明细，并按照数据权限开展结果导出。",
            ),
            (
                "定期风险报告管理",
                "支持相关职能部门开展定期风险报告编制、上传、提交、审阅、反馈、发布和归档，并引用指标运行、预警处置和重大风险事件等相关信息。",
            ),
        ],
    },
    {
        "title": "（八）AI应用",
        "overview": (
            "结合集团并表管理数字化、智能化建设要求，系统拟逐步引入人工智能应用能力，辅助相关职能部门开展数据查询、报表编制、风险识别和分析研判。"
            "当前DEMO中的智能助手以全局组件形态提供服务，尚未设置独立AI应用页面。"
        ),
        "rows": [
            (
                "暂无独立功能界面",
                "当前智能助手作为全局辅助组件使用，不作为独立页面列示；智能问数、智能填表和舆情监测等规划能力待形成实际页面后再补充。",
            ),
        ],
    },
    {
        "title": "（九）集团基础与系统管理",
        "overview": (
            "基于集团风险并表管理系统运行支撑及基础配置要求，系统拟设置系统管理功能界面，支持相关职能部门开展基础信息维护、"
            "用户及权限配置、业务参数管理和操作审计等业务活动，实现系统权限管控和操作过程可追溯。"
        ),
        "rows": [
            (
                "系统管理",
                "支持指定系统管理人员开展基础信息、用户、角色、权限和业务参数维护，查询操作日志及审计记录，并按照职责分工实施后台管理。",
            ),
        ],
    },
]


def find_paragraph(document: Document, text: str) -> Paragraph:
    return next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == text)


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def new_paragraph_like(reference: Paragraph, text: str, page_break_before: bool = False) -> Paragraph:
    element = OxmlElement("w:p")
    if reference._p.pPr is not None:
        element.append(deepcopy(reference._p.pPr))
    paragraph = Paragraph(element, reference._parent)
    run = paragraph.add_run(text)
    if reference.runs:
        source_rpr = reference.runs[0]._r.rPr
        if source_rpr is not None:
            run._r.insert(0, deepcopy(source_rpr))
    paragraph.paragraph_format.keep_with_next = True
    if page_break_before:
        paragraph.paragraph_format.page_break_before = True
    return paragraph


def copy_cell_properties(source: _Cell, target: _Cell) -> None:
    source_tcpr = source._tc.tcPr
    target_tcpr = target._tc.get_or_add_tcPr()
    for child in list(target_tcpr):
        if child.tag != qn("w:tcW"):
            target_tcpr.remove(child)
    if source_tcpr is not None:
        for child in source_tcpr:
            if child.tag != qn("w:tcW"):
                target_tcpr.append(deepcopy(child))


def copy_paragraph_run_format(source: _Cell, target: _Cell, text: str, centered: bool) -> None:
    source_paragraph = source.paragraphs[0]
    target_paragraph = target.paragraphs[0]
    if source_paragraph._p.pPr is not None:
        target_paragraph._p.insert(0, deepcopy(source_paragraph._p.pPr))
    target_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    target_paragraph.paragraph_format.space_before = Pt(0)
    target_paragraph.paragraph_format.space_after = Pt(0)
    run = target_paragraph.add_run(text)
    if source_paragraph.runs and source_paragraph.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(source_paragraph.runs[0]._r.rPr))
    run.font.size = Pt(11)
    target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell: _Cell, width_twips: int) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:type"), "dxa")
    tcw.set(qn("w:w"), str(width_twips))


def set_table_geometry(table: Table, total_width: int, widths: tuple[int, int]) -> None:
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(total_width))

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row_pr = row._tr.get_or_add_trPr()
        if row_pr.find(qn("w:cantSplit")) is None:
            row_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])

    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header = header_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        header_pr.append(header)
    header.set(qn("w:val"), "true")


def build_table(document: Document, template: Table, rows: list[tuple[str, str]]) -> Table:
    table = document.add_table(rows=len(rows) + 1, cols=2)
    if template.style is not None:
        table.style = template.style
    table.autofit = False

    header_refs = (template.rows[0].cells[1], template.rows[0].cells[2])
    data_refs = (template.rows[1].cells[1], template.rows[1].cells[2])

    values = [("功能界面", "功能说明"), *rows]
    for row_index, row_values in enumerate(values):
        refs = header_refs if row_index == 0 else data_refs
        for col_index, value in enumerate(row_values):
            cell = table.rows[row_index].cells[col_index]
            copy_cell_properties(refs[col_index], cell)
            cell.text = ""
            copy_paragraph_run_format(refs[col_index], cell, value, centered=(col_index == 0))

    usable_width = int(document.sections[-1].page_width.twips - document.sections[-1].left_margin.twips - document.sections[-1].right_margin.twips)
    first_width = int(usable_width * 0.30)
    widths = (first_width, usable_width - first_width)
    set_table_geometry(table, usable_width, widths)
    return table


def main() -> None:
    document = Document(SOURCE)
    target_heading = find_paragraph(document, "拟设功能介绍")
    end_heading = find_paragraph(document, "文档核验要求")
    heading2_reference = find_paragraph(document, "集团并表风险偏好及目标拟设功能表")
    body_reference = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith("本章按照四个一期核心模块")
    )
    table_template = document.tables[7]

    current = target_heading._p.getnext()
    while current is not None and current is not end_heading._p:
        next_element = current.getnext()
        current.getparent().remove(current)
        current = next_element

    intro = new_paragraph_like(
        body_reference,
        "本章按照最新系统DEMO页面结构，将风险并表功能划分为九类核心功能。各类功能采用“功能概述+功能界面说明表”的方式编写，功能界面名称与DEMO保持一致，同一页面仅列示一次，页面内部操作不拆分为独立功能界面。",
    )
    end_heading._p.addprevious(intro._p)

    for index, section in enumerate(SECTIONS):
        heading = new_paragraph_like(
            heading2_reference,
            section["title"],
            page_break_before=index > 0,
        )
        overview = new_paragraph_like(body_reference, section["overview"])
        overview.paragraph_format.keep_with_next = True
        table = build_table(document, table_template, section["rows"])
        end_heading._p.addprevious(heading._p)
        end_heading._p.addprevious(overview._p)
        end_heading._p.addprevious(table._tbl)

    replacements = {
        "2. 正文仅围绕集团并表风险偏好及目标、金融机构风险限额、风险并表指标监测预警和风险报告四个核心模块编写。":
            "2. 风险并表功能章节按照门户与工作台、并表管理核心功能、业务规则管理、任务与业务流程管理、预警与整改闭环、重大风险事件管理、报表与分析中心、AI应用、集团基础与系统管理九类核心功能编排。",
        "8. 第四章功能点与第五章功能名称保持一一对应。":
            "8. 第五章功能界面名称与最新系统DEMO保持一致，同一页面仅列示一次。",
        "9. 第五章保留“数据来源”列并暂时留空。":
            "9. 第五章功能界面说明表仅保留“功能界面”和“功能说明”两列。",
    }
    for paragraph in document.paragraphs:
        key = paragraph.text.strip()
        if key in replacements:
            replace_paragraph_text(paragraph, replacements[key])

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    main()
