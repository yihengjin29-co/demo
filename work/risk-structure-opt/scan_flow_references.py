import sys
from pathlib import Path
from docx import Document

root = Path(r"C:\Users\Yiheng.jin\demo")
if len(sys.argv) > 1:
    path = Path(sys.argv[1])
    doc = Document(path)
    starts = [i for i, p in enumerate(doc.paragraphs) if any(key in p.text for key in ("传导流程", "场景描述", "说明·流程", "说明・流程"))]
    for hit in starts:
        print(f"\nHIT {hit}")
        for i in range(max(0, hit - 16), min(len(doc.paragraphs), hit + 22)):
            p = doc.paragraphs[i]
            drawings = len(p._p.xpath(".//w:drawing")) + len(p._p.xpath(".//w:pict"))
            print(f"{i}\t{p.style.name if p.style else ''}\tdraw={drawings}\t{p.text.strip()[:500]}")
    raise SystemExit

for path in sorted(root.rglob("*.docx")):
    try:
        doc = Document(path)
    except Exception:
        continue
    hits = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if any(key in text for key in ("场景描述", "说明·流程", "说明・流程", "传导流程", "业务流程说明")):
            hits.append((i, p.style.name if p.style else "", text))
    if hits:
        print(f"FILE\t{path}")
        for i, style, text in hits[:40]:
            print(f"{i}\t{style}\t{text[:300]}")
