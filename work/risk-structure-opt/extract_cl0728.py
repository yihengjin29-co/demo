import json
from pathlib import Path
from openpyxl import load_workbook

xlsx = Path(r"C:\Users\Yiheng.jin\Desktop\SHGJ\附件2：上海国际集团并表管理系统建设_系统功能清单_CL0728.xlsx")
out = Path(r"C:\Users\Yiheng.jin\demo\work\risk-structure-opt\cl0728_business.json")
wb = load_workbook(xlsx, read_only=False, data_only=True)
ws = wb["业务功能需求"]

last = [None] * 8
rows = []
for row_idx in range(3, ws.max_row + 1):
    vals = [ws.cell(row_idx, col).value for col in range(1, 9)]
    for i in range(8):
        if vals[i] not in (None, ""):
            last[i] = vals[i]
    rows.append(
        {
            "row": row_idx,
            "序号": vals[0],
            "功能层级": vals[1] or last[1],
            "一级模块": vals[2] or last[2],
            "二级模块": vals[3],
            "功能点": vals[4],
            "系统功能简介": vals[5],
            "复用外购": vals[6],
            "实现阶段": vals[7],
        }
    )

out.write_text(json.dumps(rows, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

groups = {}
for item in rows:
    groups.setdefault(item["一级模块"], []).append(item)
print(json.dumps({k: len(v) for k, v in groups.items()}, ensure_ascii=False, indent=2))
for key in ["门户与工作台", "风险并表管理", "业务规则管理", "任务与业务流程管理", "预警与整改闭环", "重大风险事件管理", "报表与分析中心", "AI应用", "集团基础与系统管理"]:
    print("\n##", key)
    for item in groups.get(key, []):
        print(json.dumps(item, ensure_ascii=False, default=str))
