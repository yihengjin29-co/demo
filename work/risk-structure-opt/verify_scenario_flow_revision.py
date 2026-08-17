import hashlib
import json
from pathlib import Path

from docx import Document

root = Path(r"C:\Users\Yiheng.jin\demo")
base = root / "deliverables" / "金控并表管理系统说明书_风险并表功能优化版.docx"
out = root / "work" / "risk-structure-opt" / "scenario-flow-revision.docx"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


d = Document(out)
texts = [p.text.strip() for p in d.paragraphs]
titles = [
    "（一）门户与工作台",
    "（二）并表管理核心功能",
    "（三）业务规则管理",
    "（四）任务与业务流程管理",
    "（五）预警与整改闭环",
    "（六）重大风险事件管理",
    "（七）报表与分析中心",
    "（八）AI应用",
    "（九）集团基础与系统管理",
]
expected_interfaces = [
    "并表管理驾驶舱", "工作台", "风险偏好及目标", "指标新增与维护", "指标版本管理", "最新期指标状态", "集中度风险监测",
    "风险预警规则管理", "任务与流程管理（能力嵌入工作台及相关业务页面）", "预警提示与处置", "新建专项风险提示",
    "专项风险提示与管理", "专项风险查看与反馈", "重大风险事件定义管理", "重大风险事件列表", "历史指标数据查询",
    "定期风险报告管理", "智能助手（全局组件，暂无独立页面）", "系统管理",
]

target_start = texts.index("拟设功能介绍")
target_end = texts.index("文档核验要求")
target_texts = texts[target_start:target_end]

function_tables = []
flow_tables = []
for idx, table in enumerate(d.tables):
    header = [c.text.strip() for c in table.rows[0].cells]
    if header[:4] == ["核心功能", "拟设子功能界面", "功能名称", "功能介绍"]:
        function_tables.append({"index": idx, "rows": len(table.rows), "cols": len(table.columns)})
    elif len(table.rows) == 1 and any("→" in c.text for c in table.rows[0].cells):
        flow_tables.append({"index": idx, "cells": len(table.rows[0].cells)})

function_names = []
interfaces = []
for table in d.tables:
    header = [c.text.strip() for c in table.rows[0].cells]
    if header[:4] != ["核心功能", "拟设子功能界面", "功能名称", "功能介绍"]:
        continue
    for row in table.rows[1:]:
        vals = [c.text.strip() for c in row.cells]
        if vals[1] and vals[1] not in interfaces:
            interfaces.append(vals[1])
        if vals[2]:
            function_names.append(vals[2])

report = {
    "base_hash": sha(base),
    "output_hash": sha(out),
    "paragraph_count": len(d.paragraphs),
    "table_count": len(d.tables),
    "title_counts": {t: texts.count(t) for t in titles},
    "all_titles_once": all(texts.count(t) == 1 for t in titles),
    "label_counts_in_target": {label: target_texts.count(label) for label in ["场景描述", "传导流程", "流程：", "功能需求表"]},
    "function_tables": function_tables,
    "flow_tables": flow_tables,
    "interfaces": interfaces,
    "missing_interfaces": [x for x in expected_interfaces if x not in interfaces],
    "unexpected_interfaces": [x for x in interfaces if x not in expected_interfaces],
    "function_name_count": len(function_names),
    "blank_function_names": sum(1 for x in function_names if not x),
    "legacy_two_col_function_tables": sum(1 for t in d.tables if [c.text.strip() for c in t.rows[0].cells][:2] == ["功能界面", "功能说明"]),
    "placeholder_hits": [x for x in texts if "待补充" in x or "TODO" in x or "[[" in x],
}
print(json.dumps(report, ensure_ascii=False, indent=2))

assert report["all_titles_once"]
assert report["label_counts_in_target"] == {"场景描述": 9, "传导流程": 9, "流程：": 9, "功能需求表": 9}
assert len(function_tables) == 9
assert len(flow_tables) == 9
assert not report["missing_interfaces"]
assert not report["unexpected_interfaces"]
assert report["legacy_two_col_function_tables"] == 0
assert not report["placeholder_hits"]
