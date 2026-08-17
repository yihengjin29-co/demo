from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(document: Document):
    body = document.element.body
    paragraph_by_id = {id(p._p): p for p in document.paragraphs}
    table_by_id = {id(t._tbl): t for t in document.tables}
    for index, child in enumerate(body.iterchildren()):
        if child.tag.endswith("}p"):
            paragraph = paragraph_by_id.get(id(child))
            if paragraph is not None:
                yield index, "paragraph", paragraph
        elif child.tag.endswith("}tbl"):
            table = table_by_id.get(id(child))
            if table is not None:
                yield index, "table", table


def summarize(path: Path) -> dict:
    document = Document(path)
    blocks = []
    for body_index, kind, obj in iter_blocks(document):
        if kind == "paragraph":
            text = obj.text.strip()
            style = obj.style.name if obj.style is not None else ""
            if text and (
                "风险并表" in text
                or "门户" in text
                or "工作台" in text
                or "驾驶舱" in text
                or "风险偏好" in text
                or "指标" in text
                or "预警" in text
                or "重大风险" in text
                or "报表" in text
                or "定期报告" in text
                or "AI" in text
                or "系统管理" in text
                or "集中度" in text
                or style.lower().startswith("heading")
                or "标题" in style
            ):
                blocks.append({
                    "body_index": body_index,
                    "kind": kind,
                    "style": style,
                    "text": text[:300],
                })
        else:
            rows = []
            for row in obj.rows[:3]:
                rows.append([cell.text.strip()[:100] for cell in row.cells[:6]])
            flat = " ".join(value for row in rows for value in row)
            if any(term in flat for term in ["风险", "指标", "预警", "功能界面", "功能说明"]):
                blocks.append({
                    "body_index": body_index,
                    "kind": kind,
                    "rows": rows,
                })
    return {
        "path": str(path.resolve()),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "section_count": len(document.sections),
        "blocks": blocks,
    }


def main() -> None:
    output = [summarize(Path(arg)) for arg in sys.argv[1:]]
    print(json.dumps(output, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
