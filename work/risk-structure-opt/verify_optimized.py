from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
SOURCE_ORIGINAL = ROOT / "金控并表系统业务功能需求说明书（一期）_框架调整稿.docx"
SOURCE_COPY = Path(__file__).with_name("source.docx")
OUTPUT = Path(__file__).with_name("optimized.docx")

EXPECTED_SECTIONS = [
    "（一）门户与工作台",
    "（二）并表管理核心功能",
    "（三）业务规则管理",
    "（四）任务与业务流程管理",
    "（五）预警与整改闭环",
    "（六）重大风险事件管理",
    "（七）报表与分析中心",
    "（八）AI应用",
    "（九）集团基础与系统管理",
]

EXPECTED_PAGES = [
    "并表管理驾驶舱",
    "工作台",
    "风险偏好及目标",
    "指标新增与维护",
    "指标版本管理",
    "最新期指标状态",
    "集中度风险监测",
    "风险预警规则管理",
    "预警提示与处置",
    "新建专项风险提示",
    "专项风险提示与管理",
    "专项风险查看与反馈",
    "重大风险事件定义管理",
    "重大风险事件列表",
    "历史指标数据查询",
    "定期风险报告管理",
    "系统管理",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    document = Document(OUTPUT)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    section_counts = {title: paragraphs.count(title) for title in EXPECTED_SECTIONS}

    start_index = paragraphs.index("拟设功能介绍")
    end_index = paragraphs.index("文档核验要求")
    target_tables = document.tables[7:16]
    page_names = []
    table_shapes = []
    for table in target_tables:
        table_shapes.append([len(table.rows), len(table.columns)])
        for row in table.rows[1:]:
            name = row.cells[0].text.strip()
            if name and name != "暂无独立功能界面":
                page_names.append(name)

    result = {
        "source_hash": sha256(SOURCE_ORIGINAL),
        "source_copy_hash": sha256(SOURCE_COPY),
        "source_unchanged": sha256(SOURCE_ORIGINAL) == sha256(SOURCE_COPY),
        "output_hash": sha256(OUTPUT),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "target_section_paragraph_span": [start_index, end_index],
        "section_counts": section_counts,
        "all_sections_once": all(count == 1 for count in section_counts.values()),
        "table_shapes": table_shapes,
        "page_names": page_names,
        "missing_pages": sorted(set(EXPECTED_PAGES) - set(page_names)),
        "unexpected_pages": sorted(set(page_names) - set(EXPECTED_PAGES)),
        "duplicate_pages": sorted({name for name in page_names if page_names.count(name) > 1}),
        "old_function_titles_present": [
            title
            for title in [
                "集团并表风险偏好及目标拟设功能表",
                "金融机构风险限额拟设功能表",
                "风险并表指标监测预警拟设功能表",
                "风险报告拟设功能表",
            ]
            if title in paragraphs
        ],
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
