from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


ROOT = Path(r"C:\Users\Yiheng.jin\demo")
SOURCE = ROOT / "deliverables" / "金控并表管理系统说明书_风险并表功能优化版.docx"
WORK_OUTPUT = ROOT / "work" / "risk-structure-opt" / "scenario-flow-revision.docx"


SECTIONS = [
    {
        "title": "（一）门户与工作台",
        "core": "门户与工作台",
        "scenario": (
            "集团、国资公司及各金融机构人员进入并表管理系统后，需要在同一入口下快速掌握与本岗位相关的风险态势和业务事项。"
            "系统根据用户所属机构、管理层级、岗位职责及数据权限，差异化展示风险偏好与限额指标、预警处置、重大风险事件、"
            "风险报告、待办任务和消息提醒等信息，并将用户引导至相应业务页面。各金融机构仅查看和办理本机构范围内事项，"
            "集团及国资公司按照管理权限开展汇总监测、任务协同和进度督导，从而形成统一入口、分级展示、协同办理的日常工作场景。"
        ),
        "flow_note": (
            "用户登录后，系统先识别机构、角色和数据权限，再生成差异化驾驶舱与工作台内容；用户查看风险态势或接收待办后进入对应业务页面办理，"
            "事项状态、处理结果和提醒信息同步回传工作台。"
        ),
        "flow": ["用户登录", "识别角色与权限", "展示驾驶舱/工作台", "查看或办理事项", "反馈状态与提醒"],
        "interfaces": [
            {
                "name": "并表管理驾驶舱",
                "functions": [
                    ("查看并表综合态势", "按照用户权限查看风险并表核心指标、风险事项及整体运行情况，形成集团风险态势总览。"),
                    ("查看风险监测信息", "查看指标亮灯、预警处置、重大风险事件和定期报告等重点监测信息。"),
                    ("查看专题分析", "查看机构、风险类型及期间等维度的专题分析结果，并进入相关明细页面。"),
                ],
            },
            {
                "name": "工作台",
                "functions": [
                    ("查看统一待办", "集中查看上报、补正、审核、预警处置、整改和报告确认等待办事项。"),
                    ("处理协同任务", "接收并办理本角色任务，支持提交、退回、催办、转办及批量处理等协同活动。"),
                    ("查看消息提醒", "查看预警、流程、机构反馈和业务更新等消息，并管理已读状态和提醒。"),
                    ("查看业务进度", "查看任务、流程及近期业务事项的当前进度和处理状态。"),
                ],
            },
        ],
    },
    {
        "title": "（二）并表管理核心功能",
        "core": "并表管理核心功能",
        "scenario": (
            "在年度风险偏好、金融机构风险限额和日常风险监测管理过程中，国资公司需要结合监管要求、集团风险承受能力及各金融机构经营情况，"
            "设置风险偏好目标、指标口径、适用机构、监测频率和红黄灯阈值，并形成可生效、可追溯的指标版本。各金融机构按期报送数据后，"
            "系统基于生效口径计算并展示最新期指标状态，对客户、行业、区域等集中度情况进行汇总分析；当外部环境、监管要求或经营情况发生变化时，"
            "相关部门可发起指标或阈值重检调整，使方案制定、指标维护、版本发布、运行监测和风险识别相互衔接。"
        ),
        "flow_note": (
            "国资公司制定风险偏好和限额要求，维护指标、阈值及适用范围并发布生效版本；金融机构按期报送数据，系统计算最新期指标和集中度结果，"
            "监测结果用于触发预警处置或发起指标重检。"
        ),
        "flow": ["制定偏好与限额", "维护指标及阈值", "发布生效版本", "报送数据并计算", "监测状态与集中度", "预警或重检"],
        "interfaces": [
            {
                "name": "风险偏好及目标",
                "functions": [
                    ("查看风险偏好方案", "查看风险偏好方案、定性陈述、偏好指标、覆盖机构、监测频率及生效状态。"),
                    ("维护风险偏好方案", "维护年度风险偏好方案及相关指标配置，形成待审阅的风险偏好设置方案。"),
                    ("维护目标及阈值", "维护偏好目标、机构分解结果以及红灯、黄灯阈值等管理要求。"),
                    ("发起重检与调整", "针对监管、市场或经营变化发起重检，维护调整原因、支持材料和调整后的方案。"),
                ],
            },
            {
                "name": "指标新增与维护",
                "functions": [
                    ("查看指标目录", "查看风险及资本指标分类、指标名称、定义、状态和适用范围等目录信息。"),
                    ("维护指标定义", "维护指标名称、业务定义、统计口径、数据来源和计算规则等基础信息。"),
                    ("维护适用范围与频率", "维护指标适用机构、风险类型、监测频率和生效时间等管理属性。"),
                    ("维护指标阈值", "维护指标红灯、黄灯阈值及相关限额要求，并与预警规则保持衔接。"),
                ],
            },
            {
                "name": "指标版本管理",
                "functions": [
                    ("查看指标历史版本", "查看指标各历史版本、生效期间、调整内容及对应审批记录。"),
                    ("维护指标版本", "基于指标调整形成新版本，维护版本说明、生效时间和适用范围。"),
                    ("发布与回滚指标版本", "对通过审核的指标版本进行发布、生效或回滚，并保留版本变更记录。"),
                ],
            },
            {
                "name": "最新期指标状态",
                "functions": [
                    ("查看最新期指标", "按机构、指标类型和期次查看最新指标值及运行状态。"),
                    ("查看亮灯与越限状态", "查看指标红灯、黄灯、绿灯状态以及阈值突破情况。"),
                    ("查看累计亮灯记录", "查看同一指标、同一机构截至最新期的历史亮灯记录和持续亮灯情况。"),
                    ("查看指标明细", "查看指标适用口径、阈值版本、计算期次及相关业务明细。"),
                ],
            },
            {
                "name": "集中度风险监测",
                "functions": [
                    ("查看集中度指标", "查看客户、集团客户、行业、区域等集中度指标及其限额执行情况。"),
                    ("查看多维集中度分析", "按机构和不同业务维度查看集中度分布、变化趋势及跨机构汇总结果。"),
                    ("维护集中度限额", "维护集中度指标红灯、黄灯阈值以及适用机构和生效期间。"),
                    ("查看集中度明细", "查看集中度指标对应的客户、业务或机构明细，支持风险识别和结果追溯。"),
                    ("跟踪集中度风险", "查看集中度指标预警和处置状态，并衔接预警整改闭环。"),
                ],
            },
        ],
    },
    {
        "title": "（三）业务规则管理",
        "core": "业务规则管理",
        "scenario": (
            "当监管口径、风险管理制度、指标定义或预警标准发生变化时，国资公司需要及时调整风险预警规则，并明确规则参数、适用机构、执行顺序和生效时间。"
            "规则维护人员对新增或修改内容进行完整性校验和历史数据测试，经审核发布后由系统按生效版本执行；运行过程中持续关注规则命中、执行异常和变更影响，"
            "必要时进行停用、回滚或重新发布，确保预警判断口径统一、过程可控、结果可追溯。"
        ),
        "flow_note": (
            "相关部门根据制度或业务变化提出规则调整需求，完成规则及适用范围维护后开展校验测试，经审核发布形成生效版本；系统执行规则并记录命中和异常情况，"
            "运行结果用于后续优化或版本回滚。"
        ),
        "flow": ["提出规则需求", "维护规则与范围", "校验与测试", "审核并发布", "执行监控与追溯", "优化或回滚"],
        "interfaces": [
            {
                "name": "风险预警规则管理",
                "functions": [
                    ("查看预警规则", "查看规则目录、规则内容、适用指标、适用机构、生效期间和运行状态。"),
                    ("维护预警规则", "新增、修改、复制、启停风险预警规则，并维护规则参数和判断条件。"),
                    ("维护规则模板与范围", "维护常用规则模板、决策条件、优先级及适用机构和指标范围。"),
                    ("测试与发布规则", "开展规则完整性校验、样例或历史数据测试、审核发布和版本回滚。"),
                    ("查看规则执行情况", "查看执行批次、命中结果、运行状态和异常信息，并追溯生效规则版本。"),
                ],
            }
        ],
    },
    {
        "title": "（四）任务与业务流程管理",
        "core": "任务与业务流程管理",
        "scenario": (
            "风险偏好与限额设置、指标数据报送、预警处置、重大风险事件报送和定期报告管理均涉及集团、国资公司及各金融机构之间的任务传递和审核协同。"
            "业务发起部门需要明确任务内容、责任机构、完成时限和审核要求，各金融机构按任务提交数据、说明和附件，审核人员结合校验结果进行确认或退回补正。"
            "系统持续记录任务接收、办理、退回、催办、转办、会签和归档状态，使跨机构业务能够按照统一流程推进并保留全过程痕迹。"
        ),
        "flow_note": (
            "业务部门配置并下发周期性或临时任务，责任机构接收后填报和提交；系统校验并进入审核流程，存在问题时退回补正，审核通过后完成归档，"
            "全过程同步进度、催办和处理记录。当前相关能力嵌入工作台及各业务页面。"
        ),
        "flow": ["配置并下发任务", "责任机构接收", "填报与提交", "校验/退回补正", "审核或会签", "完成与归档"],
        "interfaces": [
            {
                "name": "任务与流程管理（能力嵌入工作台及相关业务页面）",
                "functions": [
                    ("查看待办任务", "查看本角色待接收、待填报、待审核、待补正及即将逾期事项。"),
                    ("维护任务下发", "维护任务类型、责任机构、数据范围、时间要求和审核人，并下发周期或临时任务。"),
                    ("跟踪任务进度", "查看任务接收、提交、退回、完成和逾期情况，开展催办、转派和统计。"),
                    ("维护业务流程", "按照业务类型维护提交、审核、会签、退回、升级和归档等流转规则。"),
                    ("查看流程运行记录", "查看当前节点、处理人、处理意见、时限、状态及全过程流转记录。"),
                ],
            }
        ],
    },
    {
        "title": "（五）预警与整改闭环",
        "core": "预警与整改闭环",
        "scenario": (
            "当风险指标突破红灯或黄灯阈值、指标出现异常波动，或相关部门发现专项风险线索时，系统需要及时形成风险提示并明确责任机构和反馈时限。"
            "金融机构收到提示后，对风险原因、影响程度和已采取措施进行分析，提交答复及应对处置方案，并持续更新整改进度和证明材料。国资公司对反馈内容进行审核、"
            "跟踪和评价，指标恢复正常或整改验证通过后解除并关闭事项；未达到关闭条件的继续亮灯、退回整改或重新开启，实现预警生成、分派、反馈、整改、验证和关闭的闭环。"
        ),
        "flow_note": (
            "系统依据指标阈值、异常波动或专项风险线索生成并分级风险提示，分派至责任机构；责任机构反馈原因并提交处置方案，国资公司持续跟踪整改、验证结果，"
            "符合条件后解除和关闭，必要时退回、升级或重开。"
        ),
        "flow": ["生成并分级预警", "分派责任机构", "反馈原因与影响", "提交并执行处置方案", "验证整改结果", "解除/关闭或重开"],
        "interfaces": [
            {
                "name": "预警提示与处置",
                "functions": [
                    ("查看预警事项", "按机构、指标、预警等级和状态查看风险预警事项及相关信息。"),
                    ("反馈预警原因", "维护指标突破原因、影响程度、初步判断和相关说明材料。"),
                    ("维护处置方案", "维护应对措施、责任人、完成时限和处置方案附件。"),
                    ("跟踪整改进度", "更新并查看整改计划、执行进度、阶段成果及超期情况。"),
                    ("解除与关闭预警", "根据指标恢复及整改验证结果办理预警解除、关闭、退回或重开。"),
                ],
            },
            {
                "name": "新建专项风险提示",
                "functions": [
                    ("新建专项风险提示", "根据日常管理、外部信息或风险线索创建专项风险提示事项。"),
                    ("维护提示内容与范围", "维护风险事项、涉及机构、风险等级、反馈要求和完成时限。"),
                    ("下发专项风险任务", "向责任机构下发专项风险提示及反馈任务，并生成协同记录。"),
                ],
            },
            {
                "name": "专项风险提示与管理",
                "functions": [
                    ("查看专项风险台账", "按机构、风险类型、状态和时间查看专项风险提示及处置情况。"),
                    ("维护专项风险事项", "维护专项风险内容、责任范围、反馈要求及事项状态。"),
                    ("跟踪反馈与处置", "查看机构反馈、应对措施、整改进度和证明材料。"),
                    ("催办、评价与关闭", "对未按期反馈或整改事项进行催办，并开展结果评价和关闭管理。"),
                ],
            },
            {
                "name": "专项风险查看与反馈",
                "functions": [
                    ("查看本机构专项风险", "查看分派至本机构的专项风险提示、反馈要求和办理时限。"),
                    ("反馈风险原因与影响", "维护风险成因、影响范围和风险程度等分析内容。"),
                    ("维护应对措施与材料", "维护应对处置措施、责任分工及相关证明材料。"),
                    ("更新处置进展", "持续更新措施执行、整改进度及事项完成情况。"),
                ],
            },
        ],
    },
    {
        "title": "（六）重大风险事件管理",
        "core": "重大风险事件管理",
        "scenario": (
            "各金融机构发现符合重大风险事件定义和识别标准的事项后，需要第一时间开展分析研判并提交事件首报及处置方案，说明事件基本情况、影响范围、发展趋势和已采取措施。"
            "国资公司对事件及报送材料进行核实、审核和必要的会签升级，并按照管理权限组织集团管理层或董事会审阅。事件处置期间，金融机构通过续报持续更新风险变化和措施执行情况，"
            "处置结束后提交终报；尚未完全消除的风险转入常态化跟踪，最终形成事件定义、首报、审核、处置、续报、终报和归档相衔接的管理场景。"
        ),
        "flow_note": (
            "金融机构依据生效的重大风险事件定义识别事件并提交首报，国资公司核实审核后组织审阅；责任机构执行处置方案并持续续报，事项结束后提交终报，"
            "经确认后归档或转入常态化跟踪。"
        ),
        "flow": ["事件识别与首报", "核实与审核", "管理层审阅", "执行处置方案", "续报与跟踪", "终报与归档"],
        "interfaces": [
            {
                "name": "重大风险事件定义管理",
                "functions": [
                    ("查看事件定义", "查看重大风险事件类型、定义摘要、识别标准、参考依据和生效状态。"),
                    ("维护事件定义", "新增或修改重大风险事件定义及其适用范围。"),
                    ("维护识别标准与依据", "维护事件识别条件、报送要求、参考制度和相关说明。"),
                    ("启停与版本管理", "办理事件定义启用、停用和版本更新，并保留历史引用关系。"),
                ],
            },
            {
                "name": "重大风险事件列表",
                "functions": [
                    ("查看事件台账", "按机构、风险类型、事件状态和发生时间查看重大风险事件及处置进度。"),
                    ("维护事件首报", "创建重大风险事件首报，维护事件情况、影响范围、趋势判断和相关材料。"),
                    ("维护处置方案", "维护已采取措施、后续处置安排、责任分工和完成时限。"),
                    ("维护续报与终报", "持续更新事件发展、风险变化和措施执行情况，并在处置结束后提交终报。"),
                    ("审核与跟踪事件", "开展事件核实、审核、退回、会签、升级和处置跟踪。"),
                    ("归档重大风险事件", "对处理完成的事件进行确认归档，并保留全过程材料和流转记录。"),
                ],
            },
        ],
    },
    {
        "title": "（七）报表与分析中心",
        "core": "报表与分析中心",
        "scenario": (
            "在风险并表日常监测和定期报告编制过程中，相关部门需要按机构、指标、风险类型和期间查询历史数据，分析指标变化、亮灯记录、预警处置和重大风险事件情况。"
            "各金融机构按照要求编制并报送本机构风险管理报告，国资公司结合监测结果及报送材料进行汇总分析，形成集团风险并表管理报告并逐级提交审阅。"
            "报告确认后完成发布和归档，同时保留报告版本、审阅意见及所引用指标数据，使历史查询、分析研判、报告编制和结果追溯形成统一场景。"
        ),
        "flow_note": (
            "系统汇集已确认的指标、预警、重大风险事件和机构报告数据，相关部门开展历史查询及分析并形成报告初稿；报告经机构报送、国资公司汇总审阅和集团逐级确认后发布归档，"
            "相关结果可回溯至指标和业务明细。"
        ),
        "flow": ["汇集监测与报送数据", "历史查询与分析", "编制/上传报告", "汇总与审阅", "逐级确认", "发布归档与追溯"],
        "interfaces": [
            {
                "name": "历史指标数据查询",
                "functions": [
                    ("查看历史指标", "按指标、机构和期次查看往期风险及资本指标数据。"),
                    ("查看趋势与亮灯历史", "查看指标变化趋势、历史亮灯记录和持续亮灯情况。"),
                    ("查看指标明细与来源", "查看指标口径、阈值版本、相关明细和数据来源信息。"),
                    ("导出指标查询结果", "按照数据权限导出历史指标查询和分析结果。"),
                ],
            },
            {
                "name": "定期风险报告管理",
                "functions": [
                    ("查看定期风险报告", "按报告名称、机构、类型、周期和报送日期查看定期风险报告。"),
                    ("维护报告及附件", "新增或修改报告基本信息、内容概述和报告附件。"),
                    ("提交与审阅报告", "完成机构报告提交、国资公司收集审阅、意见反馈和逐级确认。"),
                    ("发布与归档报告", "对确认后的报告进行发布、下载和归档，并保留版本及审阅记录。"),
                ],
            },
        ],
    },
    {
        "title": "（八）AI应用",
        "core": "AI应用",
        "scenario": (
            "在风险、资本和财务数据持续积累，制度文件、监管材料、公告及舆情信息快速增长的情况下，业务人员需要通过自然语言快速查询授权范围内的数据和知识，"
            "并获得数据处理、异常识别、趋势分析和报告编制方面的辅助支持。系统拟在后续阶段通过全局智能助手接收用户问题，在完成权限校验后调用已审批的指标口径、业务数据及知识材料，"
            "生成查询结果、风险线索、分析结论或报告初稿，由业务人员进行专业复核后使用。AI输出仅作为辅助依据，不替代正式计算、阈值判断、审批决策和监管报送。"
        ),
        "flow_note": (
            "用户通过智能助手提出查询或分析需求，系统校验身份和数据权限后检索相关数据及知识材料，形成问答、分析或报告草稿；业务人员复核并决定是否采用，"
            "系统保留调用和结果记录。当前 DEMO 以全局组件承载，尚未设置独立 AI 页面。"
        ),
        "flow": ["提出智能需求", "校验角色与权限", "检索数据与知识", "生成辅助结果", "业务人员复核", "采用结果并留痕"],
        "interfaces": [
            {
                "name": "智能助手（全局组件，暂无独立页面）",
                "functions": [
                    ("智能问答", "基于制度、法规、指标口径、规则和历史案例提供解释、依据和参考材料。"),
                    ("查询并表数据", "通过自然语言查询授权范围内的风险、资本和财务指标及相关明细。"),
                    ("辅助处理数据", "辅助开展数据识别、提取、清洗、分类、汇总和异常识别。"),
                    ("辅助分析风险", "辅助识别异常趋势和结构性问题，形成机构对比、原因研判和管理建议。"),
                    ("监测外部风险", "辅助识别舆情、公告、处罚和司法信息中的主体、事件及潜在风险线索。"),
                    ("辅助编制报告", "结合已授权数据和业务材料生成预警摘要、分析结论和报告初稿。"),
                ],
            }
        ],
    },
    {
        "title": "（九）集团基础与系统管理",
        "core": "集团基础与系统管理",
        "scenario": (
            "随着集团组织架构、股权关系、金融牌照、岗位人员和业务管理口径变化，系统管理员需要及时维护机构、法人、组织、股权控制关系及公共业务维度，"
            "并按照机构、部门和岗位配置功能与数据权限。业务参数或权限调整经确认后发布至相关模块，系统在用户访问和业务处理时自动执行权限与参数控制，"
            "同时记录查询、导出、修改、审批、规则变更和结果发布等关键操作，为系统稳定运行、权限管控和责任追溯提供基础支撑。"
        ),
        "flow_note": (
            "管理员根据组织、制度或业务变化维护基础信息、公共维度、用户权限和业务参数，完成校验确认后发布生效；各业务模块调用最新配置开展访问控制和业务处理，"
            "关键操作同步形成审计日志。"
        ),
        "flow": ["提出配置变更", "维护基础信息/维度", "配置权限与参数", "校验确认并发布", "业务模块调用", "审计记录与追溯"],
        "interfaces": [
            {
                "name": "系统管理",
                "functions": [
                    ("查看与维护机构信息", "查看和维护机构、法人、组织、股权控制、合并层级、金融牌照和经营范围。"),
                    ("查看与维护公共维度", "查看和维护客户、行业、区域、产品、币种、风险分类等公共维度及映射关系。"),
                    ("维护用户权限", "按机构、部门和岗位维护用户、角色、功能权限、数据权限及敏感信息权限。"),
                    ("维护业务参数", "维护业务周期、数据字典、提醒规则和模块运行参数。"),
                    ("查看操作审计", "查看查询、导出、修改、审批、调整、规则变更和结果发布等关键操作记录。"),
                ],
            }
        ],
    },
]


def find_paragraph(document: Document, text: str) -> Paragraph:
    return next(p for p in document.paragraphs if p.text.strip() == text)


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def new_paragraph_like(reference: Paragraph, text: str = "", page_break_before: bool = False) -> Paragraph:
    element = OxmlElement("w:p")
    if reference._p.pPr is not None:
        element.append(deepcopy(reference._p.pPr))
    paragraph = Paragraph(element, reference._parent)
    if text:
        run = paragraph.add_run(text)
        if reference.runs and reference.runs[0]._r.rPr is not None:
            run._r.insert(0, deepcopy(reference.runs[0]._r.rPr))
    paragraph.paragraph_format.page_break_before = page_break_before
    return paragraph


def add_label(reference: Paragraph, label: str) -> Paragraph:
    paragraph = new_paragraph_like(reference)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(label)
    if reference.runs and reference.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(reference.runs[0]._r.rPr))
    run.bold = True
    run.font.size = Pt(11)
    return paragraph


def add_prefixed_paragraph(reference: Paragraph, prefix: str, text: str) -> Paragraph:
    paragraph = new_paragraph_like(reference)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    lead = paragraph.add_run(prefix)
    body = paragraph.add_run(text)
    source_rpr = reference.runs[0]._r.rPr if reference.runs else None
    if source_rpr is not None:
        lead._r.insert(0, deepcopy(source_rpr))
        body._r.insert(0, deepcopy(source_rpr))
    lead.bold = True
    return paragraph


def set_cell_width(cell: _Cell, width: int) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:type"), "dxa")
    tcw.set(qn("w:w"), str(width))


def set_cell_margins(cell: _Cell, top: int = 80, start: int = 90, bottom: int = 80, end: int = 90) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = margins.find(qn(f"w:{tag}"))
        if item is None:
            item = OxmlElement(f"w:{tag}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def set_cell_shading(cell: _Cell, fill: str | None) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if fill is None:
        if shd is not None:
            tcpr.remove(shd)
        return
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_borders(cell: _Cell, color: str = "808080", size: str = "6", visible: bool = True) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single" if visible else "nil")
        if visible:
            node.set(qn("w:sz"), size)
            node.set(qn("w:color"), color)


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "华文细黑"
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), "华文细黑")


def set_table_geometry(table: Table, widths: list[int], repeat_header: bool = False) -> None:
    total = sum(widths)
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(total))
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        rowpr = row._tr.get_or_add_trPr()
        if rowpr.find(qn("w:cantSplit")) is None:
            rowpr.append(OxmlElement("w:cantSplit"))
        for idx, cell in enumerate(row.cells[: len(widths)]):
            set_cell_width(cell, widths[idx])
    if repeat_header:
        headerpr = table.rows[0]._tr.get_or_add_trPr()
        marker = headerpr.find(qn("w:tblHeader"))
        if marker is None:
            marker = OxmlElement("w:tblHeader")
            headerpr.append(marker)
        marker.set(qn("w:val"), "true")


def write_cell(cell: _Cell, text: str, size: float = 9.5, bold: bool = False, center: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def build_flow_table(document: Document, steps: list[str], usable_width: int) -> Table:
    columns = len(steps) * 2 - 1
    table = document.add_table(rows=1, cols=columns)
    arrow_width = 250
    step_width = (usable_width - arrow_width * (len(steps) - 1)) // len(steps)
    widths = []
    for idx in range(columns):
        widths.append(step_width if idx % 2 == 0 else arrow_width)
    widths[-1] += usable_width - sum(widths)
    set_table_geometry(table, widths)
    for idx, cell in enumerate(table.rows[0].cells):
        if idx % 2 == 0:
            write_cell(cell, steps[idx // 2], size=9, bold=True, center=True)
            set_cell_shading(cell, "EAF2F8")
            set_cell_borders(cell, color="7F8C8D", size="6", visible=True)
        else:
            write_cell(cell, "→", size=11, bold=True, center=True)
            set_cell_shading(cell, None)
            set_cell_borders(cell, visible=False)
    return table


def build_function_table(document: Document, section: dict, usable_width: int) -> Table:
    records = []
    interface_ranges = []
    start = 1
    for interface in section["interfaces"]:
        first = start
        for fn_name, fn_intro in interface["functions"]:
            records.append((section["core"], interface["name"], fn_name, fn_intro))
            start += 1
        interface_ranges.append((first, start - 1, interface["name"]))

    table = document.add_table(rows=len(records) + 1, cols=4)
    table.style = "Table Grid"
    headers = ["核心功能", "拟设子功能界面", "功能名称", "功能介绍"]
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        write_cell(cell, header, size=10, bold=True, center=True)
        set_cell_shading(cell, "D9E2F3")
        set_cell_borders(cell, color="000000", size="8", visible=True)

    for row_idx, record in enumerate(records, start=1):
        for col_idx, value in enumerate(record):
            center = col_idx < 3
            write_cell(table.rows[row_idx].cells[col_idx], value, size=9.5, bold=False, center=center)
            set_cell_borders(table.rows[row_idx].cells[col_idx], color="000000", size="6", visible=True)

    widths = [1050, 1550, 1650, usable_width - 4250]
    set_table_geometry(table, widths, repeat_header=True)

    if len(records) > 1:
        merged = table.cell(1, 0).merge(table.cell(len(records), 0))
        write_cell(merged, section["core"], size=9.5, center=True)
        set_cell_borders(merged, color="000000", size="6", visible=True)
    for first, last, interface_name in interface_ranges:
        if last > first:
            merged = table.cell(first, 1).merge(table.cell(last, 1))
            write_cell(merged, interface_name, size=9.5, center=True)
            set_cell_borders(merged, color="000000", size="6", visible=True)
    return table


def main() -> None:
    document = Document(SOURCE)
    target_heading = find_paragraph(document, "拟设功能介绍")
    end_heading = find_paragraph(document, "文档核验要求")
    heading_reference = find_paragraph(document, "（一）门户与工作台")
    body_reference = next(p for p in document.paragraphs if p.text.strip().startswith("本章按照最新系统DEMO页面结构"))

    current = target_heading._p.getnext()
    while current is not None and current is not end_heading._p:
        next_element = current.getnext()
        current.getparent().remove(current)
        current = next_element

    intro = new_paragraph_like(
        body_reference,
        "本章按照最新系统 DEMO 页面结构，将风险并表功能划分为九类核心功能。每类核心功能依次采用“场景描述—传导流程—功能需求表”的方式编写；"
        "功能需求结合《附件2：上海国际集团并表管理系统建设_系统功能清单_CL0728》进行归并，按查看、维护、处理和跟踪等较粗颗粒度列示，不展开到具体按钮、字段或页面操作步骤。",
    )
    end_heading._p.addprevious(intro._p)

    usable_width = int(
        document.sections[-1].page_width.twips
        - document.sections[-1].left_margin.twips
        - document.sections[-1].right_margin.twips
    )

    for idx, section in enumerate(SECTIONS):
        heading = new_paragraph_like(heading_reference, section["title"], page_break_before=idx > 0)
        heading.paragraph_format.keep_with_next = True
        end_heading._p.addprevious(heading._p)

        label = add_label(body_reference, "场景描述")
        end_heading._p.addprevious(label._p)
        scenario = new_paragraph_like(body_reference, section["scenario"])
        scenario.paragraph_format.space_after = Pt(6)
        end_heading._p.addprevious(scenario._p)

        label = add_label(body_reference, "传导流程")
        end_heading._p.addprevious(label._p)
        note = add_prefixed_paragraph(body_reference, "说明：", section["flow_note"])
        end_heading._p.addprevious(note._p)
        flow_label = add_label(body_reference, "流程：")
        flow_label.paragraph_format.space_before = Pt(0)
        flow_label.paragraph_format.space_after = Pt(3)
        end_heading._p.addprevious(flow_label._p)
        flow_table = build_flow_table(document, section["flow"], usable_width)
        end_heading._p.addprevious(flow_table._tbl)

        label = add_label(body_reference, "功能需求表")
        label.paragraph_format.space_before = Pt(10)
        end_heading._p.addprevious(label._p)
        function_table = build_function_table(document, section, usable_width)
        end_heading._p.addprevious(function_table._tbl)

    replacements = {
        "8. 第五章功能界面名称与最新系统DEMO保持一致，同一页面仅列示一次。":
            "8. 第五章拟设子功能界面名称与最新系统 DEMO 保持一致；CL0728 功能条目已归并到对应界面，不拆分为按钮级功能。",
        "9. 第五章功能界面说明表仅保留“功能界面”和“功能说明”两列。":
            "9. 第五章各核心功能按照“场景描述—传导流程—功能需求表”编写，功能需求表包含“核心功能、拟设子功能界面、功能名称、功能介绍”四列。",
        "13. 表格正文统一使用华文细黑11号字体。":
            "13. 第五章功能需求表正文统一使用华文细黑9.5号字体，其他原有表格格式保持不变。",
    }
    for paragraph in document.paragraphs:
        key = paragraph.text.strip()
        if key in replacements:
            replace_paragraph_text(paragraph, replacements[key])

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.save(WORK_OUTPUT)
    print(WORK_OUTPUT)


if __name__ == "__main__":
    main()
