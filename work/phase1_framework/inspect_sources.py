import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from docx import Document


def cell_value(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value


def inspect_workbook(path: Path, output: Path):
    wb = load_workbook(path, read_only=False, data_only=False)
    data = {"path": str(path), "sheets": []}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            values = [cell_value(c.value) for c in row]
            if any(v != "" for v in values):
                rows.append({"row": row[0].row, "values": values})
        sheet = {
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "rows": rows,
        }
        data["sheets"].append(sheet)
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def extract_workbook(path: Path, output: Path):
    wb = load_workbook(path, read_only=False, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"### SHEET {ws.title} ({ws.max_row}x{ws.max_column})")
        for row in ws.iter_rows():
            values = [str(cell_value(c.value)).replace("\n", " ") for c in row]
            if any(v for v in values):
                lines.append(f"{row[0].row}\t" + "\t".join(values))
        lines.append("")
    output.write_text("\n".join(lines), encoding="utf-8-sig")


def inspect_docx(path: Path, output: Path):
    doc = Document(path)
    paragraphs = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            paragraphs.append({"index": i, "style": p.style.name if p.style else "", "text": text})
    tables = []
    for ti, table in enumerate(doc.tables):
        rows = []
        for ri, row in enumerate(table.rows):
            vals = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
            if any(vals):
                rows.append({"row": ri, "values": vals})
        tables.append({"index": ti, "style": table.style.name if table.style else "", "rows": rows})
    data = {
        "path": str(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "paragraphs": paragraphs,
        "tables": tables,
    }
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def extract_docx_text(path: Path, output: Path):
    doc = Document(path)
    lines = [f"# {path.name}", "", "## PARAGRAPHS"]
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            lines.append(f"P{i}\t[{p.style.name if p.style else ''}]\t{text}")
    lines.append("")
    lines.append("## TABLES")
    for ti, table in enumerate(doc.tables):
        lines.append(f"### TABLE {ti} [{table.style.name if table.style else ''}]")
        for ri, row in enumerate(table.rows):
            vals = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            lines.append(f"R{ri}\t" + "\t".join(vals))
        lines.append("")
    output.write_text("\n".join(lines), encoding="utf-8-sig")


def main():
    kind, src, dst = sys.argv[1:4]
    if kind == "xlsx":
        inspect_workbook(Path(src), Path(dst))
    elif kind == "xlsx_text":
        extract_workbook(Path(src), Path(dst))
    elif kind == "docx":
        inspect_docx(Path(src), Path(dst))
    elif kind == "docx_text":
        extract_docx_text(Path(src), Path(dst))
    else:
        raise SystemExit(f"unknown kind: {kind}")


if __name__ == "__main__":
    main()
