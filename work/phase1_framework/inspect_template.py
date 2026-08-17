import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SRC = Path(r"C:\Users\Yiheng.jin\Desktop\SHGJ\金控并表系统功能需求说明书_风险并表板块..docx")
OUT = Path(r"C:\Users\Yiheng.jin\demo\work\phase1_framework\template_metrics.json")


def style_info(style):
    font = style.font
    pf = style.paragraph_format
    rpr = style.element.find(qn("w:rPr"))
    fonts = None
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is not None:
            fonts = {k.split("}")[-1]: v for k, v in rfonts.attrib.items()}
    return {
        "name": style.name,
        "type": str(style.type),
        "base_style": style.base_style.name if style.base_style else None,
        "font_name": font.name,
        "font_size_pt": font.size.pt if font.size else None,
        "bold": font.bold,
        "italic": font.italic,
        "font_xml": fonts,
        "alignment": str(pf.alignment),
        "left_indent_cm": pf.left_indent.cm if pf.left_indent else None,
        "first_line_indent_cm": pf.first_line_indent.cm if pf.first_line_indent else None,
        "space_before_pt": pf.space_before.pt if pf.space_before else None,
        "space_after_pt": pf.space_after.pt if pf.space_after else None,
        "line_spacing": str(pf.line_spacing),
        "keep_with_next": pf.keep_with_next,
        "page_break_before": pf.page_break_before,
    }


doc = Document(SRC)
sections = []
for i, s in enumerate(doc.sections):
    sections.append({
        "index": i,
        "width_cm": s.page_width.cm,
        "height_cm": s.page_height.cm,
        "top_cm": s.top_margin.cm,
        "bottom_cm": s.bottom_margin.cm,
        "left_cm": s.left_margin.cm,
        "right_cm": s.right_margin.cm,
        "header_cm": s.header_distance.cm,
        "footer_cm": s.footer_distance.cm,
        "start_type": str(s.start_type),
        "header_text": " | ".join(p.text for p in s.header.paragraphs if p.text),
        "footer_text": " | ".join(p.text for p in s.footer.paragraphs if p.text),
    })

style_names = [
    "Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4",
    "样式2", "toc 1", "toc 2", "toc 3", "Table Grid"
]
styles = {}
for name in style_names:
    try:
        styles[name] = style_info(doc.styles[name])
    except KeyError:
        styles[name] = None

with zipfile.ZipFile(SRC) as zf:
    names = set(zf.namelist())
    settings = zf.read("word/settings.xml").decode("utf-8", errors="replace")
    numbering = zf.read("word/numbering.xml").decode("utf-8", errors="replace") if "word/numbering.xml" in names else ""
    footers = {n: zf.read(n).decode("utf-8", errors="replace") for n in sorted(names) if n.startswith("word/footer") and n.endswith(".xml")}
    headers = {n: zf.read(n).decode("utf-8", errors="replace") for n in sorted(names) if n.startswith("word/header") and n.endswith(".xml")}

result = {
    "source": str(SRC),
    "sections": sections,
    "styles": styles,
    "tables": [{"index": i, "style": t.style.name if t.style else None, "rows": len(t.rows), "cols": len(t.columns)} for i, t in enumerate(doc.tables[:12])],
    "header_parts": list(headers),
    "footer_parts": list(footers),
    "settings_has_update_fields": "updateFields" in settings,
    "numbering_size": len(numbering),
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8-sig")
