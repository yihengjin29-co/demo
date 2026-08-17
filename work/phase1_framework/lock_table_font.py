from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

src = Path(r"C:\Users\Yiheng.jin\demo\work\phase1_framework\final.docx")
dst = Path(r"C:\Users\Yiheng.jin\demo\work\phase1_framework\final_patched.docx")
doc = Document(src)

style = doc.styles["Table Grid"]
style.font.name = "华文细黑"
style.font.size = Pt(11)
style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "华文细黑")
style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "华文细黑")
style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "华文细黑")

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "华文细黑"
                    run.font.size = Pt(11)
                    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                    rfonts.set(qn("w:eastAsia"), "华文细黑")
                    rfonts.set(qn("w:ascii"), "华文细黑")
                    rfonts.set(qn("w:hAnsi"), "华文细黑")

doc.save(dst)
print(dst)
