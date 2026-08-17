from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


WORK = Path(r"C:\Users\Yiheng.jin\demo\work\phase1_framework")
SOURCE = WORK / "reference.docx"
OUTPUT = WORK / "金控并表系统业务功能需求说明书（一期）_框架调整稿.docx"
FONT = "华文细黑"

ROLES = [
    "集团董事会", "集团经理层", "集团金融机构管理部门", "集团其他相关职能部门",
    "国资公司董事会", "国资公司经理层", "国资公司金融机构管理部门", "国资公司其他相关职能部门",
    "各金融机构董事会", "各金融机构经理层", "各金融机构主要责任部门", "各金融机构其他相关职能部门",
]


def P(*checked: str, pending: tuple[str, ...] = ()) -> list[str]:
    values = []
    for role in ROLES:
        if role in checked:
            values.append("√")
        elif role in pending:
            values.append("待确认")
        else:
            values.append("")
    return values


G_VIEW = ("集团董事会", "集团经理层", "集团金融机构管理部门", "集团其他相关职能部门")
Z_VIEW = ("国资公司董事会", "国资公司经理层", "国资公司金融机构管理部门", "国资公司其他相关职能部门")
F_VIEW = ("各金融机构董事会", "各金融机构经理层", "各金融机构主要责任部门", "各金融机构其他相关职能部门")
ALL_VIEW = G_VIEW + Z_VIEW + F_VIEW
CORE_VIEW = ("集团董事会", "集团经理层", "集团金融机构管理部门", "国资公司董事会", "国资公司经理层", "国资公司金融机构管理部门", "各金融机构董事会", "各金融机构经理层", "各金融机构主要责任部门")


def flow_items() -> dict[str, list[dict]]:
    appetite = [
        ("维护风险偏好陈述", "支持国资公司金融机构管理部门依据集团风险偏好方案维护定性风险偏好陈述，形成年度方案的定性管理要求。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("维护风险偏好指标", "支持维护风险偏好指标名称、定义、计量口径、单位及管理属性，形成风险偏好指标清单。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("维护适用机构及监测频率", "支持为风险偏好指标配置适用金融机构范围和监测频率，形成差异化传导与监测要求。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("设置红灯、黄灯阈值", "支持根据监管要求、集团风险承受能力及历史运行情况设置红灯、黄灯阈值，形成指标分级判断标准。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("查询监管要求及历史数据", "支持相关管理部门查询指标监管要求和历史运行数据，为风险偏好目标与阈值设置提供参考。", P(*G_VIEW, *Z_VIEW)),
        ("依据历史数据形成参考阈值", "系统根据已确认的历史数据及计算规则形成指标参考阈值，供方案编制人员分析使用。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("调整指标参考阈值", "支持国资公司金融机构管理部门在参考阈值基础上提出调整值并记录调整依据，形成待确认阈值。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("形成风险偏好设置方案", "系统汇总风险偏好陈述、指标、适用机构、监测频率及阈值，形成年度风险偏好设置方案。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("提交审阅", "支持方案编制部门提交风险偏好设置方案进入审阅流程，并留存提交版本和相关材料。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("会签审批", "支持集团和国资公司按照管理权限对风险偏好设置方案开展会签、审阅或审批，形成处理意见和审批结果。", P("集团董事会", "集团经理层", "集团金融机构管理部门", "国资公司经理层", "国资公司金融机构管理部门", pending=("集团其他相关职能部门", "国资公司董事会", "国资公司其他相关职能部门"))),
        ("发布风险偏好方案", "支持授权管理部门在方案审阅通过后发布风险偏好方案，并向适用主体传导生效要求。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("风险偏好版本生效", "支持将审阅通过并发布的方案置为生效版本，作为风险限额设置和指标监测的依据。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("查看风险偏好方案", "支持各主体按照数据权限查看已发布或流程中的风险偏好方案；各金融机构仅查看适用于本机构的内容。", P(*ALL_VIEW)),
        ("查看风险偏好指标", "支持各主体按照权限查看指标定义、适用机构、监测频率和生效阈值等信息。", P(*ALL_VIEW)),
        ("开展风险偏好监测", "系统依据生效方案和规定频率开展风险偏好指标监测，形成指标状态及限额执行结果。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门")),
        ("形成风险偏好监测报告", "系统汇总风险偏好指标运行、阈值突破和处置情况，形成风险偏好监测报告。", P("集团金融机构管理部门", "国资公司金融机构管理部门", pending=("集团经理层", "国资公司经理层"))),
        ("提交风险指标重检申请", "支持金融机构主要责任部门对明显不适应经营实际的指标、阈值或频率提交重检申请。", P("各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("填写重检原因", "支持申请主体说明监管、市场环境或经营发展变化及其对指标监测结果的影响。", P("各金融机构主要责任部门")),
        ("上传重检支持材料", "支持申请主体上传重检说明、监管依据、历史数据及其他支持材料，并与申请关联留存。", P("各金融机构主要责任部门", "各金融机构其他相关职能部门")),
        ("审核重检申请", "支持管理部门核实重检申请并形成通过、退回或维持原方案的审核意见。", P("集团金融机构管理部门", "国资公司金融机构管理部门", pending=("集团经理层", "国资公司经理层"))),
        ("动态调整风险偏好方案", "对审核确认需要调整的事项，支持按原审议程序更新指标、阈值或监测频率并形成调整版本。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门", "集团经理层", "国资公司经理层"))),
        ("查询历史版本", "支持按照方案年度、生效期间和版本状态查询风险偏好历史版本及调整记录。", P(*ALL_VIEW)),
        ("下载风险偏好方案及相关材料", "支持相关主体按照数据权限下载风险偏好方案、审阅材料和附件；各金融机构仅限本机构适用范围。", P(*CORE_VIEW, "集团其他相关职能部门", "国资公司其他相关职能部门")),
    ]

    limits = [
        ("发起年度风险限额设置工作", "支持国资公司金融机构管理部门按年度发起纳入并表金融机构风险限额设置与更新工作，明确范围、要求和完成时限。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("维护风险限额指标", "支持维护金融机构风险限额指标名称、管理类型、计量口径和单位，形成差异化限额指标清单。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("维护指标定义及适用机构", "支持维护风险限额指标定义、适用金融机构及生效期间，形成可执行的机构限额要求。", P("国资公司金融机构管理部门")),
        ("维护监测频率", "支持根据指标属性和管理要求维护日、月、季或其他监测频率。", P("国资公司金融机构管理部门")),
        ("设置红灯、黄灯阈值", "支持为各机构限额指标设置红灯、黄灯阈值，形成分级预警判断标准。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("查询监管要求及历史数据", "支持查询各机构适用监管要求及指标历史数据，为限额设置提供依据。", P(*G_VIEW, *Z_VIEW, "各金融机构主要责任部门")),
        ("依据历史数据形成参考阈值", "系统根据历史数据及经确认的规则计算限额参考阈值，供限额方案编制使用。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("调整限额参考阈值", "支持管理部门结合机构实际和沟通结果对参考阈值进行调整并记录依据。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("与金融机构沟通确认", "支持国资公司金融机构管理部门与相关金融机构就限额指标、阈值和频率进行沟通并留存意见。", P("国资公司金融机构管理部门", "各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("确认最终限额及阈值", "支持相关管理部门依据沟通结果确认最终限额指标及红黄灯阈值，形成方案编制基础。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门", "各金融机构经理层"))),
        ("形成金融机构风险限额设置方案", "系统按机构汇总指标、定义、频率、阈值和生效期间，形成年度金融机构风险限额设置方案。", P("国资公司金融机构管理部门")),
        ("提交会签", "支持方案编制部门将风险限额设置方案提交相关部门会签，并留存会签意见。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("提交审批", "支持会签完成后的风险限额设置方案按照管理权限提交审批并留存审批结果。", P("集团经理层", "集团金融机构管理部门", "国资公司经理层", "国资公司金融机构管理部门", pending=("集团董事会", "国资公司董事会"))),
        ("发布风险限额方案", "支持授权管理部门在审批通过后发布风险限额方案，并向适用金融机构传导。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("风险限额版本生效", "支持将已发布方案设置为生效版本，作为指标监测和预警判断依据。", P("国资公司金融机构管理部门")),
        ("查看风险限额方案", "支持各主体按照权限查看风险限额设置方案；各金融机构仅查看本机构方案。", P(*ALL_VIEW)),
        ("查看风险限额指标", "支持各主体按照权限查看限额指标定义、频率、阈值和生效状态。", P(*ALL_VIEW)),
        ("开展风险限额监测", "系统依据生效限额方案和监测频率计算指标并判断限额执行状态。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门")),
        ("形成风险限额运行报告", "系统汇总各金融机构限额指标运行、预警和处置情况，形成风险限额运行报告。", P("集团金融机构管理部门", "国资公司金融机构管理部门", pending=("集团经理层", "国资公司经理层"))),
        ("提交风险限额重检申请", "支持金融机构主要责任部门对不适应实际情况的限额、阈值或频率提交重检申请。", P("各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("填写重检原因", "支持申请主体说明重检事项、变化因素、影响情况和调整建议。", P("各金融机构主要责任部门")),
        ("上传支持材料", "支持申请主体上传监管依据、历史数据、测算结果及其他重检支持材料。", P("各金融机构主要责任部门", "各金融机构其他相关职能部门")),
        ("审核重检申请", "支持集团或国资公司管理部门核实重检申请并形成审核意见。", P("集团金融机构管理部门", "国资公司金融机构管理部门", pending=("集团经理层", "国资公司经理层"))),
        ("动态调整风险限额方案", "对审核确认需要调整的事项，支持更新限额、阈值或频率并履行原方案审议程序。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门", "集团经理层", "国资公司经理层"))),
        ("查询历史版本", "支持查询金融机构风险限额方案的历史版本、生效期间和调整记录。", P(*ALL_VIEW)),
        ("下载风险限额方案及相关材料", "支持相关主体按照数据权限下载风险限额方案及其附件；各金融机构仅限本机构。", P(*CORE_VIEW, "集团其他相关职能部门", "国资公司其他相关职能部门")),
    ]

    monitoring = [
        ("接收指标报送数据", "支持管理部门接收金融机构按任务和频率报送的风险指标数据，并记录报送机构、期次和版本。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("填报指标数据", "支持金融机构主要责任部门按照指标口径和报送要求填报本机构指标数据。", P("各金融机构主要责任部门")),
        ("补充或修正指标数据", "支持金融机构对校验不通过或被退回的数据补充说明、修正并重新提交。", P("各金融机构主要责任部门")),
        ("审核确认报送数据", "支持金融机构内部责任部门及管理部门按照流程审核确认报送数据。", P("国资公司金融机构管理部门", "各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("自动计算风险指标", "系统依据生效指标口径、计算规则和报送数据自动计算风险并表指标并保留计算结果。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("按监测频率开展指标监测", "系统按照已配置监测频率执行指标计算和状态监测，形成期次化监测结果。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门")),
        ("根据阈值判断指标状态", "系统将指标值与生效红黄灯阈值比较，判断红、黄、绿灯状态并留存阈值版本。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("生成红灯或黄灯预警", "指标突破相应阈值后，系统自动生成红灯或黄灯预警事项并关联指标、机构和期次。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("下发预警提示函", "支持管理部门向相关金融机构下发预警提示函，明确预警事项、反馈要求和处理时限。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("查看预警信息", "支持各主体按照权限查看预警等级、指标状态、责任机构及处置状态；金融机构仅查看本机构事项。", P(*ALL_VIEW)),
        ("填写预警原因", "支持金融机构主要责任部门填写指标突破原因、影响情况和初步判断。", P("各金融机构主要责任部门", "各金融机构其他相关职能部门")),
        ("提交预警答复", "支持金融机构提交预警答复及处理意见，并记录提交时间和版本。", P("各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("上传说明材料", "支持金融机构上传与预警原因、影响分析和处置相关的证明材料。", P("各金融机构主要责任部门", "各金融机构其他相关职能部门")),
        ("制定应对处置方案", "支持金融机构针对预警事项制定应对处置方案，明确措施、责任和计划时限。", P("各金融机构主要责任部门", "各金融机构其他相关职能部门", pending=("各金融机构经理层",))),
        ("审核应对处置方案", "支持管理部门对金融机构提交的应对处置方案进行审核、退回或确认。", P("集团金融机构管理部门", "国资公司金融机构管理部门", pending=("集团经理层", "国资公司经理层"))),
        ("跟踪处置方案执行情况", "支持集团、国资公司和金融机构按照职责跟踪处置措施、责任人、节点和完成状态。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门")),
        ("提交执行情况报告", "支持金融机构按要求提交处置方案执行情况报告并更新未完成措施。", P("各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("查看预警处置进度", "支持相关主体查看预警答复、方案审核、措施执行及事项关闭进度。", P(*CORE_VIEW, "集团其他相关职能部门", "国资公司其他相关职能部门", "各金融机构其他相关职能部门")),
        ("自动解除亮灯", "指标恢复绿灯区间且满足解除规则时，系统自动解除亮灯并保留状态变更记录。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门")),
        ("手动解除亮灯", "支持授权管理部门在符合管理要求时确认解除亮灯并记录解除依据。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("提交指标重检申请", "支持金融机构对指标定义、阈值或频率显著不适应实际情况的事项提交重检申请。", P("各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("审核指标重检申请", "支持管理部门核实指标重检申请并形成维持原方案或启动调整的审核意见。", P("集团金融机构管理部门", "国资公司金融机构管理部门", pending=("集团经理层", "国资公司经理层"))),
        ("调整指标定义、阈值或频率", "对审核确认需要调整的事项，支持按指标设置与更新流程调整定义、阈值或频率并生成新版本。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("查询预警事项", "支持按照指标、机构、等级、状态和期次查询预警事项；金融机构仅查询本机构事项。", P(*ALL_VIEW)),
        ("查看预警详情", "支持查看预警指标、阈值、提示函、原因答复、处置方案、执行记录和解除信息。", P(*ALL_VIEW)),
        ("下载预警函件及相关材料", "支持按照权限下载预警提示函、答复、处置方案及附件。", P(*CORE_VIEW, "集团其他相关职能部门", "国资公司其他相关职能部门", "各金融机构其他相关职能部门")),
        ("导出监测预警结果", "支持按照查询条件和数据权限导出指标监测及预警处置结果。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门", pending=("集团经理层", "国资公司经理层"))),
        ("监测客户集中度", "支持按单一客户和单一集团客户归集跨机构投融资余额，计算客户集中度并判断亮灯状态。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门")),
        ("监测行业集中度", "支持按行业归集并表范围业务余额，计算行业集中度并判断亮灯状态。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门")),
        ("监测区域集中度", "支持按区域归集并表范围业务余额，计算区域集中度并判断亮灯状态。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门")),
        ("查看集中度趋势", "支持按照客户、行业或区域查看集中度及相关业务余额的历史变化趋势。", P(*CORE_VIEW, "集团其他相关职能部门", "国资公司其他相关职能部门")),
        ("下钻查询集中度明细", "支持从集中度汇总结果下钻至往来机构、业务类型及底层业务明细；金融机构仅限本机构数据。", P("集团经理层", "集团金融机构管理部门", "国资公司经理层", "国资公司金融机构管理部门", "各金融机构经理层", "各金融机构主要责任部门", pending=("集团董事会", "国资公司董事会", "各金融机构董事会"))),
    ]

    reports = [
        ("根据指标运行情况生成报表", "系统根据风险偏好、风险限额和指标监测结果生成期次化风险报表，并固化数据口径和指标版本。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("形成风险监测报告", "支持管理部门汇总指标运行、预警处置和重大风险事件情况，形成风险监测报告。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("查看风险报表", "支持各主体按照权限查看风险报表；各金融机构仅查看与本机构相关内容。", P(*ALL_VIEW)),
        ("提交风险报告审阅", "支持报告编制部门将风险报告及附件提交相关部门、经理层或董事会审阅。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("审阅风险报告", "支持集团和国资公司相关部门、经理层及董事会按照管理权限审阅风险报告并形成意见。", P("集团董事会", "集团经理层", "集团金融机构管理部门", "国资公司董事会", "国资公司经理层", "国资公司金融机构管理部门", pending=("集团其他相关职能部门", "国资公司其他相关职能部门"))),
        ("反馈确认风险报告", "支持相关机构或部门对报告内容反馈意见、补充说明并确认最终版本。", P("集团金融机构管理部门", "国资公司金融机构管理部门", "各金融机构主要责任部门", pending=("集团其他相关职能部门", "国资公司其他相关职能部门"))),
        ("发布风险报告", "支持授权管理部门在审阅通过后发布风险报告并记录发布版本和时间。", P("国资公司金融机构管理部门", pending=("集团金融机构管理部门",))),
        ("下载风险报告", "支持相关主体按照数据权限下载已发布风险报告及附件。", P(*CORE_VIEW, "集团其他相关职能部门", "国资公司其他相关职能部门")),
        ("归档风险报告", "支持对已发布风险报告、审阅意见、引用数据和附件统一归档并保留版本。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("创建重大风险事件", "支持金融机构发现重大或潜在重大风险事件后创建事件记录，形成事件管理编号。", P("各金融机构主要责任部门")),
        ("填写重大风险事件首报", "支持金融机构填写事件基本情况、影响范围、发展趋势、已采取措施及初步研判并提交首报。", P("各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("上传事件材料", "支持金融机构上传与重大风险事件及处置方案相关的说明、证明和报送材料。", P("各金融机构主要责任部门", "各金融机构其他相关职能部门")),
        ("核实重大风险事件", "支持金融机构管理部门核实事件事实、报送材料和影响情况，并形成核实意见。", P("集团金融机构管理部门", "国资公司金融机构管理部门")),
        ("组织重大风险事件汇报", "支持金融机构管理部门根据事件程度组织形成汇报材料并协调相关部门参与。", P("集团金融机构管理部门", "国资公司金融机构管理部门", pending=("集团其他相关职能部门", "国资公司其他相关职能部门"))),
        ("提交管理层或董事会审阅", "支持按管理权限将重大风险事件及处置方案提交集团或国资公司经理层、董事会审阅。", P("集团董事会", "集团经理层", "集团金融机构管理部门", "国资公司董事会", "国资公司经理层", "国资公司金融机构管理部门")),
        ("制定事件处置方案", "支持金融机构针对重大风险事件制定处置目标、措施、责任安排和计划完成时间。", P("各金融机构主要责任部门", "各金融机构其他相关职能部门", pending=("各金融机构经理层",))),
        ("审核事件处置方案", "支持集团或国资公司金融机构管理部门审核事件处置方案并记录意见。", P("集团金融机构管理部门", "国资公司金融机构管理部门", pending=("集团经理层", "国资公司经理层"))),
        ("跟踪事件处置执行情况", "支持相关主体持续跟踪事件变化、处置措施、责任人、完成进度及未解决风险。", P("集团经理层", "集团金融机构管理部门", "国资公司经理层", "国资公司金融机构管理部门", "各金融机构经理层", "各金融机构主要责任部门")),
        ("填写重大风险事件续报", "支持金融机构在事件处置期间提交续报，更新事件变化、风险影响和措施执行情况。", P("各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("填写重大风险事件终报", "支持金融机构在处置完成后提交终报，总结处置结果、剩余风险及后续安排。", P("各金融机构主要责任部门", pending=("各金融机构经理层",))),
        ("查询重大风险事件", "支持按照机构、风险类型、状态和发生时间查询重大风险事件；金融机构仅查询本机构事件。", P(*ALL_VIEW)),
        ("查看事件处置进度", "支持相关主体查看首报、核实、审阅、处置、续报、终报和跟踪归档进度。", P(*CORE_VIEW, "集团其他相关职能部门", "国资公司其他相关职能部门", "各金融机构其他相关职能部门")),
        ("下载事件材料", "支持相关主体按照权限下载事件报告、处置方案、审阅意见及附件。", P(*CORE_VIEW, "集团其他相关职能部门", "国资公司其他相关职能部门")),
        ("归档重大风险事件", "支持在终报审核确认后将事件及全过程材料归档；风险未完全消除的可转入常态化跟踪。", P("集团金融机构管理部门", "国资公司金融机构管理部门", pending=("各金融机构主要责任部门",))),
    ]

    def pack(rows):
        return [{"name": n, "desc": d, "perm": p} for n, d, p in rows]

    return {
        "集团并表风险偏好及目标": pack(appetite),
        "金融机构风险限额": pack(limits),
        "风险并表指标监测预警": pack(monitoring),
        "风险报告": pack(reports),
    }


TERMS = [
    ("集团并表风险偏好", "集团在并表管理范围内，结合战略目标、监管要求和风险承受能力，对可接受风险类型、水平及边界作出的总体表述。"),
    ("风险偏好目标", "对集团并表风险偏好进行量化或可监测化分解后形成的阶段性管理目标。"),
    ("风险偏好指标", "用于反映集团并表风险偏好目标执行情况，并配置适用机构、监测频率和阈值的指标。"),
    ("金融机构风险限额", "依据集团风险偏好和金融机构业务特征，对纳入并表金融机构相关风险指标设定的管理边界。"),
    ("风险限额指标", "用于度量金融机构风险限额执行情况并开展阈值判断的具体指标。"),
    ("风险并表指标", "按照统一口径对纳入并表金融机构相关风险数据进行计算、汇总或展示的指标。"),
    ("风险监测", "按照规定频率采集或接收数据，计算风险指标并持续识别指标状态变化的管理活动。"),
    ("风险预警", "风险指标达到或突破设定阈值后，系统形成预警事项并推动反馈、处置、跟踪和解除的管理过程。"),
    ("红灯预警", "风险指标达到红灯阈值所形成的较高等级预警，原则上应及时反馈原因、制定处置方案并持续跟踪。"),
    ("黄灯预警", "风险指标达到黄灯阈值所形成的提示性预警，用于提醒相关主体分析原因并采取必要管理措施。"),
    ("风险指标重检", "因监管要求、市场环境或经营发展变化，申请重新核实指标定义、阈值或监测频率适用性的业务活动。"),
    ("风险报告", "基于风险偏好、风险限额、指标监测预警和重大风险事件等信息形成的风险管理报告。"),
    ("重大风险事件", "对金融机构经营、财务、声誉、合规或持续稳健运行可能产生重大影响，并需要按规定报告和处置的风险事项。"),
    ("首报", "重大风险事件发生或发现后首次提交的报告，主要说明事件基本情况、初步影响、已采取措施和处置安排。"),
    ("续报", "重大风险事件处置期间提交的阶段性报告，用于更新事件变化、风险影响和措施执行情况。"),
    ("终报", "重大风险事件处置基本完成后提交的总结性报告，用于说明处置结果、剩余风险和后续安排。"),
    ("集团", "上海国际集团本部及其按照并表管理职责参与风险统筹、管理审阅和重大事项决策的相关主体。"),
    ("国资公司", "在集团并表管理架构中承担监测分析、方案管理、流程组织及协同职责的金控公司（国资公司）。"),
    ("纳入并表的金融机构", "按照集团并表管理范围纳入统一风险监测、限额管理、报告和重大风险事件管理的金融机构。"),
]


def set_run_font(run, size=11, bold=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_font(paragraph, size=11, bold=None):
    for run in paragraph.runs:
        set_run_font(run, size, bold)


def replace_para_text(paragraph, text):
    if paragraph.runs:
        for run in paragraph.runs:
            run.text = ""
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def clear_after_first_heading1(doc: Document):
    body = doc._element.body
    delete = False
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:p"):
            pstyle = child.find("./w:pPr/w:pStyle", namespaces=child.nsmap)
            if pstyle is not None and pstyle.get(qn("w:val")) == doc.styles["Heading 1"].style_id:
                delete = True
        if delete:
            body.remove(child)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def fill_cell(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    set_run_font(run, size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "000000")


def set_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).twips)))
    tc_w.set(qn("w:type"), "dxa")


def add_body(doc, text):
    p = doc.add_paragraph(text, style="样式2")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_heading(doc, text, level, new_page=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if new_page:
        p.paragraph_format.page_break_before = True
    return p


def add_simple_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        fill_cell(table.rows[0].cells[i], h, WD_ALIGN_PARAGRAPH.CENTER, True)
        if widths:
            set_cell_width(table.rows[0].cells[i], widths[i])
    set_repeat_table_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            fill_cell(row.cells[i], value, align)
            if widths:
                set_cell_width(row.cells[i], widths[i])
    set_table_borders(table)
    return table


def add_permission_matrix(doc, items):
    table = doc.add_table(rows=2, cols=14)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # Use the full landscape text width so long role headers and function names
    # wrap less aggressively; this avoids sparsely populated continuation pages.
    widths = [0.8, 5.0] + [1.70] * 12
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])

    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(1, 1))
    table.cell(0, 2).merge(table.cell(0, 5))
    table.cell(0, 6).merge(table.cell(0, 9))
    table.cell(0, 10).merge(table.cell(0, 13))
    fill_cell(table.cell(0, 0), "序号", WD_ALIGN_PARAGRAPH.CENTER, True)
    fill_cell(table.cell(0, 1), "功能点", WD_ALIGN_PARAGRAPH.CENTER, True)
    fill_cell(table.cell(0, 2), "集团", WD_ALIGN_PARAGRAPH.CENTER, True)
    fill_cell(table.cell(0, 6), "国资公司", WD_ALIGN_PARAGRAPH.CENTER, True)
    fill_cell(table.cell(0, 10), "各金融机构", WD_ALIGN_PARAGRAPH.CENTER, True)
    sub = ["董事会", "经理层", "金融机构管理部门", "其他相关职能部门"] * 2 + ["董事会", "经理层", "主要责任部门", "其他相关职能部门"]
    for i, title in enumerate(sub, start=2):
        fill_cell(table.cell(1, i), title, WD_ALIGN_PARAGRAPH.CENTER, True)
    set_repeat_table_header(table.rows[0])
    set_repeat_table_header(table.rows[1])

    for idx, item in enumerate(items, 1):
        row = table.add_row()
        fill_cell(row.cells[0], idx, WD_ALIGN_PARAGRAPH.CENTER)
        fill_cell(row.cells[1], item["name"], WD_ALIGN_PARAGRAPH.LEFT)
        for j, value in enumerate(item["perm"], start=2):
            fill_cell(row.cells[j], value, WD_ALIGN_PARAGRAPH.CENTER)
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
    set_table_borders(table)
    return table


def add_function_table(doc, items):
    rows = [[i, item["name"], item["desc"], ""] for i, item in enumerate(items, 1)]
    return add_simple_table(doc, ["序号", "功能名称", "功能描述", "数据来源"], rows, [1.0, 4.0, 8.8, 1.7])


def add_landscape_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)
    sec.header.is_linked_to_previous = True
    sec.footer.is_linked_to_previous = True
    return sec


def add_portrait_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.175)
    sec.right_margin = Cm(3.175)
    sec.header_distance = Cm(1.499)
    sec.footer_distance = Cm(1.499)
    sec.header.is_linked_to_previous = True
    sec.footer.is_linked_to_previous = True
    return sec


def update_front_matter(doc):
    replacements = {
        "系统功能需求说明书": "业务功能需求说明书（一期）",
        "版本：V2.0": "版本：V3.0",
        "2026年8月": "2026年8月",
    }
    for p in doc.paragraphs[:20]:
        if p.text.strip() in replacements:
            replace_para_text(p, replacements[p.text.strip()])

    info = doc.tables[0]
    values = {
        (1, 1): "SHGJ-RISK-BFRS-P1-V3.0", (1, 3): "内部",
        (2, 1): "业务功能需求说明书（一期）", (2, 3): "风险并表板块",
        (3, 1): "集团并表管理系统", (3, 3): "V3.0",
        (4, 1): "项目组", (4, 3): "2026-08-06",
    }
    for (r, c), value in values.items():
        fill_cell(info.cell(r, c), value, WD_ALIGN_PARAGRAPH.LEFT)
    for c in range(4):
        fill_cell(info.cell(0, c), ["信息项", "内容", "信息项", "内容"][c], WD_ALIGN_PARAGRAPH.CENTER, True)
    for r in range(1, 5):
        for c in (0, 2):
            fill_cell(info.cell(r, c), info.cell(r, c).text, WD_ALIGN_PARAGRAPH.CENTER)
    set_table_borders(info)

    rev = doc.tables[1]
    headers = ["日期", "版本", "修订说明", "修订人"]
    for c, text in enumerate(headers):
        fill_cell(rev.cell(0, c), text, WD_ALIGN_PARAGRAPH.CENTER, True)
    row = rev.rows[1]
    data = ["2026-08-06", "V3.0", "调整一期文档框架，明确四个核心模块、权限矩阵及拟设功能说明。", "项目组"]
    for c, text in enumerate(data):
        fill_cell(row.cells[c], text, WD_ALIGN_PARAGRAPH.CENTER if c != 2 else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_borders(rev)


def build():
    doc = Document(SOURCE)
    update_front_matter(doc)
    clear_after_first_heading1(doc)
    modules = flow_items()

    # Chapter 1
    add_heading(doc, "名词解释", 1, new_page=True)
    add_body(doc, "本章对一期业务功能需求中使用的主要业务术语进行统一说明。")
    add_simple_table(doc, ["名词", "解释"], TERMS, [4.0, 11.5])

    # Chapter 2
    add_heading(doc, "概述", 1, new_page=True)
    add_body(doc, "一期业务功能需求聚焦集团风险并表管理的核心闭环，围绕集团并表风险偏好及目标、金融机构风险限额、风险并表指标监测预警和风险报告四个模块建设。系统支持集团、国资公司和纳入并表的金融机构按照职责协同开展方案设置、数据报送、指标监测、预警处置、报告审阅和重大风险事件管理。")
    add_heading(doc, "集团并表风险偏好及目标", 2)
    add_body(doc, "系统围绕集团并表风险偏好及目标，支持偏好方案及目标维护、指标和阈值设置、适用机构及监测频率维护、监管要求和历史数据参考、参考阈值形成、年度设置方案形成、审阅审批、发布生效、监测报告、重检申请及动态调整等业务活动，并保留方案和指标的历史版本。")
    add_heading(doc, "金融机构风险限额", 2)
    add_body(doc, "系统围绕纳入并表金融机构的风险限额，支持限额指标及阈值维护、历史数据参考、限额设置方案形成、与金融机构沟通确认、会签审批、发布生效、日常监测、运行报告、重检申请及动态调整等业务活动，并按照机构数据权限展示相关方案和指标。")
    add_heading(doc, "风险并表指标监测预警", 2)
    add_body(doc, "系统根据生效的风险偏好及风险限额方案，接收或采集相关数据，完成指标计算、监测频率执行、阈值判断、红黄灯预警、提示函下发、原因反馈、处置方案制定、执行跟踪、预警解除及指标重检调整等闭环管理。一期集中度监测覆盖客户、行业和区域维度，并支持趋势查看和业务明细下钻。")
    add_heading(doc, "风险报告", 2)
    add_body(doc, "系统支持根据风险偏好、风险限额和指标监测结果形成风险报表及风险报告，完成报告提交、审阅、反馈、发布和归档；同时支持重大风险事件的首报、核实、组织汇报、管理层或董事会审阅、处置方案、执行跟踪、续报、终报和归档管理。")

    # Chapter 3
    add_heading(doc, "适用对象", 1, new_page=True)
    add_body(doc, "一期系统适用于集团、国资公司及纳入并表的各金融机构。当前并行期按照“集团实质管控、国资公司监测分析、金融机构执行反馈”的总体模式配置职责和数据权限；细分审批层级及职责下沉安排以最终确认的管理制度和流程为准。")
    add_heading(doc, "集团", 2)
    add_body(doc, "集团侧包括董事会、经理层、金融机构管理部门和其他相关职能部门。集团主要承担并表风险统筹监督、重大事项决策、方案和报告审阅、风险处置协调及集团范围风险信息查看等职责。")
    add_heading(doc, "国资公司", 2)
    add_body(doc, "国资公司侧包括董事会、经理层、金融机构管理部门和其他相关职能部门。国资公司主要承担风险偏好和限额方案管理、指标监测分析、预警及重大风险事件流程组织、报告编制和跨机构协同等职责。")
    add_heading(doc, "各金融机构", 2)
    add_body(doc, "各金融机构侧包括董事会、经理层、主要责任部门和其他相关职能部门。各金融机构主要承担本机构数据报送、方案沟通确认、预警原因反馈、处置方案执行、重大风险事件报告和相关材料提交等职责，并仅访问本机构权限范围内的数据。")

    # Chapter 4 intro in portrait, matrices in landscape.
    add_heading(doc, "权限管理", 1, new_page=True)
    add_heading(doc, "权限矩阵说明", 2)
    add_body(doc, "权限矩阵按照集团、国资公司和各金融机构三类适用对象设置，每类对象进一步划分为董事会、经理层、金融机构管理部门（或主要责任部门）和其他相关职能部门四个层级。单元格填写“√”表示现有材料可支持该层级具备相应职责或权限；留空表示现有材料未体现其职责；材料对具体审批层级或职责归属尚未明确的填写“待确认”。各金融机构的数据权限原则上限定为本机构范围。")
    add_landscape_section(doc)
    headings = [
        ("集团并表风险偏好及目标权限表", "集团并表风险偏好及目标"),
        ("金融机构风险限额权限表", "金融机构风险限额"),
        ("风险并表指标监测预警权限表", "风险并表指标监测预警"),
        ("风险报告权限表", "风险报告"),
    ]
    for i, (heading, key) in enumerate(headings):
        add_heading(doc, heading, 2, new_page=i > 0)
        add_permission_matrix(doc, modules[key])

    # Chapter 5 back in portrait.
    add_portrait_section(doc)
    add_heading(doc, "拟设功能介绍", 1)
    add_body(doc, "本章按照四个一期核心模块说明拟设业务功能。功能名称与第四章权限矩阵一一对应，数据来源暂由业务部门和相关机构后续补充确认。")
    function_headings = [
        ("集团并表风险偏好及目标拟设功能表", "集团并表风险偏好及目标"),
        ("金融机构风险限额拟设功能表", "金融机构风险限额"),
        ("风险并表指标监测预警拟设功能表", "风险并表指标监测预警"),
        ("风险报告拟设功能表", "风险报告"),
    ]
    for i, (heading, key) in enumerate(function_headings):
        add_heading(doc, heading, 2, new_page=i > 0)
        add_function_table(doc, modules[key])

    # Chapter 6
    add_heading(doc, "文档核验要求", 1, new_page=True)
    checks = [
        "文档仅保留第一阶段业务功能需求。",
        "正文仅围绕集团并表风险偏好及目标、金融机构风险限额、风险并表指标监测预警和风险报告四个核心模块编写。",
        "未纳入风险评价和压力测试功能。",
        "未混入资本并表、会计并表和关联交易功能。",
        "权限管理已按四个核心模块分别形成四张权限矩阵。",
        "每张权限矩阵均包括集团、国资公司、各金融机构三类对象及其四个层级。",
        "材料无法明确的细分权限已标记为“待确认”。",
        "第四章功能点与第五章功能名称保持一一对应。",
        "第五章保留“数据来源”列并暂时留空。",
        "正文未编写按钮级操作说明、页面布局和UI设计内容。",
        "功能范围与附件2一期功能及现有集团并表管理材料保持一致。",
        "文档沿用原需求说明书的页面、标题编号、表格和页眉页脚样式。",
        "表格正文统一使用华文细黑11号字体。",
        "“集团”“国资公司”“各金融机构”等主体名称在全文保持一致。",
        "最终文档目录、标题编号、分页和表格跨页显示应正常。",
    ]
    for i, item in enumerate(checks, 1):
        p = add_body(doc, f"{i}. {item}")
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        set_paragraph_font(p, 12)

    # Ensure every table run meets the explicit font requirement.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    set_paragraph_font(p, 11)

    # Ask Word to refresh fields on open as a fallback.
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
