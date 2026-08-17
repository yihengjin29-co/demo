from pathlib import Path
import sys
from docx import Document
from docx.oxml.ns import qn

path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(r"C:\Users\Yiheng.jin\demo\work\phase1_framework\final.docx")
doc = Document(path)
expected_rows = [5, 2, 20, 25, 28, 34, 26, 24, 27, 33, 25]
assert len(doc.tables) == 11, len(doc.tables)
assert [len(t.rows) for t in doc.tables] == expected_rows, [len(t.rows) for t in doc.tables]

for permission_index, function_index in ((3, 7), (4, 8), (5, 9), (6, 10)):
    permission_names = [r.cells[1].text.strip() for r in doc.tables[permission_index].rows[2:]]
    function_names = [r.cells[1].text.strip() for r in doc.tables[function_index].rows[1:]]
    assert permission_names == function_names, (permission_index, function_index)
    assert all(not r.cells[3].text.strip() for r in doc.tables[function_index].rows[1:])

font_errors = []
for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for run in p.runs:
                    if not run.text:
                        continue
                    size = run.font.size.pt if run.font.size else None
                    rfonts = run._element.get_or_add_rPr().rFonts
                    east_asia = rfonts.get(qn("w:eastAsia")) if rfonts is not None else None
                    if size != 11 or east_asia != "华文细黑":
                        font_errors.append((ti, ri, ci, run.text[:20], size, east_asia))
assert not font_errors, font_errors[:20]

assert len(doc.sections) == 3, len(doc.sections)
sizes = [(round(s.page_width.cm, 1), round(s.page_height.cm, 1)) for s in doc.sections]
assert sizes == [(21.0, 29.7), (29.7, 21.0), (21.0, 29.7)], sizes

text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
    cell.text for t in doc.tables for row in t.rows for cell in row.cells
)
for required in ("名词解释", "概述", "适用对象", "权限管理", "拟设功能介绍", "文档核验要求"):
    assert required in text

print("PASS")
print(f"tables={len(doc.tables)} rows={expected_rows} sections={sizes}")
