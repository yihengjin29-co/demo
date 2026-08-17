from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def clean(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").split())


def main() -> None:
    path = Path(sys.argv[1]).resolve()
    doc = Document(path)
    print(f"PATH\t{path}")
    print(f"PARAGRAPHS\t{len(doc.paragraphs)}\tTABLES\t{len(doc.tables)}\tSECTIONS\t{len(doc.sections)}")
    print("HEADINGS")
    for index, paragraph in enumerate(doc.paragraphs):
        text = clean(paragraph.text)
        style = paragraph.style.name or ""
        if text and ("Heading" in style or "标题" in style or "权限" in text):
            print(f"P{index}\t{style}\t{text[:180]}")
    print("TABLES")
    for table_index, table in enumerate(doc.tables):
        row_count = len(table.rows)
        col_count = max((len(row.cells) for row in table.rows), default=0)
        first = [clean(cell.text) for cell in table.rows[0].cells] if table.rows else []
        joined = " | ".join(first)
        permission_like = col_count >= 10 or "权限" in joined or any(
            mark in joined for mark in ("董事会", "经理层", "功能点", "功能名称")
        )
        print(f"T{table_index}\t{row_count}x{col_count}\t{table.style.name if table.style else ''}\t{joined[:300]}")
        if permission_like:
            for row_index, row in enumerate(table.rows):
                cells = [clean(cell.text) for cell in row.cells]
                print(f"  R{row_index}\t" + " | ".join(cells))


if __name__ == "__main__":
    main()
