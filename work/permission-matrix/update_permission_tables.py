from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document


# Matrix columns after 序号、功能点.
GB, GM, GD, GO = 2, 3, 4, 5
SB, SM, SD, SO = 6, 7, 8, 9
FB, FM, FD, FO = 10, 11, 12, 13

G_ALL = {GB, GM, GD, GO}
S_ALL = {SB, SM, SD, SO}
F_ALL = {FB, FM, FD, FO}
ALL = G_ALL | S_ALL | F_ALL
G_OPER = {GD, GO}
S_OPER = {SD, SO}
F_OPER = {FD, FO}


def rows(*items: tuple[str, set[int]]) -> list[tuple[str, set[int]]]:
    return list(items)


TABLE_ROWS: list[list[tuple[str, set[int]]]] = [
    rows(
        ("并表驾驶舱页面（查看）", G_ALL | S_ALL | {FB, FM, FD}),
        ("并表驾驶舱页面（机构下钻）", G_ALL | S_ALL | {FM, FD}),
        ("工作台页面（集团视图）", {GM, GD, GO}),
        ("工作台页面（国资公司视图）", {SM, SD, SO}),
        ("工作台页面（本机构视图）", {FM, FD, FO}),
        ("纳入机构列表页面（查看）", G_ALL | S_ALL | {FM, FD}),
        ("纳入机构详情页面（查看）", G_ALL | S_ALL | F_ALL),
        ("风险偏好及目标页面（查看）", G_ALL | S_ALL),
        ("风险偏好及目标页面（导出）", G_OPER | S_OPER),
        ("风险偏好及目标页面（新增）", {SD}),
        ("风险偏好方案编辑页面（维护）", {SD}),
        ("风险偏好方案编辑页面（暂存）", {SD}),
        ("风险偏好方案编辑页面（提交）", {SD}),
        ("风险偏好方案编辑页面（附件）", {SD}),
        ("风险偏好方案详情页面（查看）", G_ALL | S_ALL),
        ("风险偏好方案详情页面（流程）", G_ALL | S_ALL),
        ("风险偏好方案详情页面（部门审核）", {GD}),
        ("风险偏好方案详情页面（经理层审阅）", {GM}),
        ("风险偏好方案详情页面（董事会审阅）", {GB}),
        ("风险偏好方案详情页面（国资复核）", {SD}),
        ("风险偏好方案详情页面（国资审阅）", {SM}),
        ("风险偏好方案详情页面（编辑入口）", {SD}),
        ("风险偏好方案详情页面（历史版本）", G_ALL | S_ALL),
    ),
    rows(
        ("风险预警规则管理页面（查看）", ALL),
        ("风险预警规则管理页面（导出）", ALL),
        ("风险预警规则管理页面（操作记录）", ALL),
        ("预警规则维护窗口（进入）", {SD}),
        ("预警规则状态编辑窗口（维护）", {SD}),
        ("提示函下发方式窗口（维护）", {SD}),
        ("预警规则配置新增页面（维护）", {SD}),
        ("预警规则配置编辑页面（维护）", {SD}),
        ("预警规则配置查看页面（查看）", ALL),
        ("预警规则审核页面（审核）", {SM, SD}),
        ("风险预警规则管理页面（批量上传）", {SD}),
        ("预警规则配置查看页面（历史版本）", ALL),
        ("指标新增与维护页面（查看）", ALL),
        ("指标新增与维护页面（操作记录）", ALL),
        ("指标新建页面（维护）", {SD}),
        ("指标编辑页面（维护）", {SD}),
        ("指标详情页面（查看）", ALL),
        ("指标新增与维护页面（发布生效）", {SD}),
        ("指标新增与维护页面（停用作废）", {SD}),
        ("指标新增与维护页面（删除草稿）", {SD}),
        ("指标版本管理页面（查询）", ALL),
        ("指标版本管理页面（历史版本）", ALL),
        ("指标版本管理页面（下载）", ALL),
        ("最新期指标状态页面（查看）", ALL),
        ("最新期指标状态页面（亮灯历史）", ALL),
        ("最新期指标状态页面（机构数据）", ALL),
    ),
    rows(
        ("预警提示与处置页面（查询）", ALL),
        ("预警提示与处置页面（查看）", ALL),
        ("预警提示与处置页面（手动发函）", {SD}),
        ("预警处置进度总览页面（查看）", ALL),
        ("预警处置进度总览页面（流程）", ALL),
        ("预警处置节点详情页面（查看）", ALL),
        ("黄灯原因反馈页面（填写）", F_OPER),
        ("黄灯处置方案页面（填写）", F_OPER),
        ("黄灯执行进展页面（更新）", F_OPER),
        ("黄灯解除评估页面（办理）", F_OPER),
        ("红灯原因反馈页面（填写）", F_OPER),
        ("红灯处置方案页面（填写）", F_OPER),
        ("红灯集团评估页面（办理）", {GD}),
        ("红灯集团审阅页面（办理）", {GB, GM}),
        ("红灯执行进展页面（更新）", {SD} | F_OPER),
        ("红灯执行跟踪页面（办理）", {SD}),
        ("红灯解除评估页面（办理）", {SD}),
        ("预警处置节点详情页面（附件）", ALL),
        ("集中度风险监测总览页面（查看）", ALL),
        ("集中度监测总览页面（配置）", {GD, SD}),
        ("单一客户集中度汇总页面（查看）", ALL),
        ("集团客户集中度汇总页面（查看）", ALL),
        ("行业集中度汇总页面（查看）", ALL),
        ("区域集中度汇总页面（查看）", ALL),
        ("集中度分类汇总页面（导出）", ALL),
        ("集中度对象分析页面（查看）", ALL),
        ("集中度对象分析页面（趋势）", ALL),
        ("集中度对象分析页面（关联预警）", ALL),
        ("集中度业务明细页面（查看）", ALL),
        ("集中度业务明细页面（导出）", ALL),
        ("集中度页面（本机构下钻）", F_ALL),
        ("集中度页面（全机构下钻）", G_ALL | S_ALL),
    ),
    rows(
        ("重大风险事件列表页面（查看）", ALL),
        ("重大风险事件新建编辑页面（办理）", F_OPER),
        ("重大风险事件进度总览页面（查看）", ALL),
        ("重大风险事件节点页面（机构办理）", F_OPER),
        ("重大风险事件节点页面（集团核实）", {GD}),
        ("重大风险事件节点页面（方案评估）", {GD}),
        ("重大风险事件节点页面（领导审阅）", {GB, GM}),
        ("重大风险事件节点页面（跟踪关闭）", {SD}),
        ("重大风险事件定义页面（查看）", ALL),
        ("重大风险事件定义编辑页面（维护）", {SD}),
        ("报表中心页面（历史指标查看）", ALL),
        ("指标历史详情页面（查看）", ALL),
        ("报表中心页面（AI智能问数入口）", ALL),
        ("定期风险报告页面（查看下载）", ALL),
        ("定期风险报告编辑页面（报送）", {GD, SD, FD}),
        ("定期报告页面（AI生成与辅助报送）", {GD, SD, FD}),
        ("新建专项风险提示页面（草稿列表）", ALL),
        ("专项风险提示新建页面（编辑下发）", {SD}),
        ("专项风险提示管理页面（查看）", ALL),
        ("专项风险提示详情页面（工单流程）", ALL),
        ("专项风险查看与反馈页面（查看）", ALL),
        ("专项风险反馈填写页面（办理）", F_OPER),
        ("专项风险反馈评估页面（办理）", {GD}),
        ("智能助手悬浮入口（使用）", ALL),
    ),
]


EXPECTED = [23, 26, 32, 24]


def set_cell_text(cell, text: str, run_properties=None) -> None:
    paragraph = cell.paragraphs[0]
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._element.getparent().remove(extra._element)
    else:
        run = paragraph.add_run(text)
        if run_properties is not None:
            run._r.insert(0, copy.deepcopy(run_properties))
    for extra_paragraph in cell.paragraphs[1:]:
        extra_paragraph._element.getparent().remove(extra_paragraph._element)


def find_run_properties(table, permission: bool):
    cells = [cell for row in table.rows[2:] for cell in row.cells]
    if permission:
        cells = [cell for row in table.rows[2:] for cell in row.cells[2:]]
    else:
        cells = [row.cells[1] for row in table.rows[2:]]
    for cell in cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                if run._r.rPr is not None:
                    return copy.deepcopy(run._r.rPr)
    return None


def validate_source(doc: Document) -> None:
    if len(doc.tables) < 7:
        raise ValueError(f"Expected at least 7 tables, found {len(doc.tables)}")
    for offset, expected_rows in enumerate(EXPECTED):
        table = doc.tables[3 + offset]
        if len(table.rows) != expected_rows + 2:
            raise ValueError(
                f"Table {3 + offset} has {len(table.rows)} rows; expected {expected_rows + 2}"
            )
        if len(table.rows[0].cells) != 14:
            raise ValueError(f"Table {3 + offset} is not a 14-column permission matrix")
        if len(TABLE_ROWS[offset]) != expected_rows:
            raise ValueError(
                f"Replacement set {offset} has {len(TABLE_ROWS[offset])} rows; expected {expected_rows}"
            )


def update(source: Path, output: Path) -> None:
    doc = Document(source)
    validate_source(doc)
    for offset, replacement_rows in enumerate(TABLE_ROWS):
        table = doc.tables[3 + offset]
        label_rpr = find_run_properties(table, permission=False)
        mark_rpr = find_run_properties(table, permission=True)
        for sequence, (label, allowed_columns) in enumerate(replacement_rows, start=1):
            row = table.rows[sequence + 1]
            set_cell_text(row.cells[0], str(sequence), label_rpr)
            set_cell_text(row.cells[1], label, label_rpr)
            for column in range(2, 14):
                set_cell_text(row.cells[column], "√" if column in allowed_columns else "", mark_rpr)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def verify(path: Path) -> None:
    doc = Document(path)
    validate_source(doc)
    for offset, replacement_rows in enumerate(TABLE_ROWS):
        table = doc.tables[3 + offset]
        for sequence, (label, allowed_columns) in enumerate(replacement_rows, start=1):
            cells = [" ".join(cell.text.split()) for cell in table.rows[sequence + 1].cells]
            expected = [str(sequence), label] + [
                "√" if column in allowed_columns else "" for column in range(2, 14)
            ]
            if cells != expected:
                raise ValueError(f"Mismatch in table {3 + offset}, row {sequence}: {cells!r}")
            if any(value not in ("", "√") for value in cells[2:]):
                raise ValueError(f"Unexpected permission marker in table {3 + offset}, row {sequence}")
    print(f"Verified permission matrices in {path}")


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: update_permission_tables.py SOURCE.docx OUTPUT.docx")
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    update(source, output)
    verify(output)


if __name__ == "__main__":
    main()
