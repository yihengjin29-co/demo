from __future__ import annotations

import hashlib
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TARGET_TABLES = {3, 4, 5, 6}


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def clean(text: str) -> str:
    return " ".join(text.split())


def c14n(element) -> bytes:
    if element is None:
        return b""
    return etree.tostring(element, method="c14n")


def canonical_without_target_text(xml_bytes: bytes) -> bytes:
    root = etree.fromstring(xml_bytes)
    body = root.find(f"{W}body")
    table_index = 0
    for child in list(body):
        if child.tag != f"{W}tbl":
            continue
        if table_index in TARGET_TABLES:
            placeholder = etree.Element(f"{W}tbl")
            placeholder.set("target-index", str(table_index))
            body.replace(child, placeholder)
        table_index += 1
    return etree.tostring(root, method="c14n")


def main() -> None:
    source = Path(sys.argv[1]).resolve()
    candidate = Path(sys.argv[2]).resolve()
    before = Document(source)
    after = Document(candidate)

    assert len(before.paragraphs) == len(after.paragraphs)
    assert len(before.tables) == len(after.tables)
    assert len(before.sections) == len(after.sections)
    for index, (left, right) in enumerate(zip(before.paragraphs, after.paragraphs)):
        assert left.text == right.text, f"Paragraph text changed at {index}"
        assert left.style.name == right.style.name, f"Paragraph style changed at {index}"

    for table_index, (left, right) in enumerate(zip(before.tables, after.tables)):
        assert len(left.rows) == len(right.rows), f"Table {table_index} row count changed"
        assert [len(row.cells) for row in left.rows] == [len(row.cells) for row in right.rows], (
            f"Table {table_index} column structure changed"
        )
        if table_index not in TARGET_TABLES:
            left_text = [[clean(cell.text) for cell in row.cells] for row in left.rows]
            right_text = [[clean(cell.text) for cell in row.cells] for row in right.rows]
            assert left_text == right_text, f"Non-target table {table_index} text changed"
        else:
            assert [clean(cell.text) for cell in left.rows[0].cells] == [
                clean(cell.text) for cell in right.rows[0].cells
            ], f"Table {table_index} top header changed"
            assert [clean(cell.text) for cell in left.rows[1].cells] == [
                clean(cell.text) for cell in right.rows[1].cells
            ], f"Table {table_index} role header changed"
            assert c14n(left._tbl.tblPr) == c14n(right._tbl.tblPr), f"Table {table_index} properties changed"
            assert c14n(left._tbl.tblGrid) == c14n(right._tbl.tblGrid), f"Table {table_index} grid changed"
            for row_index, (left_row, right_row) in enumerate(zip(left.rows, right.rows)):
                assert c14n(left_row._tr.trPr) == c14n(right_row._tr.trPr), (
                    f"Table {table_index} row {row_index} properties changed"
                )
                for cell_index, (left_cell, right_cell) in enumerate(zip(left_row.cells, right_row.cells)):
                    assert c14n(left_cell._tc.tcPr) == c14n(right_cell._tc.tcPr), (
                        f"Table {table_index} row {row_index} cell {cell_index} properties changed"
                    )
                    assert c14n(left_cell.paragraphs[0]._p.pPr) == c14n(right_cell.paragraphs[0]._p.pPr), (
                        f"Table {table_index} row {row_index} cell {cell_index} paragraph formatting changed"
                    )

    with zipfile.ZipFile(source) as old_zip, zipfile.ZipFile(candidate) as new_zip:
        old_names = set(old_zip.namelist())
        new_names = set(new_zip.namelist())
        assert old_names == new_names, "DOCX package members changed"
        changed = []
        for name in sorted(old_names):
            if digest(old_zip.read(name)) != digest(new_zip.read(name)):
                changed.append(name)
        allowed = {"word/document.xml", "docProps/core.xml"}
        unexpected = [name for name in changed if name not in allowed]
        assert not unexpected, f"Unexpected package parts changed: {unexpected}"
        old_doc = canonical_without_target_text(old_zip.read("word/document.xml"))
        new_doc = canonical_without_target_text(new_zip.read("word/document.xml"))
        assert old_doc == new_doc, "Document XML changed outside target table text"

    target_text = "\n".join(
        clean(row.cells[1].text)
        for index in TARGET_TABLES
        for row in after.tables[index].rows[2:]
    )
    required = [
        "并表驾驶舱", "工作台", "风险偏好及目标", "风险预警规则管理", "预警提示与处置",
        "集中度风险监测", "指标新增与维护", "指标版本管理", "最新期指标状态",
        "重大风险事件列表", "重大风险事件定义", "报表中心", "定期风险报告",
        "专项风险", "AI智能问数", "智能助手",
    ]
    missing = [item for item in required if item not in target_text]
    assert not missing, f"Required DEMO pages missing: {missing}"
    forbidden = ["风险限额", "待确认"]
    present = [item for item in forbidden if item in target_text]
    assert not present, f"Old permission content remains in target matrices: {present}"
    assert not re.search(r"(?<![（(])旧版", target_text)

    print("Scope verification passed")
    print(f"Changed package parts: {changed}")
    print(f"Paragraphs: {len(after.paragraphs)}; tables: {len(after.tables)}; sections: {len(after.sections)}")


if __name__ == "__main__":
    main()
