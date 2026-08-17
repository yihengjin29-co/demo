from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document


def compact(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").split())


def main() -> None:
    path = Path(sys.argv[1]).resolve()
    doc = Document(path)
    payload: dict[str, object] = {
        "path": str(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "sections": len(doc.sections),
        "paragraph_hits": [],
        "tables": [],
    }
    keywords = ("权限", "页面功能", "建设方案", "并表驾驶舱", "工作台")
    for index, paragraph in enumerate(doc.paragraphs):
        text = compact(paragraph.text)
        if text and any(keyword in text for keyword in keywords):
            payload["paragraph_hits"].append(
                {"index": index, "style": paragraph.style.name, "text": text}
            )
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row_index, row in enumerate(table.rows):
            cells = [compact(cell.text) for cell in row.cells]
            rows.append({"row": row_index, "cells": cells})
        payload["tables"].append(
            {
                "index": table_index,
                "rows": len(table.rows),
                "columns_first_row": len(table.rows[0].cells) if table.rows else 0,
                "style": table.style.name if table.style is not None else None,
                "content": rows,
            }
        )
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
