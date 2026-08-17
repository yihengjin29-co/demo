from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn


def main() -> None:
    path = Path(sys.argv[1]).resolve()
    doc = Document(path)
    errors = []
    fixed_rows = []
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            if row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY:
                fixed_rows.append((table_index, row_index, row.height))
            for cell_index, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if not run.text:
                            continue
                        rpr = run._element.get_or_add_rPr()
                        fonts = rpr.rFonts
                        east_asia = fonts.get(qn("w:eastAsia")) if fonts is not None else None
                        size = run.font.size.pt if run.font.size else None
                        if size != 11 or east_asia != "华文细黑":
                            errors.append(
                                (table_index, row_index, cell_index, run.text[:24], size, east_asia)
                            )
    if errors:
        raise AssertionError(f"Unexpected table font formatting: {errors[:20]}")
    if fixed_rows:
        raise AssertionError(f"Fixed-height rows may clip wrapped text: {fixed_rows[:20]}")
    print("Font/layout audit passed")
    print("All table text uses 11 pt 华文细黑; no exact-height rows found")


if __name__ == "__main__":
    main()
