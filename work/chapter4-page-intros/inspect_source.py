import json
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

src = Path(__file__).with_name("source.docx")
out = Path(__file__).with_name("source_report.json")
doc = Document(src)

paragraphs = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    drawings = len(p._p.xpath(".//w:drawing")) + len(p._p.xpath(".//w:pict"))
    if text or drawings:
        paragraphs.append(
            {
                "i": i,
                "style": p.style.name if p.style else "",
                "text": text,
                "drawings": drawings,
            }
        )

tables = []
for i, table in enumerate(doc.tables):
    header = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
    tables.append({"i": i, "rows": len(table.rows), "cols": len(table.columns), "header": header})

report = {
    "paragraph_count": len(doc.paragraphs),
    "nonempty_or_drawings": len(paragraphs),
    "table_count": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "section_count": len(doc.sections),
    "paragraphs": paragraphs,
    "tables": tables,
}
out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

print(json.dumps({k: report[k] for k in ("paragraph_count", "nonempty_or_drawings", "table_count", "inline_shapes", "section_count")}, ensure_ascii=False))
print("HEADINGS")
for p in paragraphs:
    if p["style"].startswith("Heading"):
        print(f"{p['i']}\t{p['style']}\t{p['text']}")
print("MARKERS")
for p in paragraphs:
    if any(k in p["text"] for k in ("界面介绍", "界面规则", "操作说明", "页面介绍", "传导流程", "流程图", "权限清单")):
        print(f"{p['i']}\t{p['style']}\tdraw={p['drawings']}\t{p['text'][:300]}")
print("TABLES")
for t in tables:
    print(t)
