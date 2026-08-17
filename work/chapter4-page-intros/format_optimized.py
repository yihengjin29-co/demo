from __future__ import annotations

import copy
import hashlib
import importlib.util
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\Yiheng.jin\demo")
WORK = ROOT / "work" / "chapter4-page-intros"
DELIVERABLES = ROOT / "deliverables"
OUTPUT = DELIVERABLES / "金控并表系统业务功能需求说明书（一期）_页面说明精简版_格式优化版.docx"
BASE_SCRIPT = WORK / "rebuild_doc.py"
REPORT = WORK / "format_optimized_report.json"


def load_rebuilder():
    spec = importlib.util.spec_from_file_location("rebuild_doc", BASE_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def norm(text: str) -> str:
    return " ".join(text.split())


def set_run_font(run, size_pt: float) -> None:
    run.font.name = "仿宋"
    run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{attr}"), "仿宋")


def set_paragraph_font(paragraph, size_pt: float) -> None:
    for run in paragraph.runs:
        set_run_font(run, size_pt)


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as zf:
        return {
            name: hashlib.sha256(zf.read(name)).hexdigest()
            for name in zf.namelist()
            if name.startswith("word/media/")
        }


def main() -> None:
    DELIVERABLES.mkdir(parents=True, exist_ok=True)

    # 从原始说明书重新生成精简内容，但另存为新版本，避免覆盖已交付的精简版。
    rebuild = load_rebuilder()
    rebuild.OUTPUT = OUTPUT
    rebuild.QA_REPORT = WORK / "format_optimized_base_report.json"
    rebuild.main()

    doc = Document(OUTPUT)
    source_media = media_hashes(OUTPUT)
    paragraphs = list(doc.paragraphs)

    # 保存界面标题对象，后续用于双图分页。
    ui_title_texts = [
        "4.2.1.1 风险偏好及目标",
        "4.2.1.2 风险偏好及目标—新建方案",
        "4.2.1.3 风险偏好及目标—概览",
        "4.2.1.4 风险偏好及目标—风险偏好指标明细",
        "4.2.1.5 风险偏好及目标—重检申请",
        "4.2.2.1 预警规则管理",
        "4.2.2.2 预警规则管理—规则配置（提交页面）",
        "4.2.2.3 预警规则管理—规则配置（审批页面）",
        "4.2.2.4 预警规则管理—状态编辑",
        "4.2.3.1 预警提示与处置",
        "4.2.3.2 预警提示与处置—重大风险提示函",
        "4.2.3.3 预警提示与处置—重大风险提示函详情",
        "4.2.3.4 预警提示与处置—预警提示函",
        "4.2.3.5 预警提示与处置—预警提示函详情",
        "4.3.1 指标新增与维护",
        "4.3.1.1 指标新增与维护—新增指标",
        "4.3.1.2 指标新增与维护—维护指标",
        "4.3.1.3 指标新增与维护—维护记录",
        "4.3.2 指标版本管理",
        "4.3.2.1 指标版本管理—查看历史版本",
        "4.4.1 重大风险事件列表",
        "4.4.2 重大风险事件新增及报送",
        "4.5 报表中心",
        "4.5.1 报表中心—报送明细",
        "4.5.2 报表中心—报表上传",
        "4.6 定期风险报告管理",
        "4.6.1 定期风险报告管理—报告上传",
        "4.7.1 新建专项风险提示",
        "4.7.1.1 新建专项风险提示—新建提示",
        "4.7.2 专项风险提示与管理",
        "4.7.2.1 专项风险提示与管理—查看提示",
        "4.7.3 专项风险查看与反馈",
        "4.7.3.1 专项风险查看与反馈—反馈填写",
        "4.7.3.2 专项风险提示与管理—反馈评估",
    ]
    by_text = {norm(p.text): p for p in paragraphs if norm(p.text)}
    ui_titles = [by_text[text] for text in ui_title_texts]

    # 章节与界面标题统一为中文层级编号。
    heading_replacements = {
        "4.1 并表看板": "（二）并表看板",
        "4.1.1 并表管理驾驶舱": "1、并表管理驾驶舱",
        "4.2 预警管理": "（三）预警管理",
        "4.2.1 风险偏好与预警管理": "1、风险偏好与预警管理",
        "4.2.1.1 风险偏好及目标": "（1）风险偏好及目标",
        "4.2.1.2 风险偏好及目标—新建方案": "（2）风险偏好及目标—新建方案",
        "4.2.1.3 风险偏好及目标—概览": "（3）风险偏好及目标—概览",
        "4.2.1.4 风险偏好及目标—风险偏好指标明细": "（4）风险偏好及目标—风险偏好指标明细",
        "4.2.1.5 风险偏好及目标—重检申请": "（5）风险偏好及目标—重检申请",
        "4.2.2 预警规则管理": "2、预警规则管理",
        "4.2.2.1 预警规则管理": "（1）预警规则管理",
        "4.2.2.2 预警规则管理—规则配置（提交页面）": "（2）预警规则管理—规则配置（提交页面）",
        "4.2.2.3 预警规则管理—规则配置（审批页面）": "（3）预警规则管理—规则配置（审批页面）",
        "4.2.2.4 预警规则管理—状态编辑": "（4）预警规则管理—状态编辑",
        "4.2.3 预警提示与处置": "3、预警提示与处置",
        "4.2.3.1 预警提示与处置": "（1）预警提示与处置",
        "4.2.3.2 预警提示与处置—重大风险提示函": "（2）预警提示与处置—重大风险提示函",
        "4.2.3.3 预警提示与处置—重大风险提示函详情": "（3）预警提示与处置—重大风险提示函详情",
        "4.2.3.4 预警提示与处置—预警提示函": "（4）预警提示与处置—预警提示函",
        "4.2.3.5 预警提示与处置—预警提示函详情": "（5）预警提示与处置—预警提示函详情",
        "4.3 指标管理": "（四）指标管理",
        "4.3.1 指标新增与维护": "1、指标新增与维护",
        "4.3.1.1 指标新增与维护—新增指标": "（1）指标新增与维护—新增指标",
        "4.3.1.2 指标新增与维护—维护指标": "（2）指标新增与维护—维护指标",
        "4.3.1.3 指标新增与维护—维护记录": "（3）指标新增与维护—维护记录",
        "4.3.2 指标版本管理": "2、指标版本管理",
        "4.3.2.1 指标版本管理—查看历史版本": "（1）指标版本管理—查看历史版本",
        "4.4 重大风险事件管理": "（五）重大风险事件管理",
        "4.4.1 重大风险事件列表": "1、重大风险事件列表",
        "4.4.2 重大风险事件新增及报送": "2、重大风险事件新增及报送",
        "4.5 报表中心": "（六）报表中心",
        "4.5.1 报表中心—报送明细": "1、报表中心—报送明细",
        "4.5.2 报表中心—报表上传": "2、报表中心—报表上传",
        "4.6 定期风险报告管理": "（七）定期风险报告管理",
        "4.6.1 定期风险报告管理—报告上传": "1、定期风险报告管理—报告上传",
        "4.7 专项风险管理": "（八）专项风险管理",
        "4.7.1 新建专项风险提示": "1、新建专项风险提示",
        "4.7.1.1 新建专项风险提示—新建提示": "（1）新建专项风险提示—新建提示",
        "4.7.2 专项风险提示与管理": "2、专项风险提示与管理",
        "4.7.2.1 专项风险提示与管理—查看提示": "（1）专项风险提示与管理—查看提示",
        "4.7.3 专项风险查看与反馈": "3、专项风险查看与反馈",
        "4.7.3.1 专项风险查看与反馈—反馈填写": "（1）专项风险查看与反馈—反馈填写",
        "4.7.3.2 专项风险提示与管理—反馈评估": "（2）专项风险提示与管理—反馈评估",
    }
    for paragraph in paragraphs:
        old = norm(paragraph.text)
        if old in heading_replacements:
            for child in list(paragraph._p):
                if child.tag != qn("w:pPr"):
                    paragraph._p.remove(child)
            paragraph.add_run(heading_replacements[old])

    # 标题统一仿宋三号（16 磅），正文统一仿宋小三（15 磅）。
    start = next(i for i, p in enumerate(paragraphs) if norm(p.text) == "一、名词解释")
    for paragraph in paragraphs[start:]:
        if paragraph.style.name.startswith("Heading"):
            set_paragraph_font(paragraph, 16)
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = True
        elif norm(paragraph.text):
            set_paragraph_font(paragraph, 15)

    # 页面简介采用一致的正文节奏；与对应截图保持相邻，但不强制单图独占一页。
    for title in ui_titles:
        node = title._p.getnext()
        if node is not None and node.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            summary = Paragraph(node, title._parent)
            summary.paragraph_format.keep_with_next = True
            summary.paragraph_format.space_before = Pt(0)
            summary.paragraph_format.space_after = Pt(2)
            summary.paragraph_format.line_spacing = 1.15
            set_paragraph_font(summary, 15)

    # 识别 38 个内嵌图片。4 张流程图保持原尺寸，其余 34 张界面图缩放后双图排版。
    shape_paragraphs = []
    for paragraph in doc.paragraphs:
        for _ in paragraph._p.xpath(".//wp:inline"):
            shape_paragraphs.append(paragraph)
    assert len(doc.inline_shapes) == 38, len(doc.inline_shapes)
    assert len(shape_paragraphs) == 38, len(shape_paragraphs)
    flow_shape_indices = {0, 1, 2, 23}
    ui_shape_indices = [i for i in range(38) if i not in flow_shape_indices]
    assert len(ui_shape_indices) == 34

    max_width = Inches(4.45)
    max_height = Inches(2.42)
    column_width = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
    for index in ui_shape_indices:
        shape = doc.inline_shapes[index]
        old_width = shape.width
        old_height = shape.height
        scale = min(max_width / shape.width, max_height / shape.height)
        shape.width = int(shape.width * scale)
        shape.height = int(shape.height * scale)
        image_paragraph = shape_paragraphs[index]

        # 部分 DEMO 截图由一张主图和若干浮动遮罩图层组成。主图缩放时，
        # 同步缩放遮罩尺寸与相对位置，避免 Word 中出现错位白块。
        old_left = (column_width - old_width) / 2
        new_left = (column_width - shape.width) / 2
        for anchor in image_paragraph._p.xpath(".//wp:anchor"):
            extent = anchor.find(qn("wp:extent"))
            if extent is not None:
                extent.set("cx", str(int(int(extent.get("cx")) * scale)))
                extent.set("cy", str(int(int(extent.get("cy")) * scale)))
            for aext in anchor.xpath(".//a:xfrm/a:ext"):
                if aext.get("cx") is not None:
                    aext.set("cx", str(int(int(aext.get("cx")) * scale)))
                if aext.get("cy") is not None:
                    aext.set("cy", str(int(int(aext.get("cy")) * scale)))
            position_h = anchor.find(qn("wp:positionH"))
            if position_h is not None:
                offset_h = position_h.find(qn("wp:posOffset"))
                if offset_h is not None and offset_h.text:
                    old_offset = int(offset_h.text)
                    offset_h.text = str(int(new_left + scale * (old_offset - old_left)))
            position_v = anchor.find(qn("wp:positionV"))
            if position_v is not None:
                offset_v = position_v.find(qn("wp:posOffset"))
                if offset_v is not None and offset_v.text:
                    offset_v.text = str(int(int(offset_v.text) * scale))

        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_before = Pt(0)
        image_paragraph.paragraph_format.space_after = Pt(4)
        image_paragraph.paragraph_format.keep_with_next = False

    # 每两个界面截图为一组；组首标题另起一页，组内第二个界面紧随其后。
    for i, title in enumerate(ui_titles):
        title.paragraph_format.page_break_before = (i % 2 == 0)
    # 第一组连同上级标题一起换页，避免“风险偏好与预警管理”孤悬在流程图页底。
    first_parent = next(p for p in paragraphs if norm(p.text) == "1、风险偏好与预警管理")
    first_parent.paragraph_format.page_break_before = True
    ui_titles[0].paragraph_format.page_break_before = False

    doc.save(OUTPUT)

    final = Document(OUTPUT)
    report = {
        "output": str(OUTPUT),
        "paragraphs": len(final.paragraphs),
        "tables": len(final.tables),
        "inline_shapes": len(final.inline_shapes),
        "media_count": len(media_hashes(OUTPUT)),
        "media_hashes_identical": source_media == media_hashes(OUTPUT),
        "ui_screenshots": len(ui_shape_indices),
        "paired_pages_expected": len(ui_shape_indices) // 2,
        "heading_replacements": len(heading_replacements),
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
