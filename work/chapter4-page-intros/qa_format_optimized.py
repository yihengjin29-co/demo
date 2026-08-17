from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\Yiheng.jin\demo\work\chapter4-page-intros")
SOURCE = ROOT / "source.docx"
OUTPUT = Path(r"C:\Users\Yiheng.jin\demo\deliverables\金控并表系统业务功能需求说明书（一期）_页面说明精简版_格式优化版.docx")


def media_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            name.removeprefix("word/media/"): hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        }


doc = Document(OUTPUT)
started = False
in_chapter4 = False
font_errors: list[dict[str, str]] = []
arabic_chapter4: list[str] = []

for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text == "一、名词解释":
        started = True
    if text == "四、界面说明":
        in_chapter4 = True
    if not started or not text:
        continue

    is_heading = paragraph.style and paragraph.style.name.startswith("Heading")
    expected_size = 16 if is_heading else 15
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        east_asia = run._element.rPr.rFonts.get(qn("w:eastAsia")) if run._element.rPr is not None and run._element.rPr.rFonts is not None else None
        size = run.font.size.pt if run.font.size else None
        if east_asia != "仿宋" or size != expected_size:
            font_errors.append({"text": text[:50], "font": str(east_asia), "size": str(size), "style": paragraph.style.name})
    if in_chapter4 and re.match(r"^4(?:\.|、)", text):
        arabic_chapter4.append(text)

source_media = media_hashes(SOURCE)
output_media = media_hashes(OUTPUT)
report = {
    "output_exists": OUTPUT.exists(),
    "output_bytes": OUTPUT.stat().st_size,
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "media_count": len(output_media),
    "media_hashes_identical": source_media == output_media,
    "media_names_identical": set(source_media) == set(output_media),
    "media_mismatches": [name for name in sorted(set(source_media) & set(output_media)) if source_media[name] != output_media[name]],
    "font_error_count": len(font_errors),
    "font_error_samples": font_errors[:10],
    "arabic_chapter4_numbering": arabic_chapter4,
    "toc_count": len(doc._element.xpath(".//w:fldSimple[contains(@w:instr, 'TOC')]")) + len(doc._element.xpath(".//w:instrText[contains(., 'TOC')]")),
}
print(json.dumps(report, ensure_ascii=False, indent=2))
