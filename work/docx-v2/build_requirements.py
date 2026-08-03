from __future__ import annotations

import re
import os
from datetime import date
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Yiheng.jin\demo")
WORK = ROOT / "work" / "docx-v2"
REFERENCE = WORK / "reference.docx"
DELIVERABLE_DIR = ROOT / "deliverables"
OUTPUT = Path(os.environ.get("REQ_OUTPUT_DOCX", str(DELIVERABLE_DIR / "上海国际集团系统功能需求说明书_风险并表板块_V2.0.docx")))
ASSET_DIR = WORK / "placeholders"
FONT_NAME = "华文细黑"
TEXT_WIDTH_IN = 5.77
TABLE_WIDTH_DXA = 8309
BULLET_NUM_ID = None


def set_run_font(run, size: float, bold: bool | None = None, color: str = "000000"):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), FONT_NAME)


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcmar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcmar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(width_dxa))
    tcw.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def set_row_cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trpr.append(cant)


def set_table_geometry(table, widths: list[int]):
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), "0")
    tblind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def style_table_cell(cell, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, fill=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    if fill:
        set_cell_shading(cell, fill)
    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        for run in p.runs:
            set_run_font(run, 11, bold)


def add_word_table(doc, title: str | None, headers: list[str], rows: Iterable[list[str]], ratios: list[int]):
    data = list(rows)
    table = doc.add_table(rows=0, cols=len(headers))
    table.style = "Table Grid"
    if title:
        title_row = table.add_row()
        merged = title_row.cells[0].merge(title_row.cells[-1])
        merged.text = title
        style_table_cell(merged, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_repeat_table_header(title_row)
        set_row_cant_split(title_row)
    header_row = table.add_row()
    for cell, value in zip(header_row.cells, headers):
        cell.text = value
        style_table_cell(cell, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(header_row)
    set_row_cant_split(header_row)
    for values in data:
        row = table.add_row()
        if len(values) == 1:
            merged = row.cells[0].merge(row.cells[-1])
            merged.text = values[0]
            style_table_cell(merged, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, fill="F2F2F2")
        else:
            for index, (cell, value) in enumerate(zip(row.cells, values)):
                cell.text = str(value)
                narrative = index == len(values) - 1 or len(headers) == 2 and index == 1
                style_table_cell(cell, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT if narrative else WD_ALIGN_PARAGRAPH.CENTER)
        set_row_cant_split(row)
    widths = [round(TABLE_WIDTH_DXA * r / sum(ratios)) for r in ratios]
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_text(doc, text: str, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, 14, bold)
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    ppr = p._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    # Keep body bullets in the same multilevel definition as the headings.
    # Word otherwise merges the two custom definitions during field updates,
    # which can turn level-one headings into bullets.
    ilvl.set(qn("w:val"), "4")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(BULLET_NUM_ID))
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.append(numpr)
    run = p.add_run(text)
    set_run_font(run, 14, False)
    return p


def add_prompt_heading(doc, text: str):
    return add_text(doc, text, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=4)


def add_heading(doc, text: str, level: int, num_id: int):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = level == 1
    ppr = p._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is not None:
        ppr.remove(numpr)
    numpr = OxmlElement("w:numPr")
    ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level - 1))
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)
    return p


def create_heading_numbering(doc) -> int:
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abs or [0]) + 1
    num_id = max(existing_num or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    for level in range(5):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        numfmt = OxmlElement("w:numFmt")
        lvltext = OxmlElement("w:lvlText")
        if level < 4:
            numfmt.set(qn("w:val"), "decimal")
            lvltext.set(qn("w:val"), ".".join(f"%{i}" for i in range(1, level + 2)))
        else:
            numfmt.set(qn("w:val"), "bullet")
            lvltext.set(qn("w:val"), "•")
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "space")
        lvl.append(start)
        lvl.append(numfmt)
        lvl.append(lvltext)
        lvl.append(suff)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def create_bullet_numbering(doc) -> int:
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abs or [0]) + 1
    num_id = max(existing_num or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    numfmt = OxmlElement("w:numFmt")
    numfmt.set(qn("w:val"), "bullet")
    lvltext = OxmlElement("w:lvlText")
    lvltext.set(qn("w:val"), "•")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(tabs)
    ppr.append(ind)
    lvl.extend([start, numfmt, lvltext, ppr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_toc_field(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("目录")
    set_run_font(run, 16, True)
    p = doc.add_paragraph()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-4" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r1 = p.add_run()
    r1._r.append(begin)
    r2 = p.add_run()
    r2._r.append(instr)
    r3 = p.add_run()
    r3._r.append(separate)
    r4 = p.add_run("目录将在文档打开时自动更新。")
    set_run_font(r4, 14, False)
    r5 = p.add_run()
    r5._r.append(end)


def add_page_number(footer):
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.clear()
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    set_run_font(run, 10, False)


def create_placeholder(title: str) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff_-]+", "_", title)[:80]
    path = ASSET_DIR / f"{safe}.png"
    if path.exists():
        return path
    image = Image.new("RGB", (1400, 620), "#E6E6E6")
    draw = ImageDraw.Draw(image)
    draw.rectangle((12, 12, 1388, 608), outline="#9B9B9B", width=4)
    font = ImageFont.truetype(r"C:\Windows\Fonts\STXIHEI.TTF", 44)
    small = ImageFont.truetype(r"C:\Windows\Fonts\STXIHEI.TTF", 26)
    text = f"请插入当前原型截图：{title}"
    box = draw.textbbox((0, 0), text, font=font)
    draw.text(((1400 - (box[2] - box[0])) / 2, 260), text, fill="#555555", font=font)
    note = "当前环境未提供可用浏览器实例，未使用旧版截图替代"
    box2 = draw.textbbox((0, 0), note, font=small)
    draw.text(((1400 - (box2[2] - box2[0])) / 2, 345), note, fill="#777777", font=small)
    image.save(path, "PNG")
    return path


def add_placeholder(doc, title: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(str(create_placeholder(title)), width=Inches(5.65))


def clear_document_body(doc):
    body = doc._element.body
    sectpr = body.sectPr
    for child in list(body):
        if child is not sectpr:
            body.remove(child)
    image_rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype == image_rel_type:
            doc.part.drop_rel(rid)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for level in range(1, 5):
        style = styles[f"Heading {level}"]
        style.font.name = FONT_NAME
        style.font.size = Pt(14)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style_ppr = style._element.get_or_add_pPr()
        style_numpr = style_ppr.find(qn("w:numPr"))
        if style_numpr is not None:
            style_ppr.remove(style_numpr)
    bullet = styles["List Paragraph"]
    bullet.font.name = FONT_NAME
    bullet.font.size = Pt(14)
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)


def configure_document(doc):
    clear_document_body(doc)
    configure_styles(doc)
    section = doc.sections[-1]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.175)
    section.right_margin = Cm(3.175)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    section.different_first_page_header_footer = False
    section.header.paragraphs[0].clear()
    add_page_number(section.footer)
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.core_properties.title = "上海国际集团系统功能需求说明书—风险并表板块"
    doc.core_properties.subject = "V2.0系统功能需求说明书"
    doc.core_properties.author = "上海国际集团"
    doc.core_properties.comments = "根据最新DEMO页面原型、角色权限、菜单结构及业务功能全面修订。"


def field_row(area: str, name: str, kind="文本", operation="自动反显", value="页面对应内容", note="展示页面对应信息。"):
    return (area, name, kind, operation, value, note)


def query_field(name: str):
    if "起止" in name or "区间" in name:
        return field_row("查询区", name, "时间", "日期区间选择", "起始值至结束值", f"支持按照{name}进行区间查询。")
    if "时间" in name or "日期" in name:
        return field_row("查询区", name, "时间", "日期选择", "日期", f"支持按照{name}进行筛选。")
    if any(key in name for key in ("状态", "类型", "机构", "频率", "周期", "等级", "亮灯")):
        return field_row("查询区", name, "文本", "下拉选择", "页面可选项", f"支持选择{name}进行精确筛选。")
    return field_row("查询区", name, "文本", "自主输入", "100个字符以内", f"支持按照{name}关键字进行模糊查询。")


def column_field(name: str):
    if name in ("序号", "数量"):
        return field_row("列表区", name, "数字", "系统计算", "正整数", "系统根据列表展示顺序或记录数量自动计算。")
    if "时间" in name or "日期" in name or "期次" in name:
        return field_row("列表区", name, "时间", "自动反显", "页面记录值", f"展示当前记录的{name}。")
    if name in ("操作", "查看", "入口", "事件流程", "方案流程", "规则配置", "操作记录", "累计亮灯情况"):
        return field_row("列表区", name, "文本", "链接操作", "页面操作入口", f"点击后进入或打开{name}对应页面。")
    return field_row("列表区", name, "文本", "自动反显", "页面记录值", f"展示当前记录的{name}。")


def op_description(name: str, target: str = ""):
    special = {
        "查询": "系统按照已填写或选择的查询条件进行筛选，并在下方列表展示查询结果。",
        "重置": "系统清空当前查询条件并恢复页面默认展示范围。",
        "分页": "系统按照选择的每页条数、页码或上一页、下一页操作切换列表记录。",
        "返回": "系统关闭当前页面或弹层并返回上一级列表，不保存未提交内容。",
        "关闭": "系统关闭当前弹层并返回背景列表，不保存未提交内容。",
        "保存草稿": "系统执行基础校验后保存当前已填写内容，状态保持为草稿，并保留操作记录。",
        "提交": "系统校验必填信息后提交当前业务事项，并进入后续处理环节。",
        "提交审核": "系统校验必填信息后形成待审核记录，现行生效数据在审核通过前保持不变。",
        "导出": "系统导出当前权限范围和查询条件下的列表数据。",
        "下载": "系统下载所选记录或附件。",
        "查看": "系统进入对应详情页面，展示基本信息、业务数据和处理记录。",
        "查看详情": "系统进入对应详情页面，展示基本信息、业务数据和处理记录。",
        "查看流程": "系统展示当前事项的流程节点、处理记录、处理意见和附件。",
        "查看操作记录": "系统展示当前事项的历史操作人、操作时间、动作和说明。",
        "查看累计亮灯情况": "系统以表格展示同一指标、同一机构截至当前期次的历史记录。",
        "新增": "系统打开新增页面，支持录入业务信息并保存或提交。",
        "编辑": "系统打开编辑页面，支持修改当前权限范围内的数据。",
        "删除": "系统确认后删除允许删除的草稿记录，已提交记录不可删除。",
    }
    text = special.get(name, f"系统执行“{name}”对应业务操作，并反馈处理结果。")
    if target:
        text = f"系统进入{target}，{text}"
    return f"通过点击“{name}”，{text}"


def permission_rows(profile: str):
    profiles = {
        "all_view": [
            ["查询", "√", "√", "√（本机构）"], ["查看", "√", "√", "√（本机构）"],
            ["导出/下载", "√", "√", "√（本机构）"],
        ],
        "dashboard": [
            ["查看驾驶舱", "√", "√", "√（本机构）"], ["查看集团范围数据", "√", "√", ""],
            ["指标展示配置", "√", "√", "√（本机构）"], ["业务卡片跳转", "√", "√", "√（本机构）"],
        ],
        "risk_preference": [
            ["菜单及路由访问", "√", "√", ""], ["查询/查看/导出", "√", "√", ""],
            ["新增/编辑/提交", "", "√", ""], ["审核/退回/审阅", "√", "√", ""], ["查看流程和历史版本", "√", "√", ""],
        ],
        "gold_manage": [
            ["查询/查看", "√（只读）", "√", "√（本机构只读）"], ["新增/编辑/维护", "", "√", ""],
            ["提交审核", "", "√", ""], ["审核", "", "√", ""], ["导出/查看记录", "√（只读）", "√", "√（本机构只读）"],
        ],
        "indicator": [
            ["查询/查看", "√（只读）", "√", "√（本机构只读）"], ["新增/编辑/修改", "", "√", ""],
            ["版本维护/发布/生效/停用", "", "√", ""], ["导入/导出", "√（只读）", "√", "√（本机构）"],
            ["查看定义/版本/操作记录", "√（只读）", "√", "√（本机构只读）"],
        ],
        "major_event": [
            ["查询/查看事件", "√", "√", "√（本机构）"], ["新增/编辑草稿/删除草稿", "", "", "√（本机构）"],
            ["提交首报/续报/终报/处置进展", "", "", "√（本机构）"], ["核实/审核/退回/催办/跟踪", "√", "√", ""],
            ["审阅/确认关闭/归档", "√", "√", ""], ["查看流程和操作记录", "√", "√", "√（本机构）"],
        ],
        "special": [
            ["查询/查看", "√", "√", "√（本机构）"], ["新建/编辑/下发专项提示", "", "√", ""],
            ["机构反馈/跟踪反馈", "", "", "√（本机构）"], ["管理评估/解除意见", "√", "", ""],
            ["查看工单流程", "√", "√", "√（本机构）"],
        ],
        "periodic": [
            ["查询/查看/下载", "√", "√", "√（本机构）"], ["报告上传/编辑", "√", "√", "√（本机构）"],
            ["删除草稿", "√", "√", "√（本机构）"], ["提交", "√", "√", "√（本机构）"],
        ],
    }
    return profiles[profile]


def add_permission_section(doc, profile: str):
    add_prompt_heading(doc, "权限清单")
    add_word_table(doc, None, ["权限事项", "集团", "金控公司", "各金融机构"], permission_rows(profile), [28, 24, 24, 24])
    if profile == "risk_preference":
        add_word_table(doc, "角色权限说明", ["角色", "权限说明"], [
            ["集团", "查询、查看、审核、退回、审阅、导出并查看流程及历史版本。"],
            ["金控公司", "创建、编辑、提交、审核、发布和维护风险偏好方案。"],
            ["各金融机构", "该角色不展示本菜单，直接访问路由时不允许进入。"],
        ], [18, 82])


def add_fields_table(doc, page):
    rows = []
    current_area = None
    all_fields = []
    all_fields.extend(query_field(name) for name in page.get("queries", []))
    all_fields.extend(column_field(name) for name in page.get("columns", []))
    all_fields.extend(page.get("fields", []))
    if page.get("pagination"):
        all_fields.append(field_row("分页区", "分页", "数字", "链接操作", "总数、每页条数、页码", "展示当前数据总数，并支持每页条数、上一页、下一页及页面跳转。"))
    for area, name, kind, operation, value, note in all_fields:
        if area != current_area:
            rows.append([area])
            current_area = area
        rows.append([name, kind, operation, value, note])
    add_word_table(doc, "一、操作界面字段规则及要求", ["字段名称", "字段类型", "操作方式", "字段值", "备注说明"], rows, [20, 10, 15, 23, 32])


def add_operations_table(doc, page):
    operations = list(page.get("operations", []))
    if page.get("queries"):
        operations = ["查询", "重置"] + operations
    if page.get("pagination"):
        operations.append("分页")
    seen = set()
    rows = []
    for item in operations:
        if isinstance(item, tuple):
            name, description = item
        else:
            name, description = item, op_description(item)
        if name not in seen:
            seen.add(name)
            rows.append([name, description])
    add_word_table(doc, "二、操作功能说明", ["操作功能", "说明"], rows, [18, 82])


def add_status_table(doc, statuses):
    if statuses:
        add_word_table(doc, "三、状态说明", ["状态名称", "状态含义", "可执行角色", "可执行操作", "后续流转"], statuses, [15, 28, 17, 20, 20])


def add_flow_table(doc, flows):
    if flows:
        rows = [[str(i + 1), *flow] for i, flow in enumerate(flows)]
        add_word_table(doc, "四、流程说明", ["序号", "流程节点", "处理角色", "处理内容", "流转条件", "下一节点"], rows, [7, 20, 13, 26, 17, 17])


def add_page_detail(doc, page, num_id, page_break=True):
    heading = add_heading(doc, page["title"], 4, num_id)
    # The preceding maintenance table exactly fills its page in Word; forcing
    # another break before this page would create an otherwise empty page.
    if page["title"] == "预警规则配置 – 新增或编辑":
        page_break = False
    heading.paragraph_format.page_break_before = page_break
    add_text(doc, f"访问路由：{page['route']}。页面类型：{page['type']}。", align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    add_prompt_heading(doc, "界面样式：")
    add_placeholder(doc, page["title"])
    add_prompt_heading(doc, "界面规则")
    for paragraph in page["rules"]:
        add_text(doc, paragraph)
    add_permission_section(doc, page.get("permission", "all_view"))
    add_fields_table(doc, page)
    add_operations_table(doc, page)
    add_status_table(doc, page.get("statuses", []))
    add_flow_table(doc, page.get("flows", []))


def form_field(area: str, name: str, required=False, read_only=False, options="页面可选项"):
    if "附件" in name or "材料" in name:
        kind, operation, value = "附件", "文件上传" if not read_only else "链接操作", "文件名称、大小、上传时间"
    elif any(key in name for key in ("时间", "日期", "期限", "完成时间", "生效", "失效")):
        kind, operation, value = "时间", "自动反显" if read_only else "日期选择", "YYYY-MM-DD"
    elif any(key in name for key in ("状态", "类型", "机构", "频率", "周期", "方式", "结果", "等级")):
        kind, operation, value = "文本", "自动反显" if read_only else "下拉选择", options
    elif any(key in name for key in ("数量", "次数", "比例", "阈值", "值")):
        kind, operation, value = "数字", "自动反显" if read_only else "自主输入", "按字段口径填写"
    else:
        kind, operation, value = "文本", "自动反显" if read_only else "自主输入", "500个字符以内" if any(key in name for key in ("说明", "意见", "原因", "情况", "分析", "措施", "安排", "目标", "定义", "依据", "内容", "概述", "陈述")) else "100个字符以内"
    prefix = "必填" if required else "选填"
    note = f"展示{name}。" if read_only else f"{prefix}；填写{name}。"
    return field_row(area, name, kind, operation, value, note)


INDICATOR_COLUMNS = ["序号", "指标编码", "指标名称", "监测频率", "指标值", "当期亮灯情况", "累计亮灯情况", "指标定义", "指标类型", "指标子类", "适用机构", "预测范围", "黄灯规则", "红灯规则", "指标期次"]

MAJOR_EVENT_FLOW = [
    ["金融机构新建事件", "各金融机构", "填写事件首报及处置方案草稿。", "完成基础信息填写", "金融机构提交首报及处置方案"],
    ["金融机构提交首报及处置方案", "各金融机构", "校验必填信息并提交首报。", "提交成功", "金融机构管理部门核实并组织汇报"],
    ["金融机构管理部门核实并组织汇报", "集团", "核实事件情况并组织汇报。", "核实通过", "管理层审阅事件方案"],
    ["管理层审阅事件方案", "集团", "审阅事件处置方案。", "审阅通过", "董事会审阅事件方案"],
    ["董事会审阅事件方案", "集团", "按照管理权限审阅方案。", "审阅通过", "金融机构执行处置并跟踪"],
    ["金融机构执行处置并跟踪", "各金融机构", "执行处置方案并提交续报或进展。", "处置完成或需持续跟踪", "终报归档或常态跟踪"],
    ["终报归档或常态跟踪", "各金融机构", "提交终报，集团确认关闭归档或转入常态跟踪。", "审核确认", "流程结束"],
]

WARNING_FLOW = [
    ["原因分析填写", "各金融机构", "填写原因、影响和初步措施。", "提交原因分析", "原因分析审核"],
    ["原因分析审核", "金控公司", "审核原因分析，支持通过或退回。", "审核通过", "处置方案填写"],
    ["处置方案填写", "各金融机构", "填写目标、措施、责任部门、责任人和完成时间。", "提交方案", "处置方案审核"],
    ["处置方案审核", "金控公司", "审核方案并形成意见。", "审核通过", "管理协同或意见反馈"],
    ["管理协同或意见反馈", "集团", "开展协调、意见反馈或组织汇报。", "协同完成", "执行进展反馈"],
    ["执行进展反馈", "各金融机构", "按期提交执行情况和下一步计划。", "达到解除条件", "跟踪确认及解除"],
    ["跟踪确认及解除", "金控公司", "确认处置效果，解除或继续跟踪。", "确认解除", "流程结束"],
]

SPECIAL_FLOW = [
    ["专项风险提示", "金控公司", "创建提示并向选定机构下发独立工单。", "提交并下发", "机构首次反馈"],
    ["机构首次反馈", "各金融机构", "填写认识、自查问题和风险应对策略。", "提交反馈", "管理评估意见"],
    ["管理评估意见", "集团", "形成解除、部分解除或退回意见。", "评估完成", "跟踪反馈或解除意见"],
    ["跟踪反馈", "各金融机构", "对未解除事项继续反馈。", "再次提交", "管理评估意见"],
    ["解除意见及管理报告", "集团", "确认解除并形成管理报告。", "确认解除", "流程结束"],
]


def build_pages():
    pages = []
    pages.append({
        "module": "系统登录与通用框架", "title": "登录页面", "route": "/login", "type": "登录页面", "permission": "all_view",
        "rules": [
            "角色选择区–页面中部为系统角色选择区域——仅展示“集团”“金控公司”“各金融机构”三类角色；用户选择角色后方可登录。",
            "登录区–页面下方为登录操作区域——点击“登录系统”后，系统保存当前演示角色并进入对应默认页面；集团默认进入并表管理驾驶舱，其他角色默认进入工作台。",
        ],
        "fields": [form_field("角色选择区", "登录角色", required=True), form_field("登录区", "登录系统", read_only=True)],
        "operations": [("选择角色", "通过点击角色卡片，系统选中集团、金控公司或各金融机构作为当前登录角色。"), ("登录系统", "通过点击“登录系统”，系统进入当前角色的默认首页。")],
    })
    pages.append({
        "module": "并表驾驶舱", "title": "并表管理驾驶舱", "route": "/dashboard", "type": "综合看板", "permission": "dashboard",
        "rules": [
            "概览区–页面上方为集团整体风险概览区域——展示纳入机构、预警规则、重大风险事件、专项风险事项和历史指标数据等汇总信息；点击业务卡片进入对应业务页面。",
            "指标区–页面中部为重点指标监测区域——展示指标名称、指标值、指标期次、监测频率和当期亮灯；点击指标卡进入趋势详情或最新期指标状态。",
            "配置区–页面右侧为指标展示配置区域——支持按指标类型筛选并选择3至8项指标，保存后刷新仍保留；页面不设置独立快捷入口和待办事项。",
        ],
        "fields": [
            form_field("概览区", "纳入机构数", read_only=True), form_field("概览区", "预警灯号概览", read_only=True),
            form_field("概览区", "重大风险事件概览", read_only=True), form_field("概览区", "专项风险提示概览", read_only=True),
            form_field("指标区", "指标名称", read_only=True), form_field("指标区", "指标值", read_only=True),
            form_field("指标区", "指标期次", read_only=True), form_field("指标区", "监测频率", read_only=True),
            form_field("指标区", "当期亮灯", read_only=True), form_field("配置区", "按指标类型筛选"), form_field("配置区", "已选择"),
        ],
        "operations": ["查看全部", "进入台账", "查看详情", ("纳入机构数", "通过点击“纳入机构数”，系统进入纳入机构信息页面。"), ("历史指标数据", "通过点击“历史指标数据”，系统进入报表中心。"), ("指标卡", "通过点击指标卡，系统进入指标趋势详情或最新期指标状态。"), "保存", "恢复默认"],
    })
    pages.append({
        "module": "并表驾驶舱", "title": "纳入机构信息", "route": "/institutions", "type": "列表页面", "permission": "all_view", "pagination": True,
        "rules": [
            "查询区–页面上方为纳入机构查询栏——支持通过机构名称、机构类型、纳入状态和纳入时间进行筛选和查询。点击“查询”后，下方列表展示符合条件的结果；点击“重置”恢复默认状态。",
            "列表区–页面主体为纳入机构列表——展示机构名称、机构简称、机构类型、纳入时间、纳入状态、有效指标数、预警事项数和操作；各金融机构仅可查看本机构记录。",
            "分页区–页面下方为列表分页区域——展示当前数据总数，并支持选择每页展示条数、上一页、下一页及页面跳转。",
        ],
        "queries": ["机构名称", "机构类型", "纳入状态", "纳入时间"],
        "columns": ["序号", "机构名称", "机构简称", "机构类型", "纳入时间", "纳入状态", "有效指标数", "预警事项数", "操作"],
        "operations": ["查看", "分页"],
    })
    pages.append({
        "module": "并表驾驶舱", "title": "纳入机构信息 – 机构详情", "route": "/institutions/:id", "type": "路由弹窗详情页面", "permission": "all_view",
        "rules": [
            "页签区–页面上方设置“机构概览”“基础指标”“监管及外部评价”“集中度信息”“关联方信息”等页签——点击不同页签查看对应机构信息。",
            "机构概览区–页面主体为机构基本信息展示区域——以单列或合理两列方式展示机构全称、简称、类型、成立时间、注册资本、统一社会信用代码、股权结构、业务范围、管理机构情况、分支机构情况、地址、官网、应急联系人、联系方式、纳入时间和纳入状态；不展示独立组织架构图或占位区域。",
            "其他页签区–页面下方展示基础指标、监管及外部评价、集中度信息和关联方信息；点击关闭或返回后回到纳入机构列表。",
        ],
        "fields": [form_field("页签区", name, read_only=True) for name in ["机构概览", "基础指标", "监管及外部评价", "集中度信息", "关联方信息"]] + [form_field("机构概览区", name, read_only=True) for name in ["机构全称", "机构简称", "机构类型", "成立时间", "注册资本", "统一社会信用代码", "股权结构", "业务范围", "管理机构情况", "分支机构情况", "企业地址", "企业官网", "应急联系人", "联系方式", "纳入并表时间", "纳入状态"]],
        "operations": ["机构概览", "基础指标", "监管及外部评价", "集中度信息", "关联方信息", "关闭", "返回"],
    })
    pages.append({
        "module": "工作台", "title": "工作台", "route": "/workbench", "type": "综合工作页面", "permission": "all_view",
        "rules": [
            "待办区–页面主体为当前角色统一待办区域——按照当前角色和数据范围展示任务事项、所属机构、截止时间、状态和处理入口，支持按业务类型切换并执行批量处理、催办和转办。",
            "消息区–页面右侧为消息提醒区域——展示预警、流程、机构反馈和业务更新消息，支持查看、标记已读和提醒设置。",
            "运行区–页面下方为运行清单、任务进度、近期操作和快捷入口——展示各业务对象运行状态并支持进入对应业务页面；各金融机构仅展示本机构事项。",
        ],
        "columns": ["任务事项", "所属机构", "截止时间", "状态", "操作", "运行对象", "业务类型", "数量", "完成进度", "入口"],
        "fields": [form_field("页签区", "全部", read_only=True), form_field("页签区", "预警处置", read_only=True), form_field("页签区", "流程审阅", read_only=True), form_field("页签区", "报表确认", read_only=True)],
        "operations": ["处理", "批量处理", "催办", "转办", "刷新待办", "消息中心", "全部已读", "提醒设置", "查看", "返回驾驶舱"],
    })

    # 风险偏好及目标
    pages.append({
        "module": "预警管理/风险偏好及目标", "title": "风险偏好及目标", "route": "/warning/risk-preference", "type": "列表页面", "permission": "risk_preference", "pagination": True,
        "rules": [
            "查询区–页面上方为风险偏好方案查询栏——支持通过方案名称、覆盖机构、方案状态、生效起止时间和指标名称进行查询，并可仅展示运行中方案。",
            "操作区–页面主体为风险偏好方案列表——展示方案编号、名称、覆盖机构范围、生效及失效时间、状态、操作和流程；金控公司可新增或编辑，集团可审核、退回和审阅，各金融机构不展示本菜单且不得直接访问。",
            "分页区–页面下方为列表分页区域——展示总数并支持翻页。",
        ],
        "queries": ["方案名称", "覆盖机构", "方案状态", "生效起止时间", "指标名称"],
        "columns": ["序号", "方案编号", "方案名称", "覆盖机构范围", "生效时间", "失效时间", "方案状态", "操作", "方案流程"],
        "fields": [field_row("查询区", "仅展示运行中方案", "布尔", "复选", "选中或未选中", "选填；选中后仅展示当前生效方案。")],
        "operations": ["新增", "编辑", "查看", "查看流程", "导出", "分页"],
        "statuses": [
            ["待提交", "方案草稿尚未提交。", "金控公司", "编辑、保存、提交", "待审核"],
            ["待审核", "方案已提交等待审核。", "集团、金控公司", "审核、退回、查看", "已生效或待提交"],
            ["已生效", "方案审核并发布后生效。", "集团、金控公司", "查看、维护、导出", "已失效"],
            ["已失效", "方案超过有效期间或被新版本替代。", "集团、金控公司", "查看、导出", "流程结束"],
        ],
    })
    pages.append({
        "module": "预警管理/风险偏好及目标", "title": "风险偏好及目标 – 新建或编辑方案", "route": "/warning/risk-preference/new、/warning/risk-preference/:id/edit", "type": "路由弹窗表单页面", "permission": "risk_preference",
        "rules": [
            "基本信息区–页面上方为方案基本信息填写区域——支持填写方案名称、覆盖机构范围、生效及失效时间、编制部门、编制摘要和定性风险偏好陈述，方案编号和状态由系统自动生成或反显。",
            "指标区–页面中部为风险偏好指标配置区域——支持逐项新增风险类别、指标名称、定义、单位、适用机构、控制频率、预警值和容忍值，或批量导入指标。",
            "附件区–页面下方为方案附件上传区域——支持模拟上传并展示文件元数据；点击保存草稿或提交审阅后形成操作记录。",
        ],
        "fields": [form_field("基本信息区", name, required=name in ["方案名称", "覆盖机构范围", "生效时间", "失效时间", "定性风险偏好陈述"], read_only=name in ["方案编号", "方案状态"]) for name in ["方案编号", "方案状态", "方案名称", "覆盖机构范围", "生效时间", "失效时间", "编制部门", "编制摘要", "定性风险偏好陈述"]] + [form_field("指标明细区", name, required=name not in ["序号"]) for name in ["风险类别", "指标名称", "指标定义", "单位", "适用机构", "控制频率", "预警值", "容忍值"]] + [form_field("附件区", "方案附件")],
        "operations": ["新增指标", "批量导入", "删除", "保存草稿", "提交审阅", "返回"],
        "flows": [
            ["方案编制与提交", "金控公司", "编制方案并提交。", "必填信息校验通过", "方案审核与重检"],
            ["方案审核与重检", "集团", "审核方案并支持退回。", "审核通过", "管理层审阅及发布"],
            ["管理层审阅及发布", "集团", "审阅并发布方案。", "审阅通过", "流程结束"],
        ],
    })
    pages.append({
        "module": "预警管理/风险偏好及目标", "title": "风险偏好及目标 – 方案详情", "route": "/warning/risk-preference/:id", "type": "路由弹窗详情页面", "permission": "risk_preference",
        "rules": [
            "页签区–页面上方设置“概览”“风险偏好指标明细”“方案附件”等页签——点击不同页签查看方案基本信息、指标明细和附件。",
            "流程区–页面上方为方案流程展示区域——展示编制、审核、审阅和发布节点，支持查看历史处理记录及已提交表单。",
            "详情区–页面主体只读展示方案编号、名称、覆盖机构、生效期间、状态、编制部门、编制依据及指标明细。",
        ],
        "fields": [form_field("页签区", name, read_only=True) for name in ["概览", "风险偏好指标明细", "方案附件"]] + [form_field("基本信息区", name, read_only=True) for name in ["方案编号", "方案名称", "覆盖机构范围", "生效时间", "失效时间", "方案状态", "编制部门", "编制依据"]] + [form_field("指标明细区", name, read_only=True) for name in ["风险类别", "指标名称", "指标定义", "单位", "适用机构", "监测频率", "预警值", "容忍值", "生效时间"]],
        "operations": ["概览", "风险偏好指标明细", "方案附件", "查看流程", "查看历史版本", "导出", "返回"],
    })

    # 预警规则与处置
    pages.extend(build_warning_pages())
    pages.extend(build_indicator_pages())
    pages.extend(build_major_event_pages())
    pages.extend(build_report_pages())
    pages.extend(build_periodic_pages())
    pages.extend(build_special_risk_pages())
    pages.append({
        "module": "本期暂不展开模块", "title": "占位模块统一页面", "route": "/empty/capital、/empty/accounting、/empty/related、/empty/knowledge、/empty/system", "type": "占位页面", "permission": "all_view",
        "rules": [
            "占位区–页面主体为统一占位信息区域——资本规划与预算、会计并表管理、关联交易管理、知识库管理和系统管理在当前菜单中保留入口，但当前DEMO未提供可评审字段和业务操作。",
            "说明区–页面下方为范围说明——本期需求不虚构上述模块字段、流程和权限；待后续提供原型和业务规则后另行补充。",
        ],
        "fields": [form_field("占位区", "模块名称", read_only=True), form_field("占位区", "本期说明", read_only=True)],
        "operations": ["返回"],
    })
    return pages


def build_warning_pages():
    pages = []
    pages.append({
        "module": "预警管理/风险预警规则管理", "title": "风险预警规则管理", "route": "/warning/rules", "type": "列表页面", "permission": "gold_manage", "pagination": True,
        "rules": [
            "查询区–页面上方为风险预警规则查询栏——支持通过生效起止时间、指标风险类型、适用机构、规则状态、规则编号和指标名称进行筛选。",
            "操作区–页面主体为风险预警规则列表——展示指标、机构、监测频率、规则编号、规则类型、提示函下发方式、黄灯规则、红灯规则、生效时间、状态、维护、操作记录和配置入口。",
            "维护区–列表操作栏设置“维护”入口——金控公司点击后选择预警状态编辑或提示函下发方式；集团和各金融机构仅查看。",
        ],
        "queries": ["生效起止时间", "指标风险类型", "适用机构", "规则状态", "规则编号", "指标名称"],
        "columns": ["序号", "指标名称", "指标风险类型", "适用机构", "监测频率", "规则编号", "规则类型", "提示函下发方式", "黄灯规则", "红灯规则", "生效时间", "规则状态", "操作", "操作记录", "规则配置"],
        "operations": ["新增", "维护", "审核", "查看操作记录", "查看", "编辑", "导出", "分页"],
        "statuses": [
            ["草稿", "规则尚未提交。", "金控公司", "编辑、提交", "待审核"], ["待审核", "规则或变更等待审核。", "金控公司", "审核、退回", "生效或被退回"],
            ["生效", "规则处于正常监测状态。", "金控公司", "维护、查看", "暂停预警、停用或作废"], ["暂停预警", "暂时停止规则触发。", "金控公司", "恢复预警、停用", "生效或停用"],
            ["停用", "规则停止使用并保留历史。", "金控公司", "查看、恢复", "生效或作废"], ["作废", "规则永久退出运行。", "金控公司", "查看", "流程结束"],
        ],
    })
    pages.append({
        "module": "预警管理/风险预警规则管理", "title": "风险预警规则管理 – 维护类型选择", "route": "/warning/rules（维护弹窗）", "type": "弹窗", "permission": "gold_manage",
        "rules": [
            "维护选择区–弹窗主体为维护类型选择区域——提供“预警状态编辑”和“提示函下发方式”两项入口，并说明各自维护范围。",
            "关闭区–弹窗右上角设置关闭按钮——点击后返回规则列表，不保存任何变更。",
        ],
        "fields": [form_field("维护选择区", "预警状态编辑", read_only=True), form_field("维护选择区", "提示函下发方式", read_only=True)],
        "operations": ["预警状态编辑", "提示函下发方式", "关闭"],
    })
    pages.append({
        "module": "预警管理/风险预警规则管理", "title": "风险预警规则管理 – 预警状态维护", "route": "/warning/rules（状态维护弹窗）", "type": "弹窗表单", "permission": "gold_manage",
        "rules": [
            "基本信息区–弹窗上方为规则只读信息——展示规则编号、指标名称、适用机构和当前状态。",
            "维护区–弹窗主体为状态变更填写区域——支持选择生效、停用、暂停预警、恢复预警或作废，填写生效时间、变更原因并上传附件。",
            "提交区–弹窗底部为提交操作区域——勾选提交审核后形成状态变更待审核记录；审核通过前现行状态保持不变。",
        ],
        "fields": [form_field("基本信息区", name, read_only=True) for name in ["规则编号", "指标名称", "适用机构", "当前状态"]] + [form_field("维护区", "目标状态", required=True, options="生效、停用、暂停预警、恢复预警、作废"), form_field("维护区", "生效时间", required=True), form_field("维护区", "状态变更原因", required=True), form_field("附件区", "附件"), field_row("提交区", "提交审核", "布尔", "复选", "选中或未选中", "选填；选中后变更进入审核环节。")],
        "operations": ["提交审核", "取消"],
    })
    pages.append({
        "module": "预警管理/风险预警规则管理", "title": "风险预警规则管理 – 提示函下发方式维护", "route": "/warning/rules（提示函方式维护弹窗）", "type": "弹窗表单", "permission": "gold_manage",
        "rules": [
            "基本信息区–弹窗上方为规则及当前方式展示区域——展示规则编号、指标名称和当前提示函下发方式。",
            "方式选择区–弹窗主体为拟调整方式选择区域——支持选择“亮灯直接发函”“只亮灯不发函”“亮灯后手动发函”，并填写计划生效日期和变更原因。",
            "提交区–弹窗底部设置提交审核和取消——提交后形成待审核变更，审核通过前当前方式保持不变。",
        ],
        "fields": [form_field("基本信息区", name, read_only=True) for name in ["规则编号", "指标名称", "当前方式"]] + [form_field("方式选择区", "拟调整方式", required=True, options="亮灯直接发函、只亮灯不发函、亮灯后手动发函"), form_field("方式选择区", "计划生效日期", required=True), form_field("方式选择区", "变更原因", required=True)],
        "operations": ["提交审核", "取消"],
    })
    pages.append({
        "module": "预警管理/风险预警规则管理", "title": "预警规则配置 – 新增或编辑", "route": "/warning/rules/new/config、/warning/rules/:id/config/edit", "type": "路由弹窗表单页面", "permission": "gold_manage",
        "rules": [
            "基本信息区–页面上方为规则基本信息填写区域——支持选择适用机构、指标、风险类型、监测频率、规则类型、生效时间和规则级提示函下发方式。",
            "规则区–页面中部为预警值分段节点及区间赋灯区域——支持维护阈值节点、黄灯规则和红灯规则，并校验红灯条件的严重程度高于黄灯条件。",
            "提交区–页面下方为推送及提交区域——支持选择邮件、平台通知或其他现有推送方式，填写配置变更原因并提交审核。",
        ],
        "fields": [form_field("提示函下发方式", "提示函下发方式", required=True, options="亮灯直接发函、只亮灯不发函、亮灯后手动发函")] + [form_field("基本信息区", name, required=name in ["规则适用机构", "指标名称", "生效时间"]) for name in ["规则适用机构", "指标名称", "指标风险类型", "监测频率", "规则类型", "生效时间"]] + [form_field("规则区", name, required=True) for name in ["节点", "阈值", "节点赋灯", "区间范围", "区间预警灯", "黄灯规则", "红灯规则"]] + [form_field("提交区", "推送方式"), form_field("提交区", "提交人", read_only=True), form_field("提交区", "配置变更原因", required=True)],
        "operations": ["新增节点", "删除节点", "生成分段区间", "提交审核", "取消"],
    })
    pages.append({
        "module": "预警管理/风险预警规则管理", "title": "预警规则配置 – 查看", "route": "/warning/rules/:id/config", "type": "路由弹窗详情页面", "permission": "all_view",
        "rules": [
            "基本信息区–页面上方为规则配置只读区域——展示适用机构、指标、风险类型、监测频率、规则类型、生效时间和提示函下发方式。",
            "历史区–页面中部为历史阈值和版本区域——展示版本、生效日期、黄灯规则、红灯规则、状态和变更原因。",
            "操作记录区–页面下方为操作记录区域——展示配置、审核和生效全过程记录。",
        ],
        "fields": [form_field("基本信息区", name, read_only=True) for name in ["规则编号", "适用机构", "指标名称", "指标风险类型", "监测频率", "规则类型", "生效时间", "提示函下发方式"]],
        "columns": ["版本号", "生效日期", "黄灯规则", "红灯规则", "状态", "变更原因", "节点", "阈值", "区间", "区间赋灯", "节点赋灯"],
        "operations": ["查看操作记录", "返回"],
    })
    pages.append({
        "module": "预警管理/风险预警规则管理", "title": "预警规则审核", "route": "/warning/rules/:id/approve", "type": "路由弹窗审核页面", "permission": "gold_manage",
        "rules": [
            "申请信息区–页面上方为待审核变更展示区域——根据申请类型展示提示函方式变更、规则配置变更或状态变更内容，并与当前生效内容对照。",
            "审批区–页面下方为审批意见填写区域——金控公司填写审核意见并选择通过或退回；通过后按生效日期更新规则，退回后保留现行版本。",
        ],
        "fields": [form_field("申请信息区", name, read_only=True) for name in ["审核类型", "当前方式", "拟变更方式", "变更原因", "计划生效日期", "提交人", "提交时间", "附件"]] + [form_field("审批区", "审批意见", required=True)],
        "columns": ["版本", "生效日期", "黄灯规则", "红灯规则", "状态", "原因"],
        "operations": ["审核通过", "退回", "返回"],
    })
    pages.append({
        "module": "预警管理/预警提示与处置", "title": "预警提示与处置", "route": "/warning/disposal", "type": "列表页面", "permission": "all_view", "pagination": True,
        "rules": [
            "查询区–页面上方为预警事项查询栏——支持通过预警等级、所属机构、风险类型、指标名称、处理状态和预警时间筛选；各金融机构的所属机构固定为本机构。",
            "操作区–页面主体为预警事项列表——展示关联规则、规则版本、提示函方式、函件状态和处理状态；仅当方式为亮灯后手动发函且函件待下发时，金控公司可执行下发提示函。",
            "运行规则–规则为亮灯直接发函时函件状态为已下发且下发人为系统；规则为只亮灯不发函时状态为不发函且不显示手动下发按钮。",
        ],
        "queries": ["预警等级", "所属机构", "风险类型", "指标名称", "处理状态", "预警时间"],
        "columns": ["序号", "预警等级", "所属机构", "风险类型", "指标名称", "关联规则", "规则版本", "提示函方式", "提示函号", "函件状态", "处理状态", "预警时间", "操作"],
        "operations": ["查看详情", "下发提示函", "导出", "分页"],
    })
    pages.append({
        "module": "预警管理/预警提示与处置", "title": "预警提示与处置 – 预警处置工单", "route": "/warning/disposal/:id", "type": "路由弹窗流程页面", "permission": "all_view",
        "rules": [
            "流程区–页面上方为预警处置流程展示区域——展示原因分析、原因审核、方案填写、方案审核、管理协同、执行反馈和跟踪解除节点，当前节点突出显示。",
            "基本信息区–页面主体为工单及预警基本信息区域——展示提示函号、预警等级、机构、指标、指标值、规则、提示函方式、函件状态、处理状态和预警时间。",
            "办理区–页面下方为当前节点办理区域——按角色展示可填写表单、处理意见、附件及通过、退回、保存或提交按钮，责任部门为自主输入。",
        ],
        "fields": [form_field("基本信息区", name, read_only=True) for name in ["提示函号", "预警等级", "所属机构", "指标名称", "指标值", "关联规则", "规则版本", "提示函方式", "函件状态", "处理状态", "预警时间"]] + [form_field("办理区", name) for name in ["原因分析", "影响情况", "初步应对措施", "处置目标", "处置措施", "责任部门", "责任人", "计划完成时间", "处理意见", "本节点附件"]],
        "operations": ["查看流程", "查看历史节点", "保存", "提交", "通过", "退回", "返回"],
        "flows": WARNING_FLOW,
    })
    return pages


def build_indicator_pages():
    pages = []
    pages.append({
        "module": "指标管理/指标新增与维护", "title": "指标新增与维护", "route": "/indicators/maintenance", "type": "列表页面", "permission": "indicator", "pagination": True,
        "rules": [
            "查询区–页面上方为指标查询栏——支持通过指标编码、指标名称、指标状态和适用机构进行筛选。",
            "操作区–页面主体为指标列表——展示指标编码、名称、状态、定义、类型、子类、适用机构及操作；金控公司可新增、编辑和维护状态，集团及各金融机构仅查看权限范围内指标。",
            "分页区–页面下方为列表分页区域——展示当前数据总数并支持翻页。",
        ],
        "queries": ["指标编码", "指标名称", "指标状态", "适用机构"],
        "columns": ["序号", "指标编码", "指标名称", "指标状态", "指标定义", "指标类型", "指标子类", "适用机构", "操作"],
        "operations": ["新增", "编辑", "查看", "生效", "停用", "作废", "查看操作记录", "导出", "分页"],
        "statuses": [
            ["草稿", "指标尚未提交或发布。", "金控公司", "编辑、提交", "未生效"], ["未生效", "指标已建立但未到生效条件。", "金控公司", "编辑、生效", "生效"],
            ["生效", "指标可用于监测和规则配置。", "金控公司", "维护、停用", "停用"], ["停用", "指标停止用于新监测并保留历史。", "金控公司", "查看、恢复或作废", "生效或作废"],
            ["作废", "指标永久退出运行。", "金控公司", "查看", "流程结束"],
        ],
    })
    pages.append({
        "module": "指标管理/指标新增与维护", "title": "指标新增与维护 – 指标维护或查看", "route": "/indicators/maintenance/new、/indicators/maintenance/:id", "type": "路由弹窗表单或详情页面", "permission": "indicator",
        "rules": [
            "基本信息区–页面上方为指标基本信息填写或展示区域——金控公司可维护指标编码、名称、适用机构、类型、子类、管理类型、定义和生效日期；其他角色以只读方式查看。",
            "运行设置区–页面中部为监测频率设置区域——金控公司选择日、周、月、季、半年、年或不定期等频率。",
            "数据来源区–页面下方为源头表选择区域——金控公司选择现有数据来源并保存；系统写入指标版本和操作记录。",
        ],
        "fields": [form_field("基本信息区", name, required=name in ["指标名称", "适用机构", "指标类型", "指标定义", "生效起始日期"], read_only=name == "指标编码") for name in ["指标编码", "指标名称", "适用机构", "指标类型", "指标子类", "指标管理类型", "指标定义", "生效起始日期"]] + [form_field("运行设置区", "监测频率", required=True), form_field("数据来源区", "源头表", required=True)],
        "operations": ["保存", "查看操作记录", "返回"],
    })
    pages.append({
        "module": "指标管理/指标版本管理", "title": "指标版本管理", "route": "/indicators/versions", "type": "列表及弹窗页面", "permission": "indicator", "pagination": True,
        "rules": [
            "查询区–页面上方为指标版本查询栏——支持通过指标编码、指标名称和版本号查询。",
            "列表区–页面主体为指标当前版本列表——展示编码、名称、状态、定义、监测频率、当前版本号和查看历史版本入口。",
            "历史版本区–点击“查看历史版本”后弹窗展示版本号、维护操作、名称、定义、生效日期、适用机构、类型、源头表和停用日期，并支持下载。",
        ],
        "queries": ["指标编码", "指标名称", "版本号"],
        "columns": ["序号", "指标编码", "指标名称", "指标状态", "指标定义", "监测频率", "当前版本号", "操作", "版本号", "维护操作", "生效起始日期", "适用机构", "指标类型", "源头表", "停用日期"],
        "operations": ["查看历史版本", "下载", "关闭", "分页"],
    })
    pages.append({
        "module": "指标管理/最新期指标状态", "title": "最新期指标状态", "route": "/indicators/latest-status", "type": "列表页面", "permission": "all_view", "pagination": True,
        "rules": [
            "查询区–页面上方为最新期指标查询栏——支持通过指标编码、名称、类型、子类、适用机构、监测频率、当期亮灯情况和指标期次查询；不得设置指标状态查询条件。",
            "列表区–页面主体为最新期指标状态列表——每条记录仅对应一个指标和一个适用机构，并仅展示最新一期；指标值显示实际值和正确单位，预测范围统一显示“待测算”。",
            "亮灯规则–系统先判断红灯规则，再判断黄灯规则，均未触发时为绿灯；当期亮灯只能为红灯、黄灯或绿灯。累计亮灯单元格仅显示“查看”。",
        ],
        "queries": ["指标编码", "指标名称", "指标类型", "指标子类", "适用机构", "监测频率", "当期亮灯情况", "指标期次"],
        "columns": INDICATOR_COLUMNS,
        "operations": ["查看累计亮灯情况", "导出", "分页"],
    })
    pages.append({
        "module": "指标管理/最新期指标状态", "title": "最新期指标状态 – 累计亮灯情况", "route": "/indicators/latest-status（累计亮灯弹窗）", "type": "弹窗列表", "permission": "all_view",
        "rules": [
            "基本信息区–弹窗上方为指标及机构基本信息——展示指标编码、名称、适用机构和统计截至期次。",
            "列表区–弹窗主体为累计亮灯明细表——按照期次倒序展示同一指标、同一机构截至最新一期的全部历史记录。",
            "关闭区–弹窗底部设置关闭按钮——点击后返回最新期指标状态列表。",
        ],
        "fields": [form_field("基本信息区", name, read_only=True) for name in ["指标编码", "指标名称", "适用机构", "统计截至"]],
        "columns": ["序号", "指标期次", "监测频率", "指标值", "黄灯规则", "红灯规则", "当期亮灯情况"],
        "operations": ["关闭"],
    })
    return pages


def build_major_event_pages():
    pages = []
    pages.append({
        "module": "重大风险事件管理/重大风险事件定义管理", "title": "重大风险事件定义管理", "route": "/major-events/definitions", "type": "列表页面", "permission": "gold_manage", "pagination": True,
        "rules": [
            "查询区–页面上方为重大风险事件定义查询栏——支持通过事件类型名称、定义状态、参考依据和更新时间查询。",
            "列表区–页面主体为事件定义列表——初始展示重大操作风险事件、关注/异常项目事件、重大合规风险事件、重大声誉风险事件和重大信息科技风险事件，并展示编码、定义摘要、参考依据、状态、版本和更新时间。",
            "操作区–金控公司可新增、编辑、启用、停用和查看操作记录；集团及各金融机构只读。生效定义进入新增事件类型下拉框，停用定义保留历史引用。",
        ],
        "queries": ["事件类型名称", "定义状态", "参考依据", "更新时间"],
        "columns": ["序号", "事件类型编码", "事件类型名称", "事件定义摘要", "参考依据", "定义状态", "版本号", "更新时间", "操作", "操作记录"],
        "operations": ["新增", "编辑", "查看", "启用", "停用", "查看操作记录", "分页"],
        "statuses": [
            ["草稿", "定义尚未启用。", "金控公司", "编辑、启用", "生效"],
            ["生效", "定义可用于新建重大风险事件。", "金控公司", "编辑、停用", "停用"],
            ["停用", "定义不再用于新建事件，但历史事件保留原类型。", "金控公司", "查看、启用", "生效"],
        ],
    })
    pages.append({
        "module": "重大风险事件管理/重大风险事件定义管理", "title": "重大风险事件定义管理 – 查看定义", "route": "/major-events/definitions/:id", "type": "路由弹窗详情页面", "permission": "all_view",
        "rules": [
            "定义区–页面主体为重大风险事件定义只读区域——展示事件类型名称、判断标准和参考依据。",
            "记录区–页面下方为操作记录区域——展示定义新增、编辑、启用和停用记录。",
        ],
        "fields": [form_field("定义区", name, read_only=True) for name in ["事件类型编码", "事件类型名称", "事件定义判断标准", "参考依据", "定义状态", "版本号", "更新时间", "备注"]],
        "operations": ["查看操作记录", "返回"],
    })
    pages.append({
        "module": "重大风险事件管理/重大风险事件定义管理", "title": "重大风险事件定义管理 – 新增或编辑定义", "route": "/major-events/definitions/new、/major-events/definitions/:id/edit", "type": "路由弹窗表单页面", "permission": "gold_manage",
        "rules": [
            "基本信息区–页面上方为定义基本信息填写区域——事件类型编码和版本号由系统生成，金控公司填写事件类型名称、定义状态、参考依据和备注。",
            "判断标准区–页面下方为事件定义判断标准区域——支持逐条新增、编辑和删除判断标准；保存后写入操作记录。",
        ],
        "fields": [form_field("基本信息区", name, required=name in ["事件类型名称", "定义状态", "参考依据"], read_only=name in ["事件类型编码", "版本号"]) for name in ["事件类型编码", "事件类型名称", "定义状态", "版本号", "参考依据", "备注"]] + [form_field("判断标准区", "事件定义判断标准", required=True)],
        "operations": ["新增标准", "删除标准", "保存", "返回"],
    })
    pages.append({
        "module": "重大风险事件管理/重大风险事件列表", "title": "重大风险事件管理", "route": "/major-events", "type": "列表页面", "permission": "major_event", "pagination": True,
        "rules": [
            "查询区–页面上方为重大风险事件查询栏——支持通过事件名称、所属机构、风险事件类型、事件状态和发生时间查询；各金融机构自动限定为本机构。",
            "操作区–页面主体为重大风险事件列表——展示事件编号、名称、机构、类型、发生时间、最新报送类型、当前环节、状态、操作和流程；仅各金融机构显示“新增事件”。",
            "列表操作–各金融机构可编辑或删除未提交草稿并提交首报、续报、终报和处置进展；集团及金控公司按当前环节执行核实、审核、退回、催办、跟踪、审阅、关闭和归档。",
        ],
        "queries": ["事件名称", "所属机构", "风险事件类型", "事件状态", "发生时间"],
        "columns": ["序号", "事件编号", "事件名称", "所属机构", "风险事件类型", "发生时间", "最新报送类型", "当前环节", "事件状态", "操作", "事件流程"],
        "operations": ["新增", "编辑", "删除", "提交首报", "更新进展 / 续报", "提交终报", "核实", "审核", "退回", "催办", "跟踪", "审阅", "关闭归档", "查看详情", "查看流程", "分页"],
        "statuses": [
            ["草稿", "事件尚未提交。", "各金融机构", "编辑、删除、提交首报", "待审核"],
            ["待审核", "首报已提交等待核实或审核。", "集团、金控公司", "核实、审核、退回、催办", "处理中或草稿"],
            ["处理中", "事件处于方案执行和续报跟踪阶段。", "集团、金控公司、各金融机构", "续报、终报、催办、跟踪", "常态跟踪或已归档"],
            ["常态跟踪", "事件结束集中处置但仍需持续跟踪。", "集团、金控公司、各金融机构", "更新进展、跟踪、关闭", "已归档"],
            ["已归档", "事件完成终报并确认关闭。", "集团、金控公司、各金融机构", "查看、查看流程", "流程结束"],
        ],
    })
    pages.append({
        "module": "重大风险事件管理/重大风险事件列表", "title": "新增重大风险事件（首报及处置方案）", "route": "/major-events/new、/major-events/:id/edit", "type": "路由弹窗表单页面", "permission": "major_event",
        "rules": [
            "流程区–页面上方为七节点重大风险事件流程——新增时“金融机构新建事件”高亮，其他节点未完成。",
            "基本信息区–页面主体为事件基本信息填写区域——事件编号自动生成，状态和报送类型自动反显；风险事件类型读取当前生效定义，所属机构对各金融机构固定为本机构。",
            "事件与方案区–页面中部依次填写事件基本情况、已采取措施、初步分析研判、发展趋势、处置目标、处置措施、责任人、责任部门和计划完成时间；责任部门为自主输入。",
            "附件区–页面下方为模拟附件上传区域——展示附件名称、大小和上传时间；保存草稿状态为草稿，提交申请状态为待审核并返回列表。",
        ],
        "fields": [form_field("基本信息区", name, required=name in ["事件名称", "风险事件类型", "所属机构", "发生时间"], read_only=name in ["事件编号", "当前状态", "报送类型"]) for name in ["事件编号", "事件名称", "风险事件类型", "所属机构", "发生时间", "当前状态", "报送类型", "影响范围", "联系人", "联系电话"]] + [form_field("事件描述区", name, required=name == "事件基本情况") for name in ["事件基本情况", "已采取措施", "初步分析研判", "发展趋势"]] + [form_field("处置方案区", name) for name in ["处置目标", "处置措施", "责任人", "责任部门", "计划完成时间"]] + [form_field("附件区", "附件")],
        "operations": ["保存草稿", "提交申请", "返回", "关闭"],
        "flows": MAJOR_EVENT_FLOW,
    })
    pages.append({
        "module": "重大风险事件管理/重大风险事件列表", "title": "重大风险事件详情", "route": "/major-events/:id", "type": "路由弹窗详情及办理页面", "permission": "major_event",
        "rules": [
            "流程区–页面上方为重大风险事件流程展示区域——当前节点突出显示，已完成节点支持查看表单、处理意见和附件。",
            "详情区–页面主体保持现有详情布局——只读展示事件基本信息、首报信息、处置方案、续报记录、终报记录和审阅及处理记录。",
            "办理区–页面下方按当前角色和节点展示处理意见、附件以及核实、审核、退回、催办、跟踪、审阅和关闭归档操作。",
        ],
        "fields": [form_field("事件基本信息区", name, read_only=True) for name in ["事件编号", "事件名称", "所属机构", "风险事件类型", "发生时间", "最新报送类型", "当前环节", "事件状态", "影响范围", "联系人", "联系电话"]] + [form_field("首报信息区", name, read_only=True) for name in ["事件基本情况", "初步分析研判", "已采取措施", "发展趋势"]] + [form_field("处置方案区", name, read_only=True) for name in ["处置目标", "处置措施", "责任部门", "责任人", "计划完成时间"]] + [form_field("办理区", "处理意见"), form_field("办理区", "本节点附件")],
        "operations": ["查看流程", "更新进展 / 续报", "提交终报", "核实", "审核", "退回", "催办", "跟踪", "审阅", "关闭归档", "查看操作记录", "返回"],
        "flows": MAJOR_EVENT_FLOW,
    })
    pages.append({
        "module": "重大风险事件管理/重大风险事件列表", "title": "重大风险事件 – 续报", "route": "/major-events/:id/follow-up", "type": "路由弹窗表单页面", "permission": "major_event",
        "rules": [
            "当前事件区–页面上方为事件基本信息只读区域——展示事件编号、名称、所属机构、当前状态和最新报送类型。",
            "续报区–页面主体为续报填写区域——各金融机构填写最新进展、风险变化、措施执行情况和下一步安排并上传附件。",
            "提交区–页面底部设置提交续报和返回——提交后追加续报记录、更新最新报送类型并写入流程记录。",
        ],
        "fields": [form_field("当前事件区", name, read_only=True) for name in ["事件编号", "事件名称", "所属机构", "当前状态", "最新报送类型"]] + [form_field("续报区", name, required=True) for name in ["最新进展", "风险变化", "措施执行情况", "下一步安排"]] + [form_field("附件区", "附件")],
        "operations": ["提交续报", "返回"],
    })
    pages.append({
        "module": "重大风险事件管理/重大风险事件列表", "title": "重大风险事件 – 终报", "route": "/major-events/:id/final-report", "type": "路由弹窗表单页面", "permission": "major_event",
        "rules": [
            "终报区–页面主体为终报填写区域——各金融机构填写最终处置结果、实际影响、风险解除情况和后续管理安排并上传附件。",
            "提交区–页面底部设置提交终报和返回——提交后形成终报记录，后续由集团或金控公司确认关闭归档或转入常态跟踪。",
        ],
        "fields": [form_field("终报区", name, required=True) for name in ["最终处置结果", "实际影响", "风险解除情况", "后续管理安排"]] + [form_field("附件区", "附件")],
        "operations": ["提交终报", "返回"],
    })
    return pages


def build_report_pages():
    return [
        {
            "module": "报表中心", "title": "报表中心 – 历史指标数据查询", "route": "/reports", "type": "列表页面", "permission": "all_view", "pagination": True,
            "rules": [
                "查询区–页面上方为历史指标数据查询栏——支持通过指标编码、名称、类型、子类、适用机构、监测频率、当期亮灯情况和指标期次起止值查询；各金融机构自动限定为本机构。",
                "列表区–页面主体为全部往期指标记录——使用与最新期指标状态相同的数据逻辑，展示实际指标值、单位、当期亮灯、预测范围、具体黄灯及红灯规则和指标期次；预测范围统一显示“待测算”。",
                "累计亮灯区–累计亮灯单元格仅显示“查看”——点击后展示同一指标、同一机构截至所选期次的历史记录，顶部注明统计截至期次，不展示所选期次之后的数据。",
            ],
            "queries": ["指标编码", "指标名称", "指标类型", "指标子类", "适用机构", "监测频率", "当期亮灯情况", "指标期次起止值"],
            "columns": INDICATOR_COLUMNS,
            "operations": ["查看累计亮灯情况", "导出", "分页"],
        },
        {
            "module": "报表中心", "title": "报表中心 – 历史指标趋势", "route": "/reports/indicator/:recordId", "type": "路由弹窗详情页面", "permission": "all_view",
            "rules": [
                "基本信息区–页面上方为指标基本信息区域——展示指标编码、名称、机构、类型、子类、监测频率、预测范围、黄灯规则、红灯规则和指标定义。",
                "趋势区–页面主体为历史趋势和明细表——按照截至所选期次的历史记录展示指标值和当期亮灯情况。",
                "累计亮灯区–页面操作区设置查看累计亮灯情况——点击后以明细表展示截至当前期次的全部历史记录。",
            ],
            "fields": [form_field("基本信息区", name, read_only=True) for name in ["指标编码", "指标名称", "当前机构", "指标类型", "指标子类", "监测频率", "预测范围", "黄灯规则", "红灯规则", "指标定义"]],
            "columns": ["指标期次", "监测频率", "指标值", "黄灯规则", "红灯规则", "当期亮灯情况"],
            "operations": ["查看累计亮灯情况", "返回"],
        },
    ]


def build_periodic_pages():
    return [
        {
            "module": "定期报告", "title": "定期风险报告管理", "route": "/periodic-reports", "type": "列表页面", "permission": "periodic", "pagination": True,
            "rules": [
                "查询区–页面上方为定期风险报告查询栏——支持通过报告名称、报送日期起止、报送机构、报告类型和周期性查询。",
                "列表区–页面主体为定期风险报告列表——展示报送日期、报告日期、报告名称、报送机构、类型、周期、概述、附件和操作；各金融机构仅查看本机构。",
                "操作区–列表上方设置报告上传——支持新增报告；列表内支持查看下载、编辑和删除允许删除的草稿记录。",
            ],
            "queries": ["报告名称", "报送日期起止", "报送机构", "报告类型", "周期性"],
            "columns": ["报送日期", "报告日期", "定期风险报告名称", "报送机构", "报告类型", "周期性", "报告概述", "附件", "操作"],
            "operations": ["报告上传", "查看下载", "编辑", "删除", "分页"],
            "statuses": [["草稿", "报告已保存但尚未提交。", "当前数据范围内角色", "编辑、删除、提交", "已提交"], ["已提交", "报告已提交并保留记录。", "当前数据范围内角色", "查看、下载", "流程结束"]],
        },
        {
            "module": "定期报告", "title": "定期风险报告 – 报告上传或编辑", "route": "/periodic-reports/upload、/periodic-reports/:id/edit", "type": "路由弹窗表单页面", "permission": "periodic",
            "rules": [
                "当前报告区–页面上方为当前报告信息区域——编辑时反显报告名称、机构、类型和日期；新增时显示新建报告。",
                "基本信息区–页面主体为报告信息填写区域——支持填写报告名称、报送机构、报告类型、报告日期、周期性和报告概述。",
                "附件及推送区–页面下方为附件上传和推送设置区域——支持模拟上传并选择现有推送方式；保存草稿或提交后写入操作记录。",
            ],
            "fields": [form_field("当前报告区", name, read_only=True) for name in ["当前报告名称", "当前报送机构", "当前报告类型", "当前报告日期"]] + [form_field("基本信息区", name, required=name in ["报告名称", "报告类型", "报告日期", "周期性"]) for name in ["报告名称", "报送机构", "报告类型", "报告日期", "周期性", "报告概述"]] + [form_field("附件区", "附件"), form_field("推送设置区", "推送方式")],
            "operations": ["保存草稿", "提交", "关闭"],
        },
    ]


def build_special_risk_pages():
    pages = []
    pages.append({
        "module": "专项风险管理/新建专项风险提示", "title": "新建专项风险提示", "route": "/special-risks/drafts", "type": "草稿列表页面", "permission": "special", "pagination": True,
        "rules": [
            "查询区–页面上方为专项风险提示查询栏——支持通过创建起止时间、专项风险类型和下发对象查询。",
            "列表区–页面主体为专项风险提示草稿及已下发记录——展示创建时间、更新时间、对象、类型、名称、状态和操作，支持展开查看机构工单。",
            "操作区–金控公司可新建提示、编辑或删除草稿并下载；其他角色按照权限查看。",
        ],
        "queries": ["创建起止时间", "专项风险类型", "下发对象"],
        "columns": ["展开", "创建时间", "上次更新时间", "下发对象", "专项风险类型", "专项风险提示名称", "状态", "操作"],
        "operations": ["新建专项风险提示", "展开", "收起", "编辑", "删除", "下载", "分页"],
    })
    pages.append({
        "module": "专项风险管理/新建专项风险提示", "title": "专项风险提示 – 新建提示", "route": "/special-risks/new", "type": "路由弹窗表单页面", "permission": "special",
        "rules": [
            "基本信息区–页面上方为专项风险提示基本信息区域——金控公司填写提示名称和一个或多个专项风险类型。",
            "下发机构区–页面中部为下发机构选择区域——可选择五家纳入机构并分别填写补充信息；提交下发后每家机构生成独立工单。",
            "提示及附件区–页面下方填写提示背景与目的、管理措施与建议、机构自检与反馈要求并上传附件。",
        ],
        "fields": [form_field("基本信息区", "专项风险提示名称", required=True), form_field("基本信息区", "专项风险提示类型", required=True), form_field("下发机构区", "下发机构", required=True), form_field("下发机构区", "机构补充信息")] + [form_field("提示内容区", name, required=True) for name in ["提示背景与目的", "专项风险提示措施与建议", "机构自检与反馈要求"]] + [form_field("附件区", "附件")],
        "operations": ["保存草稿", "提交并下发", "关闭"],
        "flows": SPECIAL_FLOW,
    })
    pages.append({
        "module": "专项风险管理/专项风险提示与管理", "title": "专项风险提示与管理", "route": "/special-risks/manage", "type": "列表页面", "permission": "special", "pagination": True,
        "rules": [
            "查询区–页面上方为专项风险提示及工单查询栏——支持通过下发起止时间、专项风险类型、状态和下发对象查询。",
            "列表区–页面主体为提示及机构工单列表——展示工单编号、下发日期、对象、类型、名称、状态、解除日期和操作，并可展开查看每家机构工单。",
            "操作区–点击状态或查看提示进入详情及流程；有反馈记录时可查看管理报告。",
        ],
        "queries": ["下发起止时间", "专项风险类型", "状态", "下发对象"],
        "columns": ["展开", "工单编号", "下发日期", "下发对象", "专项风险类型", "提示名称", "状态", "解除日期", "操作"],
        "operations": ["展开", "收起", "查看提示", "查看管理报告", "查看流程", "分页"],
        "statuses": [
            ["待反馈", "机构尚未提交首次反馈。", "各金融机构", "反馈填写", "评估中"], ["已反馈", "机构已提交反馈。", "集团", "查看、评估", "评估中"],
            ["评估中", "集团正在开展管理评估。", "集团", "解除、部分解除、退回", "已解除、部分解除或待跟踪反馈"],
            ["部分解除", "部分风险点已解除，剩余事项继续跟踪。", "集团、各金融机构", "跟踪反馈、评估", "已解除或待跟踪反馈"],
            ["待跟踪反馈", "事项被退回或需继续跟踪。", "各金融机构", "跟踪反馈", "评估中"], ["已解除", "事项完成评估并解除。", "集团", "查看、形成管理报告", "流程结束"],
        ],
    })
    pages.append({
        "module": "专项风险管理/专项风险提示与管理", "title": "专项风险提示与工单详情", "route": "/special-risks/manage/:id", "type": "路由弹窗详情及流程页面", "permission": "special",
        "rules": [
            "提示区–页面上方为专项风险提示基本信息区域——展示提示名称、类型、对象、背景与目的、管理措施与建议、反馈要求和附件。",
            "流程区–页面中部为所选机构工单流程——展示提示下发、首次反馈、管理评估、跟踪反馈和解除意见节点，可查看已完成节点的表单、意见和附件。",
            "工单区–页面下方为全部机构工单列表——展示机构、工单编号、反馈时间、联系人、状态、评估结果和流程入口。",
        ],
        "fields": [form_field("提示区", name, read_only=True) for name in ["提示名称", "专项风险类型", "下发机构", "提示背景与目的", "管理措施与建议", "机构自检与反馈要求", "附件"]],
        "columns": ["机构", "工单编号", "反馈时间", "联系人", "状态", "评估结果", "操作"],
        "operations": ["查看流程", "进入反馈填写", "进入评估办理", "附件下载", "返回"],
        "flows": SPECIAL_FLOW,
    })
    pages.append({
        "module": "专项风险管理/专项风险查看与反馈", "title": "专项风险查看与反馈", "route": "/special-risks/feedback", "type": "列表页面", "permission": "special", "pagination": True,
        "rules": [
            "查询区–页面上方为机构工单查询栏——支持通过工单编号、下发起止时间、专项风险类型和状态查询；各金融机构仅查看本机构工单。",
            "列表区–页面主体为机构反馈工单列表——展示工单编号、下发日期、上次反馈时间、类型、提示名称、状态、反馈入口、解除日期和报告查看入口。",
        ],
        "queries": ["工单编号", "下发起止时间", "专项风险类型", "状态"],
        "columns": ["关联工单编号", "下发日期", "上次反馈时间", "专项风险类型", "提示名称", "查看提示", "状态", "反馈", "解除日期", "报告查看"],
        "operations": ["查看提示", "反馈填写", "跟踪反馈填写", "查看管理报告", "查看流程", "分页"],
    })
    pages.append({
        "module": "专项风险管理/专项风险查看与反馈", "title": "专项风险反馈填写", "route": "/special-risks/feedback/:id", "type": "路由弹窗表单页面", "permission": "special",
        "rules": [
            "提示区–页面上方为专项风险提示只读区域——展示名称、类型、下发机构及完整提示内容。",
            "反馈区–页面主体为机构反馈填写区域——各金融机构填写联系人、联系电话、认识与理解、自查问题与风险点、风险应对策略和其他总结反馈。",
            "附件区–页面下方为附件上传区域——提交后更新工单状态并写入流程记录。",
        ],
        "fields": [form_field("提示区", name, read_only=True) for name in ["提示名称", "专项风险类型", "下发机构", "提示背景与目的", "管理措施与建议", "自检与反馈要求"]] + [form_field("反馈区", name, required=name in ["联系人", "对专项风险提示的认识与理解", "自查发现的问题与风险点", "未来风险预防与应对策略"]) for name in ["反馈机构", "联系人", "联系电话", "对专项风险提示的认识与理解", "自查发现的问题与风险点", "未来风险预防与应对策略", "其他总结反馈"]] + [form_field("附件区", "附件")],
        "operations": ["保存", "关闭"],
        "flows": SPECIAL_FLOW,
    })
    pages.append({
        "module": "专项风险管理/专项风险查看与反馈", "title": "专项风险反馈评估", "route": "/special-risks/feedback/:id/evaluate", "type": "路由弹窗审核页面", "permission": "special",
        "rules": [
            "反馈信息区–页面上方为机构反馈只读区域——展示联系人、认识与理解、自查问题、风险应对策略、其他反馈和附件。",
            "评估区–页面主体为管理评估区域——集团选择解除、部分解除或驳回，填写评估部门和分析总结并上传附件。",
            "提交区–页面底部支持保存草稿和保存评估结果——解除后工单关闭，部分解除进入持续跟踪，驳回后返回跟踪反馈节点。",
        ],
        "fields": [form_field("反馈信息区", name, read_only=True) for name in ["反馈机构", "联系人", "联系电话", "认识与理解", "自查问题与风险点", "未来风险预防与应对策略", "其他总结反馈", "反馈附件"]] + [form_field("评估区", "评估结果", required=True, options="解除、部分解除、驳回"), form_field("评估区", "评估部门", required=True), form_field("评估区", "分析及总结", required=True), form_field("附件区", "评估附件")],
        "operations": ["保存草稿", "保存", "关闭"],
        "flows": SPECIAL_FLOW,
    })
    return pages


MENU_ROWS = [
    ["并表驾驶舱", "并表管理驾驶舱", "/dashboard", "正式功能"],
    ["工作台", "工作台", "/workbench", "正式功能"],
    ["预警管理", "风险偏好及目标", "/warning/risk-preference", "集团、金控公司可见"],
    ["预警管理", "风险预警规则管理", "/warning/rules", "正式功能"],
    ["预警管理", "预警提示与处置", "/warning/disposal", "正式功能"],
    ["指标管理", "指标新增与维护", "/indicators/maintenance", "正式功能"],
    ["指标管理", "指标版本管理", "/indicators/versions", "正式功能"],
    ["指标管理", "最新期指标状态", "/indicators/latest-status", "正式功能"],
    ["重大风险事件管理", "重大风险事件列表", "/major-events", "正式功能"],
    ["重大风险事件管理", "重大风险事件定义管理", "/major-events/definitions", "正式功能"],
    ["报表中心", "历史指标数据查询", "/reports", "正式功能"],
    ["定期报告", "定期风险报告管理", "/periodic-reports", "正式功能"],
    ["专项风险管理", "新建专项风险提示", "/special-risks/drafts", "正式功能"],
    ["专项风险管理", "专项风险提示与管理", "/special-risks/manage", "正式功能"],
    ["专项风险管理", "专项风险查看与反馈", "/special-risks/feedback", "正式功能"],
    ["资本规划与预算", "本期暂不展开", "/empty/capital", "统一占位"],
    ["会计并表管理", "本期暂不展开", "/empty/accounting", "统一占位"],
    ["关联交易管理", "本期暂不展开", "/empty/related", "统一占位"],
    ["知识库管理", "本期暂不展开", "/empty/knowledge", "统一占位"],
    ["系统管理", "本期暂不展开", "/empty/system", "统一占位"],
]


def add_cover_and_front_matter(doc):
    for _ in range(4):
        doc.add_paragraph()
    for line in ["上海国际集团", "系统功能需求说明书", "风险并表板块"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(line)
        set_run_font(run, 26, True)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("版本：V2.0")
    set_run_font(r, 14, False)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("2026年8月")
    set_run_font(r2, 14, False)

    doc.add_page_break()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("文档信息")
    set_run_font(run, 16, True)
    add_word_table(doc, None, ["信息项", "内容", "信息项", "内容"], [
        ["文档编号", "SHGJ-RISK-FRS-V2.0", "保密等级", "内部"],
        ["文档名称", "系统功能需求说明书", "业务板块", "风险并表板块"],
        ["适用系统", "集团并表管理系统", "文档版本", "V2.0"],
        ["编制单位", "上海国际集团项目组", "编制日期", "2026-08-03"],
    ], [18, 32, 18, 32])

    doc.add_page_break()
    revision = doc.add_paragraph()
    revision.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = revision.add_run("修订记录")
    set_run_font(rr, 16, True)
    revision_table = add_word_table(doc, None, ["日期", "版本", "修订说明", "修订人"], [
        ["2026-08-03", "V2.0", "根据最新DEMO页面原型、角色权限、菜单结构及业务功能全面修订。", "项目组"],
    ], [18, 14, 52, 16])
    for row in revision_table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, 12, run.bold)

    doc.add_page_break()
    add_toc_field(doc)
    doc.add_page_break()


def add_business_chapters(doc, num_id):
    add_heading(doc, "风险并表业务描述", 1, num_id)
    add_text(doc, "集团并表管理系统面向集团、金控公司及各金融机构，覆盖并表驾驶舱、工作台、风险偏好、预警规则、预警处置、指标管理、重大风险事件、历史指标数据、定期报告和专项风险管理。系统按照机构隔离和角色权限归集业务信息，支持配置、提交、审核、退回、跟踪、关闭和归档，形成风险识别、监测、预警、处置、报告和追溯的管理闭环，为集团掌握并表范围风险状况和推动风险处置提供支撑。")
    descriptions = [
        ("并表驾驶舱与纳入机构", "集团通过并表管理驾驶舱查看整体风险态势、机构分布、指标亮灯和重大事项；系统支持业务卡片下钻至纳入机构、预警规则、重大风险事件和历史指标数据，并允许按角色保存重点指标展示配置。"),
        ("工作台", "各类角色通过工作台集中查看本角色待办、消息、运行清单和任务处理进度，按照权限进入相应业务事项并执行处理、催办、转办或查看。"),
        ("风险偏好及目标", "金控公司创建、维护并提交风险偏好方案，集团开展审核、退回、审阅和发布；系统保留方案指标、附件、版本和全过程记录。"),
        ("风险预警与处置", "金控公司维护预警规则、阈值和规则级提示函下发方式；系统根据指标值先判断红灯规则、再判断黄灯规则，生成预警事项并推动原因分析、方案审核、执行反馈和解除。"),
        ("指标管理", "金控公司新增、编辑和维护指标及版本；集团和各金融机构按照数据范围查询指标。系统按指标与适用机构形成最新期和历史期记录，并展示指标值、单位、灯号及规则。"),
        ("重大风险事件", "各金融机构新建本机构重大风险事件并提交首报、续报和终报；集团及金控公司开展核实、审核、退回、催办、跟踪、审阅、关闭和归档，系统以生效事件定义控制新建类型。"),
        ("历史指标数据", "集团、金控公司和各金融机构通过报表中心查询权限范围内全部往期指标记录，查看截至所选期次的累计亮灯明细并导出查询结果。"),
        ("定期报告", "各角色按照当前页面权限维护或查看定期风险报告，上传附件、保存草稿或提交，并保留报告基本信息和操作记录。"),
        ("专项风险管理", "金控公司创建专项风险提示并向机构生成独立工单；各金融机构提交首次反馈和跟踪反馈；集团开展管理评估并形成解除意见和管理报告。"),
    ]
    for title, body in descriptions:
        add_text(doc, title, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=2)
        add_text(doc, body)

    add_heading(doc, "适用对象", 1, num_id)
    add_word_table(doc, None, ["角色", "使用范围", "数据范围", "主要职责与操作"], [
        ["集团", "集团并表风险管理与管理审阅。", "集团范围内全部机构数据。", "查询、查看、审核、退回、审阅、协调、催办、跟踪、关闭和归档；原则上不直接维护金控公司业务配置。"],
        ["金控公司", "并表业务方案、指标、规则、定义和参数管理。", "全部并表机构数据。", "新增、编辑、维护、发布、审核，跟踪预警处置并执行提示函相关操作。"],
        ["各金融机构", "本机构报送、反馈和风险处置。", "仅本机构数据。", "提交本机构事件报告、处置进展、专项风险反馈、相关材料和附件，不维护集团层面配置。"],
    ], [14, 24, 24, 38])

    add_heading(doc, "总体规则", 1, num_id)
    rules = [
        "系统最外层登录角色仅设置为集团、金控公司和各金融机构；业务流程中的部门或管理审阅名称不作为登录角色。",
        "各金融机构仅可查询、查看和处理本机构数据；集团和金控公司可查看全部并表机构数据。",
        "风险偏好及目标菜单仅向集团和金控公司展示，各金融机构直接访问相关路由时不得进入。",
        "指标新增、编辑、版本维护、发布、生效和停用仅由金控公司操作；集团和各金融机构只读。",
        "预警规则维护和规则级提示函下发方式配置仅由金控公司操作，所有变更保留版本和审批记录。",
        "指标亮灯判断按照红灯规则优先、黄灯规则次之、其他情况为绿灯的顺序执行；指标值、单位和规则必须能够支撑灯号。",
        "重大风险事件新增仅向各金融机构开放；事件类型读取当前生效定义，停用定义不用于新建但保留历史引用。",
        "所有业务含义明确为责任部门的字段均采用自主输入；支持多个部门时使用顿号或逗号分隔。",
        "报表中心仅用于历史指标数据查询，不承担文件报送或报表生成职能。",
        "系统保留各类提交、审核、退回、发布、状态变更、处置、反馈、关闭和归档记录，记录操作人、时间、内容、意见和附件。",
        "系统演示数据通过本地持久化机制保存，刷新后保留新增与修改记录；恢复演示数据属于演示环境辅助操作，不作为生产功能。",
    ]
    for item in rules:
        add_bullet(doc, item)

    add_heading(doc, "总体流程", 1, num_id)
    add_text(doc, "总体流程按照发起主体、填写或配置、提交、审核、退回或通过、生效或发布、后续跟踪、关闭或归档的主线组织。各业务流程的处理角色均使用三类系统角色，部门名称仅用于表示具体业务环节。")
    add_word_table(doc, "风险偏好流程", ["序号", "流程节点", "处理角色", "处理内容", "流转条件", "下一节点"], [["1", "方案编制与提交", "金控公司", "编制方案、指标和附件并提交。", "必填信息完整", "方案审核与重检"], ["2", "方案审核与重检", "集团", "审核方案并支持退回。", "审核通过", "管理层审阅及发布"], ["3", "管理层审阅及发布", "集团", "审阅并发布生效。", "审阅通过", "流程结束"]], [7, 20, 13, 26, 17, 17])
    add_flow_table(doc, WARNING_FLOW)
    add_flow_table(doc, MAJOR_EVENT_FLOW)
    add_flow_table(doc, SPECIAL_FLOW)


def subsection_label(primary: str, secondary: str | None):
    if secondary:
        return f"二级菜单—{secondary}"
    mapping = {
        "系统登录与通用框架": "登录与角色切换",
        "并表驾驶舱": "并表管理驾驶舱及机构下钻",
        "工作台": "工作台",
        "报表中心": "历史指标数据查询",
        "定期报告": "定期风险报告管理",
        "本期暂不展开模块": "统一占位说明",
    }
    return mapping.get(primary, primary)


def add_operation_rules(doc, num_id, pages):
    add_heading(doc, "操作规则", 1, num_id)
    add_text(doc, "本章按照当前菜单、路由和页面组织编写。每个页面依次说明界面样式、界面规则、权限、字段、操作、状态和流程；无法生成的当前页面截图使用统一灰色占位框，不使用原说明书旧截图。")
    last_primary = None
    last_secondary = None
    for page in pages:
        parts = page["module"].split("/", 1)
        primary = parts[0]
        secondary = parts[1] if len(parts) > 1 else None
        new_primary = primary != last_primary
        label = subsection_label(primary, secondary)
        new_secondary = label != last_secondary
        if new_primary:
            primary_heading = add_heading(doc, f"一级菜单——{primary}", 2, num_id)
            primary_heading.paragraph_format.page_break_before = True
            last_primary = primary
            last_secondary = None
            new_secondary = True
        if new_secondary:
            secondary_heading = add_heading(doc, label, 3, num_id)
            if not new_primary:
                secondary_heading.paragraph_format.page_break_before = True
            last_secondary = label
        add_page_detail(doc, page, num_id, page_break=not new_secondary)


def add_supplement(doc, num_id, pages):
    add_heading(doc, "附件及补充说明", 1, num_id)
    add_heading(doc, "菜单与路由清单", 2, num_id)
    add_text(doc, "当前系统共设置13个一级菜单，其中预警管理、指标管理、重大风险事件管理和专项风险管理包含11个可展开二级菜单。纳入机构信息由驾驶舱业务卡片下钻进入，不作为独立左侧菜单。")
    add_word_table(doc, None, ["一级菜单", "二级菜单或页面", "路由", "当前说明"], MENU_ROWS, [22, 30, 28, 20])

    add_heading(doc, "机构与数据范围", 2, num_id)
    add_word_table(doc, None, ["机构名称", "机构类型", "数据范围说明"], [
        ["浦发银行", "商业银行", "纳入并表范围，按照角色权限查询。"],
        ["国际AMC", "资产管理公司", "纳入并表范围；各金融机构演示角色默认对应本机构。"],
        ["太保集团", "保险集团", "纳入并表范围，按照角色权限查询。"],
        ["上农商", "商业银行", "纳入并表范围，按照角色权限查询。"],
        ["国泰海通", "证券公司", "纳入并表范围，按照角色权限查询。"],
    ], [26, 26, 48])

    add_heading(doc, "亮灯判断示例", 2, num_id)
    add_word_table(doc, None, ["指标", "风险方向", "黄灯规则", "红灯规则", "判断顺序"], [
        ["不良资产率", "数值越高风险越大", "≥2.00%", "≥5.00%", "先判断≥5.00%，未触发再判断≥2.00%。"],
        ["流动性覆盖率", "数值越低风险越大", "<120.00%", "<100.00%", "先判断<100.00%，未触发再判断<120.00%。"],
    ], [20, 22, 17, 17, 24])

    add_heading(doc, "截图补录清单", 2, num_id)
    add_text(doc, "当前运行环境未提供可用浏览器实例，因此本版未嵌入任何旧截图。下表所列正式页面均已在正文插入统一灰色占位框，后续应在当前DEMO环境按对应角色补录截图并保持原图比例。")
    screenshot_rows = [[str(i + 1), p["title"], p["route"], "统一灰色占位框"] for i, p in enumerate(pages)]
    add_word_table(doc, None, ["序号", "页面名称", "路由", "当前处理"], screenshot_rows, [8, 34, 38, 20])

    add_heading(doc, "实现差异与人工确认事项", 2, num_id)
    add_text(doc, "本说明书按照最新业务要求优先于当前代码实现的原则编制。当前代码仍存在少量历史机构别名归一化、历史文件报送页面残留及预测范围默认计算值；正文已统一采用五家标准机构、历史指标数据查询定位和“待测算”预测范围。上述代码差异未在本次文档任务中修改，后续开发验收时应按本说明书核对。")
    add_text(doc, "驾驶舱当前路由对三类角色均可访问并允许保存指标展示配置，但页面定位为集团领导视角；如生产系统需限制菜单可见范围，应在权限设计确认后同步调整。")
    add_text(doc, "定期报告页面在当前DEMO中未对上传、编辑和删除按钮进行更细的角色限制，正文按页面现状记录为权限范围内可操作；该规则需由业务和安全负责人在生产设计阶段确认。")


def main():
    global BULLET_NUM_ID
    DELIVERABLE_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(REFERENCE))
    configure_document(doc)
    num_id = create_heading_numbering(doc)
    BULLET_NUM_ID = num_id
    add_cover_and_front_matter(doc)
    add_business_chapters(doc, num_id)
    pages = build_pages()
    add_operation_rules(doc, num_id, pages)
    add_supplement(doc, num_id, pages)
    doc.save(str(OUTPUT))
    stats = WORK / "generation-stats.txt"
    primary_menus = 13
    secondary_menus = 11
    stats.write_text(
        "\n".join([
            f"output={OUTPUT}",
            f"primary_menus={primary_menus}",
            f"secondary_menus={secondary_menus}",
            f"written_pages={len(pages)}",
            "screenshots=0",
            f"screenshot_placeholders={len(pages)}",
        ]),
        encoding="utf-8",
    )
    print(stats.read_text(encoding="utf-8"))


if __name__ == "__main__":
    main()
