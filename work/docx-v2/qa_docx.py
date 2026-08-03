import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

DOCX = Path(r"C:\Users\Yiheng.jin\demo\deliverables\上海国际集团系统功能需求说明书_风险并表板块_V2.0.docx")
PROHIBITED = [
    "集团管理层", "集团金融机构管理部门经办岗", "集团金融机构管理部门审核岗",
    "金控公司经办岗", "金控公司审核岗", "并表金融机构经办岗", "并表金融机构审核岗",
    "黄规则", "红规则", "待配置", "未评价", "合理范围", "申能", "XXX", "待进一步确认",
]


def run_font(run):
    rpr = run._element.rPr
    if rpr is None:
        return None
    fonts = rpr.rFonts
    if fonts is None:
        return run.font.name
    return fonts.get(qn("w:eastAsia")) or fonts.get(qn("w:ascii")) or run.font.name


doc = Document(str(DOCX))
body_text = "\n".join(p.text for p in doc.paragraphs)
table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
all_text = body_text + "\n" + table_text
hits = {term: all_text.count(term) for term in PROHIBITED if term in all_text}

table_font_errors = []
table_size_counts = Counter()
for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for run in p.runs:
                    if not run.text.strip():
                        continue
                    font = run_font(run)
                    size = run.font.size.pt if run.font.size else None
                    table_size_counts[str(size)] += len(run.text)
                    if font != "华文细黑":
                        table_font_errors.append([ti, ri, ci, "font", font, run.text[:30]])
                    expected = 12 if ti == 1 else 11
                    if size != expected:
                        table_font_errors.append([ti, ri, ci, "size", size, run.text[:30]])

paragraph_font_errors = []
for pi, p in enumerate(doc.paragraphs):
    for run in p.runs:
        if not run.text.strip():
            continue
        font = run_font(run)
        size = run.font.size.pt if run.font.size else None
        allowed = {10, 14, 16, 26}
        if font != "华文细黑" or size not in allowed:
            paragraph_font_errors.append([pi, font, size, run.text[:50]])

with zipfile.ZipFile(DOCX) as zf:
    names = zf.namelist()
    xml_text = "\n".join(zf.read(name).decode("utf-8", errors="ignore") for name in names if name.endswith((".xml", ".rels")))
    package_hits = {term: xml_text.count(term) for term in PROHIBITED if term in xml_text}
    features = {
        "comments": any("comments" in name for name in names),
        "tracked_insertions": len(re.findall(r"<w:ins(?:\s|>)", xml_text)),
        "tracked_deletions": len(re.findall(r"<w:del(?:\s|>)", xml_text)),
        "toc_field": " TOC " in xml_text,
        "update_fields": "updateFields" in xml_text and ('w:val="true"' in xml_text or 'w:val="1"' in xml_text),
        "page_field": " PAGE " in xml_text,
        "media_files": len([name for name in names if name.startswith("word/media/")]),
    }

headings = Counter(p.style.name for p in doc.paragraphs if p.style.name.startswith("Heading"))
report = {
    "file": str(DOCX),
    "size": DOCX.stat().st_size,
    "sections": len(doc.sections),
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "headings": dict(headings),
    "text_hits": hits,
    "package_hits": package_hits,
    "table_font_errors": table_font_errors[:100],
    "table_font_error_count": len(table_font_errors),
    "table_size_counts": dict(table_size_counts),
    "paragraph_font_errors": paragraph_font_errors[:100],
    "paragraph_font_error_count": len(paragraph_font_errors),
    "features": features,
}
(DOCX.parent.parent / "work" / "docx-v2" / "qa-structure.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
