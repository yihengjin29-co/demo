import hashlib
import json
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

DOCX = Path(r"C:\Users\Yiheng.jin\demo\work\docx-v2\reference.docx")
OUT = DOCX.parent


def font_name(run):
    rpr = run._element.rPr
    if rpr is not None and rpr.rFonts is not None:
        return (
            rpr.rFonts.get(qn("w:eastAsia"))
            or rpr.rFonts.get(qn("w:ascii"))
            or rpr.rFonts.get(qn("w:hAnsi"))
        )
    return run.font.name


doc = Document(str(DOCX))
paragraphs = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    runs = []
    for r in p.runs:
        if r.text:
            runs.append({
                "text": r.text[:100],
                "font": font_name(r),
                "size_pt": r.font.size.pt if r.font.size else None,
                "bold": r.bold,
            })
    paragraphs.append({
        "index": i,
        "style": p.style.name,
        "text": text[:500],
        "alignment": str(p.alignment),
        "runs": runs[:6],
    })

tables = []
for i, table in enumerate(doc.tables):
    rows = []
    for row in table.rows[:8]:
        rows.append([cell.text.strip().replace("\n", " / ")[:300] for cell in row.cells])
    widths = []
    if table.rows:
        widths = [cell.width for cell in table.rows[0].cells]
    tables.append({
        "index": i,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "style": table.style.name if table.style else None,
        "widths_emu": widths,
        "sample": rows,
    })

headers_footers = []
for i, section in enumerate(doc.sections):
    headers_footers.append({
        "section": i + 1,
        "header": [p.text for p in section.header.paragraphs],
        "footer": [p.text for p in section.footer.paragraphs],
        "page_width": section.page_width,
        "page_height": section.page_height,
        "left_margin": section.left_margin,
        "right_margin": section.right_margin,
        "top_margin": section.top_margin,
        "bottom_margin": section.bottom_margin,
        "header_distance": section.header_distance,
        "footer_distance": section.footer_distance,
    })

with zipfile.ZipFile(DOCX) as zf:
    package = [{"path": n, "size": zf.getinfo(n).file_size, "sha256": hashlib.sha256(zf.read(n)).hexdigest()} for n in zf.namelist()]
    names = set(zf.namelist())

report = {
    "reference": str(DOCX),
    "sha256": hashlib.sha256(DOCX.read_bytes()).hexdigest(),
    "sections": len(doc.sections),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "paragraphs": paragraphs,
    "tables": tables,
    "headers_footers": headers_footers,
    "package": package,
    "features": {
        "comments": "word/comments.xml" in names,
        "footnotes": "word/footnotes.xml" in names,
        "endnotes": "word/endnotes.xml" in names,
        "numbering": "word/numbering.xml" in names,
        "settings": "word/settings.xml" in names,
        "custom_xml": any(n.startswith("customXml/") for n in names),
    },
}
(OUT / "template-inspection.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

font_counts = Counter()
size_counts = Counter()
for p in doc.paragraphs:
    for r in p.runs:
        if r.text:
            font_counts[font_name(r) or "(inherit)"] += len(r.text)
            size_counts[str(r.font.size.pt if r.font.size else "(inherit)")] += len(r.text)

print(json.dumps({
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "sections": len(doc.sections),
    "fonts": font_counts.most_common(8),
    "sizes": size_counts.most_common(8),
    "features": report["features"],
}, ensure_ascii=False, indent=2))
